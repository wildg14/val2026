"""Läsning av Valmyndighetens data.

Två källor:
- den kurerade filen majorna-valresultat-2022.xlsx (flikarna RD, RF, KF, Sammanfattning)
- Valmyndighetens rådatafiler "Röster per distrikt" (ett blad roster_RD/RF/KF i långformat:
  en rad per valdistrikt och parti, med summeringsrader och ogiltiga röster som egna rader)

Modulen är den enda platsen där partietiketter mappas till partikoder.
"""
import openpyxl

VAL = {"rd": "Riksdag", "rf": "Region", "kf": "Kommun"}

# kod -> (visningsnamn, etiketter i rådata, normaliserade till gemener)
PARTIER = {
    "V": ("Vänsterpartiet", ["vänsterpartiet"]),
    "S": ("Socialdemokraterna", ["arbetarepartiet-socialdemokraterna", "socialdemokraterna"]),
    "MP": ("Miljöpartiet", ["miljöpartiet de gröna", "miljöpartiet"]),
    "SD": ("Sverigedemokraterna", ["sverigedemokraterna"]),
    "M": ("Moderaterna", ["moderaterna"]),
    "C": ("Centerpartiet", ["centerpartiet"]),
    "L": ("Liberalerna", ["liberalerna (tidigare folkpartiet)", "liberalerna", "folkpartiet liberalerna"]),
    "KD": ("Kristdemokraterna", ["kristdemokraterna"]),
    "D": ("Demokraterna", ["demokraterna"]),
    "FI": ("Feministiskt initiativ", ["feministiskt initiativ"]),
    "K": ("Kommunistiska Partiet", ["kommunistiska partiet"]),
}
OVRIGA = "Övriga"

# Partier med egen kolumn per val. Allt annat läggs i Övriga. Samma lista som i den kurerade filen.
NYCKELPARTIER = {
    "rd": ["V", "S", "MP", "SD", "M", "C", "L", "KD"],
    "rf": ["V", "S", "MP", "M", "SD", "L", "C", "KD", "D", "FI"],
    "kf": ["V", "S", "MP", "M", "SD", "L", "D", "C", "KD", "FI", "K"],
}

MAJORNA_KODER = [str(14800526 + i) for i in range(23)]
AVGRANSNING = ("23 valdistrikt (14800526-14800548) i valkrets Västra Centrum = primärområdena "
               "Majorna, Stigberget, Kungsladugård och Sanna")

# Rader i rådatafilen som inte är partier
SUMMARADER = {"summa giltiga röster": "giltiga", "valdeltagande": "rostande"}
OGILTIGA = {"blanka röster", "övriga ogiltiga", "ej anmält deltagande"}
OVRIGA_ANMALDA = "övriga anmälda partier"

KOLUMNER = ["Valdistriktskod", "Valdistriktnamn", "Parti", "Röster", "Röstberättigade"]
KOLUMNER_VALFRIA = ["Kommun", "Län", "Val", "Region"]


class FormatFel(ValueError):
    """Rådatafilen har inte det format som förväntas."""


class SummaFel(ValueError):
    """Siffrorna i en rad går inte ihop."""


def _s(v):
    return "" if v is None else str(v).strip()


def partikod(etikett):
    e = _s(etikett).lower()
    for kod, (_, etiketter) in PARTIER.items():
        if e in etiketter:
            return kod
    return None


def partinamn():
    namn = {kod: v[0] for kod, v in PARTIER.items()}
    namn[OVRIGA] = "Övriga partier"
    return namn


# ---------------------------------------------------------------- kurerad xlsx

def las_kurerad(path):
    """Läser majorna-valresultat-2022.xlsx.

    Returnerar {"distrikt": {kod: {...}}, "total": {val: {...}}, "sammanfattning": {...}}.
    """
    wb = openpyxl.load_workbook(path, data_only=True)
    distrikt, total = {}, {}
    for val in VAL:
        ws = wb[val.upper()]
        rader = [r for r in ws.iter_rows(values_only=True) if any(v is not None for v in r)]
        huvud = [_s(h) for h in rader[0]]
        partier = huvud[2:huvud.index("Giltiga röster")]
        for r in rader[1:]:
            post = dict(zip(huvud, r))
            roster = {p: int(post[p]) for p in partier}
            summor = {"giltiga": int(post["Giltiga röster"]), "rostande": int(post["Röstande"]),
                      "rostberattigade": int(post["Röstberättigade"])}
            if post["Distriktskod"] is None:
                total[val] = {"roster": roster, **summor}
                continue
            kod = _s(post["Distriktskod"])
            d = distrikt.setdefault(kod, {"kod": kod, "namn": _s(post["Valdistrikt"]),
                                          "giltiga": {}, "rostande": {}, "rostberattigade": {}})
            d[val] = roster
            for k, v in summor.items():
                d[k][val] = v
    return {"distrikt": distrikt, "total": total, "sammanfattning": _las_sammanfattning(wb["Sammanfattning"])}


def _las_sammanfattning(ws):
    """Fliken Sammanfattning: andelar för Majorna, Göteborg och riket (region för RF)."""
    ut = {"majorna": {}, "goteborg": {}, "riket": {}}
    val = None
    for r in ws.iter_rows(values_only=True):
        forsta = _s(r[0])
        if forsta.startswith("Riksdagsvalet"):
            val = "rd"
        elif forsta.startswith("Regionvalet"):
            val = "rf"
        elif forsta.startswith("Kommunvalet"):
            val = "kf"
        if val is None or forsta in ("", "Parti") or forsta.startswith(("Riksdagsvalet", "Regionvalet", "Kommunvalet")):
            if forsta == "Parti":
                jamforelse = _s(r[4]).replace("Andel ", "") if len(r) > 4 and r[4] else None
                ut["majorna"][val] = {"roster": {}, "andel": {}}
                ut["goteborg"][val] = {"andel": {}}
                if jamforelse:
                    ut["riket"][val] = {"andel": {}, "namn": jamforelse}
            continue
        if forsta == "Giltiga röster":
            ut["majorna"][val]["giltiga"] = int(r[1])
        elif forsta == "Valdeltagande":
            ut["majorna"][val]["valdeltagande"] = float(r[1])
            ut["goteborg"][val]["valdeltagande"] = float(r[3])
            if val in ut["riket"] and r[4] is not None:
                ut["riket"][val]["valdeltagande"] = float(r[4])
        else:
            ut["majorna"][val]["roster"][forsta] = int(r[1])
            ut["majorna"][val]["andel"][forsta] = float(r[2])
            if r[3] is not None:
                ut["goteborg"][val]["andel"][forsta] = float(r[3])
            if val in ut["riket"] and r[4] is not None:
                ut["riket"][val]["andel"][forsta] = float(r[4])
    return ut


# ---------------------------------------------------------------- rådatafiler

def validera_huvud(huvud):
    """Kontrollerar att huvudraden innehåller de kolumner parsern behöver. Returnerar namn -> index."""
    rensat = [_s(h) for h in huvud]
    kol = {}
    for namn in KOLUMNER + KOLUMNER_VALFRIA:
        if namn in rensat:
            kol[namn] = rensat.index(namn)
    saknas = [n for n in KOLUMNER if n not in kol]
    if saknas:
        raise FormatFel(f"Rådatafilen saknar kolumnerna {saknas}. Huvud: {rensat}")
    return kol


def _tom_grupp():
    return {"roster": {}, "giltiga": 0, "rostande": 0, "ogiltiga": 0, "rostberattigade": 0, "koder": set()}


def _lagg_rad(post, etikett, roster):
    if etikett in SUMMARADER:
        post[SUMMARADER[etikett]] = roster
    elif etikett in OGILTIGA:
        post["ogiltiga"] += roster
    else:
        post["roster"][etikett] = post["roster"].get(etikett, 0) + roster


def _lagg_grupp(grupp, kod, etikett, roster, rostberattigade):
    if kod not in grupp["koder"]:
        grupp["koder"].add(kod)
        grupp["rostberattigade"] += rostberattigade
    if etikett in SUMMARADER:
        grupp[SUMMARADER[etikett]] += roster
    elif etikett in OGILTIGA:
        grupp["ogiltiga"] += roster
    else:
        grupp["roster"][etikett] = grupp["roster"].get(etikett, 0) + roster


def las_rafil(path, koder=None):
    """Läser en rådatafil. Returnerar råa poster (etiketter som i filen) för valda koder
    samt aggregat för riket, Göteborg och länet.

    {"val": "rd", "distrikt": {kod: {"kod", "namn", "roster": {etikett: n}, "giltiga", "rostande",
                                     "ogiltiga", "rostberattigade"}},
     "aggregat": {"riket": {...}, "goteborg": {...}, "lan": {...}}, "huvud": [...], "fil": str}
    """
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        blad = [ws for ws in wb.worksheets if ws.title.lower().startswith("roster_")]
        if not blad:
            raise FormatFel(f"{path}: hittar inget blad vars namn börjar med 'roster_'. Blad: {wb.sheetnames}")
        ws = blad[0]
        val = ws.title.split("_", 1)[1].lower()
        if val not in VAL:
            raise FormatFel(f"{path}: bladet {ws.title!r} motsvarar inget känt val (rd, rf, kf)")
        rader = ws.iter_rows(values_only=True)
        huvud = [_s(h) for h in next(rader)]
        kol = validera_huvud(huvud)
        vill = set(koder) if koder is not None else None
        distrikt = {}
        aggregat = {"riket": _tom_grupp(), "goteborg": _tom_grupp(), "lan": _tom_grupp()}
        for rad in rader:
            kod = _s(rad[kol["Valdistriktskod"]])
            if not kod:
                continue
            etikett = _s(rad[kol["Parti"]]).lower()
            roster = int(rad[kol["Röster"]] or 0)
            rostberattigade = int(rad[kol["Röstberättigade"]] or 0)
            grupper = ["riket"]
            if "Kommun" in kol and _s(rad[kol["Kommun"]]) == "Göteborg":
                grupper.append("goteborg")
            if "Län" in kol and _s(rad[kol["Län"]]) == "Västra Götaland":
                grupper.append("lan")
            for g in grupper:
                _lagg_grupp(aggregat[g], kod, etikett, roster, rostberattigade)
            if vill is not None and kod not in vill:
                continue
            post = distrikt.setdefault(kod, {"kod": kod, "namn": _s(rad[kol["Valdistriktnamn"]]), "roster": {},
                                             "giltiga": None, "rostande": None, "ogiltiga": 0,
                                             "rostberattigade": rostberattigade})
            _lagg_rad(post, etikett, roster)
        if not distrikt and vill:
            raise FormatFel(f"{path}: inget av de {len(vill)} önskade distrikten finns i filen")
        for g in aggregat.values():
            g["antal_distrikt"] = len(g.pop("koder"))
        return {"val": val, "distrikt": distrikt, "aggregat": aggregat, "huvud": huvud, "fil": str(path)}
    finally:
        wb.close()


def _mappa(roster_ra, nycklar):
    """Etikett -> partikod. Returnerar (roster per nyckelparti + Övriga, okända etiketter med röster)."""
    roster = {p: 0 for p in nycklar}
    roster[OVRIGA] = 0
    okanda = {}
    for etikett, n in roster_ra.items():
        kod = partikod(etikett)
        if kod in nycklar:
            roster[kod] += n
        else:
            roster[OVRIGA] += n
            if n and etikett != OVRIGA_ANMALDA:
                okanda[etikett] = n
    return roster, okanda


def till_distrikt(post, val, nyckelpartier=None):
    """Mappar en rå distriktspost till partikoder och kontrollerar summorna."""
    nycklar = nyckelpartier or NYCKELPARTIER[val]
    roster, okanda = _mappa(post["roster"], nycklar)
    namn = post.get("namn") or post.get("kod")
    if post.get("giltiga") is None:
        raise SummaFel(f"{namn}: raden 'Summa giltiga röster' saknas")
    if sum(roster.values()) != post["giltiga"]:
        raise SummaFel(f"{namn}: partiröster {sum(roster.values())} != giltiga {post['giltiga']}")
    if post.get("rostande") is None:
        raise SummaFel(f"{namn}: raden 'Valdeltagande' (röstande) saknas")
    if post["giltiga"] + post["ogiltiga"] != post["rostande"]:
        raise SummaFel(f"{namn}: giltiga {post['giltiga']} + ogiltiga {post['ogiltiga']} != röstande {post['rostande']}")
    return {"roster": roster, "giltiga": post["giltiga"], "rostande": post["rostande"],
            "rostberattigade": post["rostberattigade"], "okanda": okanda}


def aggregat_andelar(ra):
    """Andelar och valdeltagande för jämförelseområdena i en rådatafil.

    rd: riket och Göteborg. rf: Västra Götaland (som "riket") och Göteborg. kf: bara Göteborg.
    """
    val = ra["val"]
    ut = {}
    for namn, g in ra["aggregat"].items():
        if not g["giltiga"]:
            continue
        roster, _ = _mappa(g["roster"], NYCKELPARTIER[val])
        if sum(roster.values()) != g["giltiga"]:
            raise SummaFel(f"aggregat {namn}: partiröster {sum(roster.values())} != giltiga {g['giltiga']}")
        ut[namn] = {"andel": {p: n / g["giltiga"] for p, n in roster.items() if p != OVRIGA},
                    "valdeltagande": g["rostande"] / g["rostberattigade"] if g["rostberattigade"] else None,
                    "giltiga": g["giltiga"], "rostande": g["rostande"], "rostberattigade": g["rostberattigade"],
                    "antal_distrikt": g["antal_distrikt"]}
    lan = ut.pop("lan", None)
    if val == "rf" and lan:
        lan["namn"] = "Västra Götaland"
        ut["riket"] = lan
    elif val == "rd" and "riket" in ut:
        ut["riket"]["namn"] = "Riket"
    elif val == "kf":
        ut.pop("riket", None)
    return ut
