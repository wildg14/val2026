#!/usr/bin/env python3
"""granskning_2014: oberoende granskning av v2014-agentens filer för valet 2014.

Läser källfilerna (Excel, XML, dbf) med egen kod, utan att importera något ur
v2014-skripten, och jämför mot CSV-filerna i data/historik. Skriver inget till
data/historik. Alla avvikelser skrivs ut med prefixet FEL eller AVVIKELSE.

Körs med scratchpad-venv (xlrd, openpyxl, dbfread):
  venv/bin/python scripts/historik/granskning_2014_verifiering.py
"""
import csv
import os
import random
import sys
import xml.etree.ElementTree as ET
from collections import defaultdict, Counter

import openpyxl
import xlrd
from dbfread import DBF

KALLA_DIR = "/Users/daniel/code/Temp/Historiska dokument"
XLS_R = os.path.join(KALLA_DIR, "2014_riksdagsval_per_valdistrikt.xls")
XLS_L = os.path.join(KALLA_DIR, "2014_landstingsval_per_valdistrikt.xls")
XLSX_K = os.path.join(KALLA_DIR, "2014_kommunval_per_valdistrikt.xlsx")
SCRATCH = "/Users/daniel/code/Temp/Historiska dokument"
XML2014 = os.path.join(SCRATCH, "unz/slutresultat")
XML2010 = os.path.join(SCRATCH, "unz/slutresultat__1_")
DBF2014 = os.path.join(SCRATCH, "unz/valgeografi_valdistrikt/valgeografi_valdistrikt.dbf")
DBF2010 = os.path.join(SCRATCH, "unz/alla_valdistrikt/alla_valdistrikt.dbf")
UT_DIR = "/Users/daniel/code/Temp/data/historik"

VAL = {"rd": ("R", XLS_R), "rf": ("L", XLS_L), "kf": ("K", XLSX_K)}
MAJORNA_2014 = ["14801011", "14801012", "14801013", "14801014", "14801015", "14801016",
                "14801021", "14801031", "14801032", "14801033", "14801034", "14801035",
                "14801036", "14801041", "14801042", "14801043", "14801044"]
NORM = {"FP": "L", "DEM": "D", "KP": "K"}
ICKE_PARTI = ("OVR", "ÖVR", "BL", "BLANK", "OG")
SLUMPFRO = 20140914

FEL = []


def fel(msg):
    FEL.append(msg)
    print("FEL:", msg)


def norm(p):
    if p in NORM:
        return NORM[p]
    if p.isdigit():
        return p.zfill(4)
    return p.upper()


def col_letter(i):
    s = ""
    i += 1
    while i:
        i, r = divmod(i - 1, 26)
        s = chr(65 + r) + s
    return s


def read_csv(name):
    with open(os.path.join(UT_DIR, name), encoding="utf-8", newline="") as f:
        return list(csv.DictReader(f, delimiter=";"))


def to_int(v):
    if v is None or v == "":
        return 0
    return int(round(float(v)))


# ------------------------------------------------------------------ Excel
def read_excel(val):
    """Returnerar (bladnamn, rubriker, [(excelrad 1-baserad, celler)]) för Göteborg."""
    bokstav, path = VAL[val]
    if path.endswith(".xls"):
        wb = xlrd.open_workbook(path)
        sh = wb.sheet_by_index(0)
        hdr = [str(sh.cell_value(2, c)).strip() for c in range(sh.ncols)]
        rows = [(r + 1, [sh.cell_value(r, c) for c in range(sh.ncols)])
                for r in range(3, sh.nrows)]
        sheet = sh.name
    else:
        wb = openpyxl.load_workbook(path, read_only=True)
        ws = wb[wb.sheetnames[0]]
        raw = [list(r) for r in ws.iter_rows(values_only=True)]
        hdr = [str(h).strip() if h is not None else "" for h in raw[2]]
        rows = [(i + 1, r) for i, r in enumerate(raw) if i >= 3]
        sheet = ws.title
    gbg = [(rn, r) for rn, r in rows if to_int(r[0]) == 14 and to_int(r[1]) == 80]
    return sheet, hdr, gbg


def excel_distrikt(val):
    sheet, hdr, gbg = read_excel(val)
    if val == "rd":
        i_vd, i_namn = hdr.index("VALDIST"), hdr.index("Valdistrikt")
        i_kvk = hdr.index("Kommunvalkrets")
    else:
        i_vd, i_namn, i_kvk = 2, 5, None
    pcols = [(h[:-4], i, hdr.index(h[:-4] + " proc") if (h[:-4] + " proc") in hdr else None)
             for i, h in enumerate(hdr) if h.endswith(" tal")]
    out = {}
    for rn, r in gbg:
        kod = "1480%04d" % to_int(r[i_vd])
        if kod in out:
            fel("%s Excel dubblett kod %s" % (val, kod))
        partier = {p: (to_int(r[i]), r[j] if j is not None else None, i, j) for p, i, j in pcols}
        out[kod] = {
            "namn": str(r[i_namn]).strip(), "rad": rn, "partier": partier,
            "kvk": r[i_kvk] if i_kvk is not None else None,
            "giltiga": to_int(r[hdr.index("Rost Giltiga")]),
            "rostande": to_int(r[hdr.index("Rostande")]),
            "rostb": r[hdr.index("Rostb")], "vdt": r[hdr.index("VDT")],
        }
    return sheet, hdr, out


def cell_raw(val, excel_row, col_idx):
    """Läser en enskild råcell på nytt ur källfilen (rad 1-baserad, kolumn 0-baserad)."""
    path = VAL[val][1]
    if path.endswith(".xls"):
        sh = xlrd.open_workbook(path).sheet_by_index(0)
        return sh.cell_value(excel_row - 1, col_idx)
    wb = openpyxl.load_workbook(path, read_only=True)
    ws = wb[wb.sheetnames[0]]
    return ws.cell(row=excel_row, column=col_idx + 1).value


# -------------------------------------------------------------------- XML
def xml_partier(el):
    """Alla partirader inklusive ÖVRIGA_GILTIGA-barn, HANDSKRIVNA och ÖVRIGA_FGVAL."""
    rows = {}
    for g in el.findall("GILTIGA"):
        rows[g.get("PARTI")] = g
    ov = el.find("ÖVRIGA_GILTIGA")
    if ov is not None:
        for g in ov.findall("GILTIGA"):
            rows[g.get("PARTI")] = g
        for tag in ("HANDSKRIVNA", "ÖVRIGA_FGVAL"):
            h = ov.find(tag)
            if h is not None:
                rows[tag] = h
    return rows


def xml_summering(el):
    og = {o.get("TEXT"): o for o in el.findall("OGILTIGA")}
    vd = el.find("VALDELTAGANDE")
    return {
        "giltiga": int(el.get("RÖSTER")),
        "giltiga_fg": el.get("RÖSTER_FGVAL"),
        "blanka": int(og["BLANK"].get("RÖSTER")), "og": int(og["OG"].get("RÖSTER")),
        "blanka_fg": og["BLANK"].get("RÖSTER_FGVAL"), "og_fg": og["OG"].get("RÖSTER_FGVAL"),
        "rostande": int(vd.get("SUMMA_RÖSTER")) if vd is not None else None,
        "rostb": int(vd.get("RÖSTBERÄTTIGADE")) if vd is not None else None,
        "vdt": vd.get("PROCENT").replace(",", ".") if vd is not None else None,
        "rostande_fg": vd.get("SUMMA_RÖSTER_FGVAL") if vd is not None else None,
        "rostb_fg": vd.get("RÖSTBERÄTTIGADE_KLARA_VALDISTRIKT_FGVAL") if vd is not None else None,
    }


def xml_distrikt(xmldir, val):
    root = ET.parse(os.path.join(xmldir, "slutresultat_1480%s.xml" % VAL[val][0])).getroot()
    out = {}
    krets_av = {}
    kommun = root.find("KOMMUN")
    for kk in kommun.findall("KRETS_KOMMUN"):
        for d in kk.iter("VALDISTRIKT"):
            out[d.get("KOD")] = d
            krets_av[d.get("KOD")] = kk.get("NAMN")
        for d in kk.findall("ONSDAGSDISTRIKT"):
            kod = "1480" + d.get("KOD").split("-")[-1].zfill(4)
            out[kod] = d
            krets_av[kod] = kk.get("NAMN")
    n_all = sum(1 for _ in root.iter("VALDISTRIKT")) + sum(1 for _ in root.iter("ONSDAGSDISTRIKT"))
    if n_all != len(out):
        fel("%s XML: %d distrikt i trädet men %d under KRETS_KOMMUN" % (val, n_all, len(out)))
    return root, out, krets_av


# ------------------------------------------------------------- kontroll 1
def kontroll_1(val, xd_excel, root, xd_xml, krets_av):
    print("\n--- (1) %s Göteborgs totaler per parti" % val)
    agg = read_csv("aggregat_2014_%s.csv" % val)
    agg_g = {r["parti"]: r for r in agg if r["niva"] == "goteborg"}
    if len(agg_g) != sum(1 for r in agg if r["niva"] == "goteborg"):
        fel("%s aggregat goteborg har dubbla parti-nycklar" % val)
    kommun = root.find("KOMMUN")
    xk = {p: int(g.get("RÖSTER")) for p, g in xml_partier(kommun).items()}
    # Excelsumma
    tot_x = defaultdict(int)
    tot_x_krets = defaultdict(lambda: defaultdict(int))
    for kod, d in xd_excel.items():
        for p, (n, a, i, j) in d["partier"].items():
            tot_x[p] += n
            tot_x_krets[krets_av[kod]][p] += n
    # XML-distriktssumma
    tot_xml = defaultdict(int)
    for kod, el in xd_xml.items():
        for p, g in xml_partier(el).items():
            tot_xml[p] += int(g.get("RÖSTER"))
    # CSV-summor
    tot_csv = {}
    for src in ("xls", "xml"):
        t = defaultdict(int)
        for r in read_csv("roster_2014_%s_%s.csv" % (val, src)):
            t[r["parti"]] += int(r["roster"])
        tot_csv[src] = t
    n = 0
    for p, s in tot_x.items():
        if p in ICKE_PARTI:
            continue
        n += 1
        np_ = norm(p)
        if np_ not in agg_g:
            if s:
                fel("%s aggregat goteborg saknar %s (Excelsumma %d)" % (val, p, s))
            continue
        if int(agg_g[np_]["roster"]) != s:
            fel("%s aggregat goteborg %s=%s, Excelsumma %d" % (val, np_, agg_g[np_]["roster"], s))
        if tot_csv["xls"].get(np_) != s:
            fel("%s roster_xls-summa %s=%s, Excelsumma %d" % (val, np_, tot_csv["xls"].get(np_), s))
    for p, s in tot_xml.items():
        np_ = norm(p)
        if xk.get(p) != s:
            fel("%s XML KOMMUN %s=%s, distriktssumma %d" % (val, p, xk.get(p), s))
        if p == "ÖVRIGA_FGVAL":
            if s != 0:
                fel("%s ÖVRIGA_FGVAL RÖSTER %d i aktuellt val" % (val, s))
            continue
        if np_ not in agg_g or int(agg_g[np_]["roster"]) != s:
            fel("%s aggregat goteborg %s=%s, XML-distriktssumma %d" % (val, np_, agg_g.get(np_, {}).get("roster"), s))
        if tot_csv["xml"].get(np_) != s:
            fel("%s roster_xml-summa %s=%s, XML-distriktssumma %d" % (val, np_, tot_csv["xml"].get(np_), s))
    for np_ in agg_g:
        if np_ not in {norm(p) for p in tot_xml}:
            fel("%s aggregat goteborg har parti %s som inte finns i XML" % (val, np_))
    ovr = tot_x.get("OVR", tot_x.get("ÖVR"))
    excel_p = {norm(p) for p in tot_x}
    rest = sum(v for p, v in tot_xml.items() if norm(p) not in excel_p and p != "ÖVRIGA_FGVAL")
    if ovr != rest:
        fel("%s Excel ÖVR-summa %d, XML övriga + handskrivna %d" % (val, ovr, rest))
    print("%s: %d Excelpartier och %d XML-partier jämförda mot aggregat goteborg; ÖVR %d = XML-rest %d"
          % (val, n, len(tot_xml) - 1, ovr, rest))
    # aggregat: andel, giltiga, röstande, röstberättigade per nivå
    ks = xml_summering(kommun)
    for r in agg:
        g = int(r["giltiga"])
        calc = round(100.0 * int(r["roster"]) / g, 2)
        if abs(calc - float(r["andel"])) > 0.011:
            fel("%s aggregat %s %s andel %s, beräknat %.2f" % (val, r["niva"], r["parti"], r["andel"], calc))
    a0 = agg_g["M"]
    if (int(a0["giltiga"]), int(a0["rostande"]), int(a0["rostberattigade"])) != (ks["giltiga"], ks["rostande"], ks["rostb"]):
        fel("%s aggregat goteborg giltiga/rostande/rostb %s/%s/%s, XML %s/%s/%s" % (
            val, a0["giltiga"], a0["rostande"], a0["rostberattigade"], ks["giltiga"], ks["rostande"], ks["rostb"]))
    dist = read_csv("distrikt_2014_%s.csv" % val)
    g_sum = sum(int(r["giltiga"]) for r in dist)
    r_sum = sum(int(r["rostande"]) for r in dist)
    b_sum = sum(int(r["rostberattigade"] or 0) for r in dist)
    b_x = sum(to_int(d["rostb"]) for d in xd_excel.values())
    print("%s giltiga: Excel %d, distrikt.csv %d, XML KOMMUN %d, aggregat %s" % (
        val, sum(d["giltiga"] for d in xd_excel.values()), g_sum, ks["giltiga"], a0["giltiga"]))
    print("%s rostande: Excel %d, distrikt.csv %d, XML KOMMUN %d, aggregat %s" % (
        val, sum(d["rostande"] for d in xd_excel.values()), r_sum, ks["rostande"], a0["rostande"]))
    print("%s rostberattigade: Excel %d, distrikt.csv %d (297 distrikt, uppsamling tomt), XML KOMMUN %d, aggregat %s"
          % (val, b_x, b_sum, ks["rostb"], a0["rostberattigade"]))
    if not (g_sum == ks["giltiga"] and r_sum == ks["rostande"] and b_sum == ks["rostb"] == b_x):
        fel("%s summor i distrikt.csv skiljer sig från XML KOMMUN" % val)
    # kretsnivåer
    per_krets = defaultdict(lambda: [0, 0, 0])
    for r in dist:
        k = per_krets[r["valkrets"]]
        k[0] += int(r["giltiga"]); k[1] += int(r["rostande"]); k[2] += int(r["rostberattigade"] or 0)
    kretsar_xml = {kk.get("NAMN"): kk for kk in kommun.findall("KRETS_KOMMUN")}
    for r in agg:
        if r["niva"] in kretsar_xml:
            kk = kretsar_xml[r["niva"]]
            xs = xml_summering(kk)
            xp = {norm(p): int(g.get("RÖSTER")) for p, g in xml_partier(kk).items()}
            if xp.get(r["parti"]) != int(r["roster"]):
                fel("%s aggregat %s %s=%s, XML KRETS %s" % (val, r["niva"], r["parti"], r["roster"], xp.get(r["parti"])))
            if r["parti"] not in ICKE_PARTI and tot_x_krets[r["niva"]].get({"L": "FP"}.get(r["parti"], r["parti"])) not in (None, int(r["roster"])):
                fel("%s aggregat %s %s=%s, Excelsumma per krets %s" % (val, r["niva"], r["parti"], r["roster"],
                    tot_x_krets[r["niva"]].get(r["parti"])))
            if (int(r["giltiga"]), int(r["rostande"]), int(r["rostberattigade"])) != (xs["giltiga"], xs["rostande"], xs["rostb"]):
                fel("%s aggregat %s giltiga/rostande/rostb skiljer från XML" % (val, r["niva"]))
            if (int(r["giltiga"]), int(r["rostande"]), int(r["rostberattigade"])) != tuple(per_krets[r["niva"]]):
                fel("%s aggregat %s giltiga/rostande/rostb %s, distriktssumma %s" % (
                    val, r["niva"], (r["giltiga"], r["rostande"], r["rostberattigade"]), per_krets[r["niva"]]))
    nivaer = Counter(r["niva"] for r in agg)
    print("%s aggregat nivåer: %s" % (val, dict(nivaer)))
    # riket och vgregion
    root00 = ET.parse(os.path.join(XML2014, "slutresultat_00%s.xml" % VAL[val][0])).getroot()
    nat = root00.find("NATION")
    lan = [l for l in root00.iter("LÄN") if l.get("KOD") == "14"][0]
    for niva, el in (("riket", nat), ("vgregion", lan)):
        xp = {norm(p): int(g.get("RÖSTER")) for p, g in xml_partier(el).items()}
        xs = xml_summering(el)
        rows = [r for r in agg if r["niva"] == niva]
        for r in rows:
            if xp.get(r["parti"]) != int(r["roster"]):
                fel("%s aggregat %s %s=%s, 00-XML %s" % (val, niva, r["parti"], r["roster"], xp.get(r["parti"])))
            if (int(r["giltiga"]), int(r["rostande"]), int(r["rostberattigade"])) != (xs["giltiga"], xs["rostande"], xs["rostb"]):
                fel("%s aggregat %s giltiga/rostande/rostb skiljer från 00-XML" % (val, niva))
        saknas = [p for p in xp if p not in {r["parti"] for r in rows} and p != "ÖVRIGA_FGVAL" and xp[p] > 0]
        if saknas:
            fel("%s aggregat %s saknar partier %s" % (val, niva, saknas))
        print("%s %s: %d rader, XML giltiga %s (%s)" % (val, niva, len(rows), xs["giltiga"], el.get("NAMN")))


# ------------------------------------------------------------- kontroll 2
def kontroll_2(val):
    print("\n--- (2) %s aritmetik per distrikt i CSV" % val)
    dist = read_csv("distrikt_2014_%s.csv" % val)
    koder = [d["kod"] for d in dist]
    if len(set(koder)) != len(koder):
        fel("%s distrikt.csv har dubbla koder" % val)
    gilt = {d["kod"]: int(d["giltiga"]) for d in dist}
    for src in ("xls", "xml"):
        roster = read_csv("roster_2014_%s_%s.csv" % (val, src))
        nyckel = Counter((r["kod"], r["parti"]) for r in roster)
        dubbla = [k for k, c in nyckel.items() if c > 1]
        if dubbla:
            fel("%s roster_%s dubbla nycklar: %s" % (val, src, dubbla[:5]))
        s = defaultdict(int)
        for r in roster:
            s[r["kod"]] += int(r["roster"])
            if r["ar"] != "2014" or r["val"] != val:
                fel("%s roster_%s fel ar/val i rad %s" % (val, src, r))
        n_ok = sum(1 for d in dist if s[d["kod"]] == int(d["giltiga"]))
        for d in dist:
            if s[d["kod"]] != int(d["giltiga"]):
                fel("%s roster_%s %s partisumma %d != giltiga %s" % (val, src, d["kod"], s[d["kod"]], d["giltiga"]))
        if set(s) != set(koder):
            fel("%s roster_%s kodmängd skiljer från distrikt.csv" % (val, src))
        n_a = n_tom = n_tom_med_roster = n_f = 0
        for r in roster:
            if r["andel"] == "":
                n_tom += 1
                if int(r["roster"]) != 0:
                    n_tom_med_roster += 1
                continue
            n_a += 1
            calc = round(100.0 * int(r["roster"]) / gilt[r["kod"]], 2) if gilt[r["kod"]] else 0.0
            if abs(calc - float(r["andel"])) > 0.011:
                n_f += 1
                if n_f <= 3:
                    fel("%s roster_%s %s %s andel %s, beräknat %.2f" % (val, src, r["kod"], r["parti"], r["andel"], calc))
        print("%s roster_%s: %d rader, partisumma = giltiga i %d/%d, andel kontrollerad i %d rader (%d tomma, varav %d med röster), %d avvikelser"
              % (val, src, len(roster), n_ok, len(dist), n_a, n_tom, n_tom_med_roster, n_f))
        if n_tom_med_roster:
            fel("%s roster_%s: %d rader med röster men tom andel" % (val, src, n_tom_med_roster))
    n_vdt = 0
    for d in dist:
        g, b, o, og, r = (int(d[k]) for k in ("giltiga", "blanka", "ogiltiga_ovriga", "ogiltiga", "rostande"))
        if b + o != og:
            fel("%s %s blanka + ogiltiga_ovriga %d != ogiltiga %d" % (val, d["kod"], b + o, og))
        if g + og != r:
            fel("%s %s giltiga + ogiltiga %d != rostande %d" % (val, d["kod"], g + og, r))
        if d["rostberattigade"]:
            rb = int(d["rostberattigade"])
            if "." in d["rostberattigade"]:
                fel("%s %s rostberattigade med decimal: %s" % (val, d["kod"], d["rostberattigade"]))
            calc = round(100.0 * r / rb, 2)
            if abs(calc - float(d["valdeltagande"])) > 0.011:
                fel("%s %s valdeltagande %s, beräknat %.2f" % (val, d["kod"], d["valdeltagande"], calc))
            n_vdt += 1
        elif d["valdeltagande"] or d["namn"] != "Uppsamlingsdistrikt":
            fel("%s %s saknar röstberättigade" % (val, d["kod"]))
    print("%s distrikt.csv: %d rader, valdeltagande omräknat i %d distrikt" % (val, len(dist), n_vdt))


# ------------------------------------------------------------- kontroll 3
def kontroll_3(val, xd_excel, xd_xml, krets_av):
    print("\n--- (3) %s Excel mot XML ur källfilerna, och CSV mot CSV" % val)
    if set(xd_excel) != set(xd_xml):
        fel("%s kodmängd Excel != XML: %s" % (val, sorted(set(xd_excel) ^ set(xd_xml))))
    n = n_fel = 0
    for kod in sorted(xd_excel):
        e, x = xd_excel[kod], xd_xml[kod]
        if e["namn"] != x.get("NAMN"):
            fel("%s %s namn Excel '%s' XML '%s'" % (val, kod, e["namn"], x.get("NAMN")))
        xp = {norm(p): int(g.get("RÖSTER")) for p, g in xml_partier(x).items()}
        xa = {norm(p): g.get("PROCENT") for p, g in xml_partier(x).items()}
        ep = {}
        for p, (t, a, i, j) in e["partier"].items():
            if p in ICKE_PARTI:
                continue
            ep[norm(p)] = t
            n += 1
            if xp.get(norm(p), 0) != t:
                n_fel += 1
                fel("%s %s %s Excel %d XML %s" % (val, kod, p, t, xp.get(norm(p))))
            if a not in (None, "") and xa.get(norm(p)) is not None:
                if abs(float(a) - float(xa[norm(p)].replace(",", "."))) > 0.005:
                    n_fel += 1
                    fel("%s %s %s proc Excel %s XML %s" % (val, kod, p, a, xa[norm(p)]))
        rest = sum(v for p, v in xp.items() if p not in ep and p != "ÖVRIGA_FGVAL")
        ovr = e["partier"].get("OVR", e["partier"].get("ÖVR"))[0]
        if ovr != rest:
            n_fel += 1
            fel("%s %s ÖVR Excel %d XML-rest %d" % (val, kod, ovr, rest))
        xs = xml_summering(x)
        bl = e["partier"].get("BL", e["partier"].get("BLANK"))[0]
        og = e["partier"]["OG"][0]
        x_rostande = xs["rostande"] if xs["rostande"] is not None else xs["giltiga"] + xs["blanka"] + xs["og"]
        if (e["giltiga"], bl, og, e["rostande"]) != (xs["giltiga"], xs["blanka"], xs["og"], x_rostande):
            n_fel += 1
            fel("%s %s giltiga/blank/og/rostande Excel %s XML %s" % (
                val, kod, (e["giltiga"], bl, og, e["rostande"]), (xs["giltiga"], xs["blanka"], xs["og"], x_rostande)))
        if xs["rostb"] is not None:
            if to_int(e["rostb"]) != xs["rostb"]:
                n_fel += 1
                fel("%s %s rostb Excel %s XML %s" % (val, kod, e["rostb"], xs["rostb"]))
            if abs(float(e["vdt"]) - float(xs["vdt"])) > 0.005:
                n_fel += 1
                fel("%s %s vdt Excel %s XML %s" % (val, kod, e["vdt"], xs["vdt"]))
        elif to_int(e["rostb"]) != 0:
            fel("%s %s uppsamling har röstberättigade i Excel: %s" % (val, kod, e["rostb"]))
        if val == "rd" and e["kvk"] != krets_av[kod]:
            fel("%s %s kommunvalkrets Excel '%s' XML '%s'" % (val, kod, e["kvk"], krets_av[kod]))
    print("%s källa: %d partital i %d distrikt jämförda, %d avvikelser" % (val, n, len(xd_excel), n_fel))
    # CSV mot CSV
    a = {(r["kod"], r["parti"]): (int(r["roster"]), r["andel"], r["parti_kalla"]) for r in read_csv("roster_2014_%s_xls.csv" % val)}
    b = {(r["kod"], r["parti"]): (int(r["roster"]), r["andel"], r["parti_kalla"]) for r in read_csv("roster_2014_%s_xml.csv" % val)}
    n_fel = n_andel = 0
    for k, v in a.items():
        if k[1] == "ÖVR":
            continue
        if k not in b:
            if v[0] != 0:
                fel("%s CSV %s %s finns i xls (%d) men inte i xml" % (val, k[0], k[1], v[0]))
            continue
        if b[k][0] != v[0]:
            n_fel += 1
            fel("%s CSV %s %s xls %d xml %d" % (val, k[0], k[1], v[0], b[k][0]))
        if v[1] and v[1] != b[k][1]:
            n_andel += 1
            fel("%s CSV %s %s andel xls %s xml %s" % (val, k[0], k[1], v[1], b[k][1]))
        if v[2] != b[k][2]:
            n_fel += 1
            fel("%s CSV %s %s parti_kalla xls '%s' xml '%s'" % (val, k[0], k[1], v[2], b[k][2]))
    # partier i xml-csv som inte täcks av xls-csv ska summera till ÖVR
    per_kod_ovr = {k[0]: v[0] for k, v in a.items() if k[1] == "ÖVR"}
    xls_p = {k[1] for k in a if k[1] != "ÖVR"}
    rest = defaultdict(int)
    for k, v in b.items():
        if k[1] not in xls_p:
            rest[k[0]] += v[0]
    for kod, o in per_kod_ovr.items():
        if rest.get(kod, 0) != o:
            n_fel += 1
            fel("%s CSV %s ÖVR xls %d, xml-rest %d" % (val, kod, o, rest.get(kod, 0)))
    # valkrets i distrikt.csv mot XML-trädet
    for r in read_csv("distrikt_2014_%s.csv" % val):
        if r["valkrets"] != krets_av[r["kod"]]:
            fel("%s distrikt.csv %s valkrets '%s', XML '%s'" % (val, r["kod"], r["valkrets"], krets_av[r["kod"]]))
    print("%s CSV xls mot xml: %d nycklar, %d avvikelser i röster/parti_kalla, %d i andel; valkrets kontrollerad"
          % (val, len(a), n_fel, n_andel))


# ------------------------------------------------------------- kontroll 4
def kontroll_4(val):
    print("\n--- (4) %s koder och antal" % val)
    dbf = {r["VD"]: r["VD_NAMN"] for r in DBF(DBF2014, encoding="latin-1") if r["VD"].startswith("1480")}
    filer = {"roster_2014_%s_xls.csv" % val: "kod", "roster_2014_%s_xml.csv" % val: "kod",
             "distrikt_2014_%s.csv" % val: "kod", "fgval_2014_%s.csv" % val: "kod_2014"}
    for name, kol in filer.items():
        # råtexten, inte via csv-modulen, för att se att inget .0 eller citattecken smugit in
        with open(os.path.join(UT_DIR, name), encoding="utf-8") as f:
            hdr = f.readline().rstrip("\n").split(";")
            i = hdr.index(kol)
            bad = set()
            for line in f:
                k = line.rstrip("\n").split(";")[i]
                if not (len(k) == 8 and k.isdigit()):
                    bad.add(k)
        if bad:
            fel("%s koder ej 8 siffror: %s" % (name, sorted(bad)[:5]))
    dist = {r["kod"]: r["namn"] for r in read_csv("distrikt_2014_%s.csv" % val)}
    vanliga = {k: v for k, v in dist.items() if not k.startswith("148000")}
    upps = {k: v for k, v in dist.items() if k.startswith("148000")}
    if set(vanliga) != set(dbf):
        fel("%s distriktskoder != dbf 2014: %s" % (val, sorted(set(vanliga) ^ set(dbf))))
    namnfel = [(k, dbf[k], vanliga.get(k)) for k in dbf if dbf[k] != vanliga.get(k)]
    if namnfel:
        fel("%s namn skiljer mot dbf: %s" % (val, namnfel[:5]))
    if sorted(upps) != ["14800001", "14800002", "14800003", "14800004"]:
        fel("%s uppsamlingskoder: %s" % (val, sorted(upps)))
    print("%s: distrikt.csv %d = %d vanliga + %d uppsamling; dbf 2014 Göteborg %d; namn lika i dbf: %s"
          % (val, len(dist), len(vanliga), len(upps), len(dbf), not namnfel))


# ------------------------------------------------------------- kontroll 5
def kontroll_5(val, sheet, hdr, xd_excel, rng):
    print("\n--- (5) %s stickprov mot råcell (slumpfrö %d)" % (val, SLUMPFRO))
    koder = rng.sample(MAJORNA_2014, 3)
    roster = {(r["kod"], r["parti_kalla"]): r for r in read_csv("roster_2014_%s_xls.csv" % val)}
    roster_xml = {(r["kod"], r["parti"]): r for r in read_csv("roster_2014_%s_xml.csv" % val)}
    for kod in koder:
        e = xd_excel[kod]
        partier = [p for p in e["partier"] if p not in ICKE_PARTI and e["partier"][p][0] > 0]
        for p in rng.sample(partier, 3):
            t, a, i, j = e["partier"][p]
            raw = cell_raw(val, e["rad"], i)
            raw_a = cell_raw(val, e["rad"], j) if j is not None else None
            c = roster.get((kod, p))
            cx = roster_xml.get((kod, norm(p)))
            ok = c is not None and to_int(raw) == int(c["roster"]) and cx is not None and int(cx["roster"]) == to_int(raw)
            if raw_a not in (None, "") and c is not None and abs(float(raw_a) - float(c["andel"])) > 0.005:
                ok = False
            print("  %s %s (%s) %s: %s, blad '%s', rad %d, kolumn %s '%s' = %r, kolumn %s '%s' = %r; CSV xls %s/%s, xml %s -> %s"
                  % (val, kod, e["namn"], p, os.path.basename(VAL[val][1]), sheet, e["rad"], col_letter(i), hdr[i], raw,
                     col_letter(j) if j is not None else "-", hdr[j] if j is not None else "-", raw_a,
                     c["roster"] if c else None, c["andel"] if c else None, cx["roster"] if cx else None,
                     "OK" if ok else "FEL"))
            if not ok:
                fel("%s stickprov %s %s: råcell %r, CSV %s" % (val, kod, p, raw, c))


# ------------------------------------------------------------- kontroll 6
def kontroll_6(val, rng, xd_xml):
    print("\n--- (6) %s FGVAL mot 2010-agentens filer" % val)
    fg = defaultdict(dict)
    for r in read_csv("fgval_2014_%s.csv" % val):
        fg[r["kod_2014"]][r["parti"]] = (int(r["roster_2014"]), int(r["roster_fgval"]), r["ar_fg"])
    match = {r["kod_2014"]: r for r in read_csv("fgval_matchning_2010_2014.csv") if r["val"] == ("kf" if val == "kf" else "rd")}
    # 2010-agentens filer
    r2010 = defaultdict(dict)
    if os.path.exists(os.path.join(UT_DIR, "roster_2010_%s_xml.csv" % val)):
        for r in read_csv("roster_2010_%s_xml.csv" % val):
            r2010[r["kod"]][r["parti"]] = int(r["roster"])
    d2010 = {r["kod"]: r for r in read_csv("distrikt_2010_%s.csv" % val)}
    # roster_2014 i fgval-filen ska vara samma som i roster_xml-filen
    rx = {(r["kod"], r["parti"]): int(r["roster"]) for r in read_csv("roster_2014_%s_xml.csv" % val)}
    for kod, pp in fg.items():
        for p, (r14, rfg, ar) in pp.items():
            if p in ("GILTIGA", "BLANK", "OG", "SUMMA_RÖSTER", "RÖSTBERÄTTIGADE", "ÖVRIGA_FGVAL"):
                continue
            if rx.get((kod, p)) != r14:
                fel("%s fgval %s %s roster_2014 %d, roster_xml %s" % (val, kod, p, r14, rx.get((kod, p))))
    ar_fg = {v[2] for pp in fg.values() for v in pp.values()}
    print("%s fgval: %d distrikt med FGVAL, ar_fg %s" % (val, len(fg), sorted(ar_fg)))
    # fem Majornadistrikt med FGVAL
    kand = [k for k in MAJORNA_2014 if k in fg]
    fem = rng.sample(kand, 5)
    for kod in fem:
        m = match.get(kod)
        k10 = m["kod_2010"] if m else ""
        namn = xd_xml[kod].get("NAMN")
        if val == "rf":
            # FGVAL ska inte matcha 2010; kontrollera att det verkligen inte gör det
            pp = fg[kod]
            lika = [p for p in ("M", "S", "V", "MP") if r2010.get(k10, {}).get(p) == pp.get(p, (0, None))[1]]
            print("  %s %s %s: FGVAL GILTIGA %s, 2010 giltiga i %s: %s, RÖSTBERÄTTIGADE fg %s mot 2010 %s; partier lika 2010: %s"
                  % (val, kod, namn, pp["GILTIGA"][1], k10, d2010.get(k10, {}).get("giltiga"),
                     pp["RÖSTBERÄTTIGADE"][1], d2010.get(k10, {}).get("rostberattigade"), lika))
            continue
        pp = fg[kod]
        diffs = []
        n = 0
        for p, (r14, rfg, ar) in pp.items():
            if p in ("GILTIGA", "BLANK", "OG", "SUMMA_RÖSTER", "RÖSTBERÄTTIGADE", "ÖVRIGA_FGVAL", "HANDSKRIVNA"):
                continue
            n += 1
            v10 = r2010.get(k10, {}).get(p)
            if v10 != rfg:
                diffs.append((p, rfg, v10))
        d = d2010.get(k10, {})
        tot_ok = (to_int(d.get("giltiga", -1)) == pp["GILTIGA"][1] and to_int(d.get("blanka", -1)) == pp["BLANK"][1]
                  and to_int(d.get("ogiltiga_ovriga", -1)) == pp["OG"][1] and to_int(d.get("rostande", -1)) == pp["SUMMA_RÖSTER"][1]
                  and to_int(d.get("rostberattigade", -1)) == pp["RÖSTBERÄTTIGADE"][1])
        print("  %s %s %s -> 2010 %s %s: %d partirader, avvikelser %s; giltiga/blank/og/röstande/röstb %s"
              % (val, kod, namn, k10, d.get("namn"), n, diffs, "lika" if tot_ok else "OLIKA"))
        if diffs or not tot_ok:
            fel("%s fgval %s mot 2010-filerna: %s tot_ok=%s" % (val, kod, diffs, tot_ok))
    if val == "rf":
        return
    # hela filen
    n = n_fel = 0
    avv = []
    tomma_rostb = []
    utan_deltagande = []
    for kod, pp in fg.items():
        m = match.get(kod)
        if not m or m["matchning"] != "exakt":
            fel("%s fgval %s saknar exakt matchning" % (val, kod))
            continue
        k10 = m["kod_2010"]
        for p, (r14, rfg, ar) in pp.items():
            if p in ("GILTIGA", "BLANK", "OG", "SUMMA_RÖSTER", "RÖSTBERÄTTIGADE", "ÖVRIGA_FGVAL", "HANDSKRIVNA"):
                continue
            n += 1
            v10 = r2010.get(k10, {}).get(p)
            if v10 != rfg:
                n_fel += 1
                avv.append((kod, k10, p, rfg, v10))
        d = d2010.get(k10)
        if d is None:
            fel("%s 2010-distrikt %s saknas i distrikt_2010" % (val, k10))
            continue
        if (int(d["giltiga"]), int(d["blanka"]), int(d["ogiltiga_ovriga"])) != (
                pp["GILTIGA"][1], pp["BLANK"][1], pp["OG"][1]):
            fel("%s fgval %s giltiga/blank/og skiljer från distrikt_2010 %s" % (val, kod, k10))
        if "SUMMA_RÖSTER" not in pp:
            utan_deltagande.append(kod)
            continue
        if int(d["rostande"]) != pp["SUMMA_RÖSTER"][1]:
            fel("%s fgval %s röstande fgval %d, distrikt_2010 %s %s" % (val, kod, pp["SUMMA_RÖSTER"][1], k10, d["rostande"]))
        if d["rostberattigade"] == "":
            tomma_rostb.append((kod, k10, pp["RÖSTBERÄTTIGADE"][1]))
        elif int(d["rostberattigade"]) != pp["RÖSTBERÄTTIGADE"][1]:
            fel("%s fgval %s röstberättigade fgval %d, distrikt_2010 %s %s" % (val, kod, pp["RÖSTBERÄTTIGADE"][1], k10, d["rostberattigade"]))
    print("%s hela fgval-filen: %d partirader mot roster_2010_%s_xml.csv, %d avvikelser" % (val, n, val, n_fel))
    for a in avv:
        print("   avvikelse %s -> %s parti %s fgval %d, 2010-fil %s" % a)
    if utan_deltagande:
        print("   distrikt utan SUMMA_RÖSTER/RÖSTBERÄTTIGADE i fgval-filen: %s" % utan_deltagande)
    if tomma_rostb:
        print("   distrikt_2010 har tomt röstberättigade för %s (FGVAL RÖSTBERÄTTIGADE i 2014-filen: %s)" % (
            [t[1] for t in tomma_rostb], [t[2] for t in tomma_rostb]))


# ------------------------------------------------------------- extra
def kontroll_fgval_matchning(val, xd_xml):
    """Egen matchning av FGVAL mot 2010 års XML med alla partier, inte bara åtta."""
    print("\n--- extra: %s egen FGVAL-matchning mot 2010 XML" % val)
    root10 = ET.parse(os.path.join(XML2010, "slutresultat_1480%s.xml" % VAL[val][0])).getroot()
    v10 = {}
    for d in root10.iter("VALDISTRIKT"):
        v10[d.get("KOD")] = {p: int(g.get("RÖSTER")) for p, g in xml_partier(d).items()}
        v10[d.get("KOD")]["_giltiga"] = int(d.get("RÖSTER"))
    for d in root10.iter("ONSDAGSDISTRIKT"):
        kod = "1480" + d.get("KOD").split("-")[-1].zfill(4)
        v10[kod] = {p: int(g.get("RÖSTER")) for p, g in xml_partier(d).items()}
        v10[kod]["_giltiga"] = int(d.get("RÖSTER"))
    match = {r["kod_2014"]: r for r in read_csv("fgval_matchning_2010_2014.csv") if r["val"] == val}
    n = n_ok = n_flera = 0
    for kod, el in xd_xml.items():
        if el.get("RÖSTER_FGVAL") is None:
            if match.get(kod, {}).get("matchning") != "ingen FGVAL":
                fel("%s matchning %s borde vara 'ingen FGVAL'" % (val, kod))
            continue
        n += 1
        fg = {p: int(g.get("RÖSTER_FGVAL")) for p, g in xml_partier(el).items() if g.get("RÖSTER_FGVAL") is not None}
        g_fg = int(el.get("RÖSTER_FGVAL"))
        # kandidater: 2010-distrikt med samma giltiga och samma röster för alla partier med FGVAL > 0
        kand = []
        for k10, v in v10.items():
            if v["_giltiga"] != g_fg:
                continue
            ok = all(v.get(p, 0) == r for p, r in fg.items() if p not in ("ÖVRIGA_FGVAL", "HANDSKRIVNA"))
            if ok:
                kand.append(k10)
        m = match.get(kod, {})
        if len(kand) == 1:
            n_ok += 1
            if m.get("kod_2010") != kand[0]:
                fel("%s matchning %s: fil säger %s, egen matchning %s" % (val, kod, m.get("kod_2010"), kand[0]))
        elif len(kand) > 1:
            n_flera += 1
            print("   %s flera kandidater: %s (fil: %s)" % (kod, kand, m.get("kod_2010")))
        else:
            print("   %s ingen kandidat (fil: %s %s)" % (kod, m.get("kod_2010"), m.get("matchning")))
    print("%s: %d distrikt med FGVAL, %d entydig egen matchning, %d flera; 2010-distrikt i XML %d" % (val, n, n_ok, n_flera, len(v10)))


def kontroll_extra():
    print("\n--- extra: övrigt")
    # 2010 dbf: SDN Majorna
    dbf10 = {r["LKFV"]: r["VDNAMN"] for r in DBF(DBF2010, encoding="latin-1") if r["LKFV"].startswith("1480")}
    maj10 = sorted(k for k in dbf10 if k.startswith("148009"))
    print("2010 dbf Göteborg %d distrikt, 148009xx: %d: %s" % (len(dbf10), len(maj10), maj10))
    # namn_2010 i matchningsfilen mot dbf 2010
    n_fel = 0
    for r in read_csv("fgval_matchning_2010_2014.csv"):
        if r["kod_2010"] and not r["kod_2010"].startswith("148000"):
            if dbf10.get(r["kod_2010"]) != r["namn_2010"]:
                n_fel += 1
                fel("matchning namn_2010 %s '%s' dbf '%s'" % (r["kod_2010"], r["namn_2010"], dbf10.get(r["kod_2010"])))
    print("fgval_matchning namn_2010 mot dbf 2010: %d avvikelser" % n_fel)
    # röstberättigade i distrikt-filerna mot rostberattigade_2014_*_bred.csv (alla distrikt)
    for val in VAL:
        name = "rostberattigade_2014_%s_bred.csv" % val
        if not os.path.exists(os.path.join(UT_DIR, name)):
            print("saknas:", name)
            continue
        bred = {r["valdistrikt_id"]: int(r["totalt"]) for r in read_csv(name)}
        dist = {r["kod"]: r for r in read_csv("distrikt_2014_%s.csv" % val)}
        n = n_f = 0
        for kod, tot in bred.items():
            if kod not in dist:
                continue
            n += 1
            if dist[kod]["rostberattigade"] != str(tot):
                n_f += 1
                if n_f <= 3:
                    fel("%s %s rostberattigade distrikt %s, bred %d" % (val, kod, dist[kod]["rostberattigade"], tot))
        print("%s röstberättigade distrikt.csv mot %s: %d distrikt, %d avvikelser; summa bred %d" % (val, name, n, n_f, sum(bred.values())))
    # mandat
    m = read_csv("mandat_2014_riksdag.csv")
    riket = sum(int(r["mandat"]) for r in m if r["niva"] == "riket")
    gbg = [(r["parti"], r["mandat"], r["varav_utjamning"]) for r in m if r["niva_kod"] == "1416"]
    print("mandat riksdag: riket %d, Göteborgs kommun (1416) %s = %d, mandat_totalt %s" % (
        riket, gbg, sum(int(g[1]) for g in gbg), {r["mandat_totalt"] for r in m if r["niva_kod"] == "1416"}))
    kretsar = defaultdict(int)
    for r in m:
        if r["niva"] != "riket":
            kretsar[r["niva"]] += int(r["mandat"])
    print("mandat riksdag: %d valkretsar, summa %d" % (len(kretsar), sum(kretsar.values())))
    mk = read_csv("mandat_2014_kf.csv")
    print("mandat kf: goteborg %d, kretsar %s" % (sum(int(r["mandat"]) for r in mk if r["niva"] == "goteborg"),
          {n: sum(int(r["mandat"]) for r in mk if r["niva"] == n) for n in {r["niva"] for r in mk} - {"goteborg"}}))
    mr = read_csv("mandat_2014_rf.csv")
    print("mandat rf: vgregion %d, kretsar %s" % (sum(int(r["mandat"]) for r in mr if r["niva"] == "vgregion"),
          {n: sum(int(r["mandat"]) for r in mr if r["niva"] == n) for n in {r["niva"] for r in mr} - {"vgregion"}}))
    # mandat riksdag mot aggregat riket
    agg = {r["parti"]: r for r in read_csv("aggregat_2014_rd.csv") if r["niva"] == "riket"}
    for r in m:
        if r["niva"] == "riket" and agg[r["parti"]]["roster"] != r["roster"]:
            fel("mandat riket %s roster %s, aggregat %s" % (r["parti"], r["roster"], agg[r["parti"]]["roster"]))
    # parti-normalisering: samma uppsättning parti_kalla -> parti i alla filer
    par = defaultdict(set)
    for name in os.listdir(UT_DIR):
        if "2014" in name and name.endswith(".csv") and (name.startswith("roster_") or name.startswith("aggregat_") or name.startswith("mandat_")):
            for r in read_csv(name):
                par[r["parti_kalla"]].add(r["parti"])
    dubbel = {k: v for k, v in par.items() if len(v) > 1}
    if dubbel:
        fel("parti_kalla med flera normaliseringar: %s" % dubbel)
    print("parti_kalla -> parti (avvikande från versaler): %s" % {k: v for k, v in par.items() if list(v)[0] != k.upper()})


def kontroll_rf_fgval(xd_rf):
    """Belägg för att landstingets FGVAL är omvalet 2011 och inte 2010."""
    print("\n--- extra: landstingets FGVAL")
    root10 = ET.parse(os.path.join(XML2010, "slutresultat_1480L.xml")).getroot()
    k10 = root10.find("KOMMUN")
    s10 = xml_summering(k10)
    root14 = ET.parse(os.path.join(XML2014, "slutresultat_1480L.xml")).getroot()
    k14 = root14.find("KOMMUN")
    s14 = xml_summering(k14)
    print("2010 L-XML KOMMUN: giltiga %s, röstande %s, röstb %s, VALDAG %s" % (s10["giltiga"], s10["rostande"], s10["rostb"], root10.get("VALDAG")))
    print("2014 L-XML KOMMUN FGVAL: giltiga %s, röstande %s, röstb %s, VALDAG_FGVAL %s" % (s14["giltiga_fg"], s14["rostande_fg"], s14["rostb_fg"], root14.get("VALDAG_FGVAL")))
    # samma för R och K som referens
    for b in "RK":
        r10 = ET.parse(os.path.join(XML2010, "slutresultat_1480%s.xml" % b)).getroot().find("KOMMUN")
        r14 = ET.parse(os.path.join(XML2014, "slutresultat_1480%s.xml" % b)).getroot().find("KOMMUN")
        a, c = xml_summering(r10), xml_summering(r14)
        print("%s: 2010 giltiga %s / 2014 FGVAL giltiga %s, röstb 2010 %s / FGVAL %s" % (b, a["giltiga"], c["giltiga_fg"], a["rostb"], c["rostb_fg"]))
    v10 = {d.get("KOD"): int(d.get("RÖSTER")) for d in root10.iter("VALDISTRIKT")}
    n = n_lika = 0
    for kod, el in xd_rf.items():
        if el.get("RÖSTER_FGVAL") is None:
            continue
        n += 1
        if int(el.get("RÖSTER_FGVAL")) in v10.values():
            n_lika += 1
    print("rf: %d distrikt med FGVAL, %d vars FGVAL-giltiga sammanfaller med något 2010-distrikts giltiga" % (n, n_lika))


def main():
    rng = random.Random(SLUMPFRO)
    xd_rf = None
    for val in VAL:
        sheet, hdr, xd_excel = excel_distrikt(val)
        root, xd_xml, krets_av = xml_distrikt(XML2014, val)
        if val == "rf":
            xd_rf = xd_xml
        print("\n=================== %s: Excel %d rader, blad '%s', %d kolumner; XML %d distrikt"
              % (val, len(xd_excel), sheet, len(hdr), len(xd_xml)))
        kontroll_1(val, xd_excel, root, xd_xml, krets_av)
        kontroll_2(val)
        kontroll_3(val, xd_excel, xd_xml, krets_av)
        kontroll_4(val)
        kontroll_5(val, sheet, hdr, xd_excel, rng)
        kontroll_6(val, rng, xd_xml)
        if val != "rf":
            kontroll_fgval_matchning(val, xd_xml)
    kontroll_rf_fgval(xd_rf)
    kontroll_extra()
    print("\n=================== antal FEL/AVVIKELSER: %d" % len(FEL))
    for f in FEL:
        print(" -", f)


if __name__ == "__main__":
    main()
