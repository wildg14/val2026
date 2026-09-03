# -*- coding: utf-8 -*-
"""
granskning_2018_kontroll.py - oberoende granskning av val2018-agentens filer.

Laser kallfilerna sjalv (openpyxl) och jamfor med data/historik/*_2018_*.csv.
Andra kallor som anvands som motkontroll:
  - slutligt-valresultat-riksdagen-jamforande-statistik-2018-2022.xlsx (Roster 2018 per riket,
    riksdagsvalkrets och kommun)
  - Mandatfordelning-jamforelser-mellan-2018-och-2022.xlsx (mandat 2018)
  - reparerade 2018_rostberattigade_R/L/K.xlsx (rostberattigade per distrikt)
  - historik.val.se distriktssidor (HTML) for tre slumpade Majornadistrikt, R, L och K
  - shapefilen alla_valdistrikt.dbf (koder och namn)

Kors med:
  <venv>/bin/python scripts/historik/granskning_2018_kontroll.py
Skriver bara till stdout (och cachar HTML-sidor i scratchpad). Andrar inga CSV-filer.
"""
import csv
import collections
import html
import os
import random
import re
import sys
import urllib.request

import openpyxl
from openpyxl.utils import get_column_letter
from dbfread import DBF

KALLMAPP = "/Users/daniel/code/Temp/Historiska dokument"
FIL = {"rd": os.path.join(KALLMAPP, "2018_R_per_valdistrikt.xlsx"),
       "rf": os.path.join(KALLMAPP, "2018_L_per_valdistrikt.xlsx"),
       "kf": os.path.join(KALLMAPP, "2018_K_per_valdistrikt.xlsx")}
FLIK = {"rd": ("R antal", "R procent"), "rf": ("L antal", "L procent"), "kf": ("K antal", "K procent")}
FIL_MANDAT = os.path.join(KALLMAPP, "2018_mandat.xlsx")
SCRATCH = "/private/tmp/claude-501/-Users-daniel-code-Temp/f347baf2-1af3-43b3-b36d-e80b868ede0e/scratchpad"
SHP_DBF = SCRATCH + "/unz/2018_valgeografi_valdistrikt/alla_valdistrikt.dbf"
FIL_ROSTBER = {"rd": SCRATCH + "/repaired/2018_rostberattigade_R.xlsx",
               "rf": SCRATCH + "/repaired/2018_rostberattigade_L.xlsx",
               "kf": SCRATCH + "/repaired/2018_rostberattigade_K.xlsx"}
FIL_JAMF = "/Users/daniel/code/Temp/slutligt-valresultat-riksdagen-jamforande-statistik-2018-2022.xlsx"
FIL_MANDAT_JAMF = "/Users/daniel/code/Temp/Mandatfordelning-jamforelser-mellan-2018-och-2022.xlsx"
UTMAPP = "/Users/daniel/code/Temp/data/historik"
HTML_CACHE = SCRATCH + "/dl2018_granskning"
HTML_URL = "https://historik.val.se/val/val2018/slutresultat/{V}/valdistrikt/14/80/{vd}/index.html"

MAJORNA = ["14801011", "14801012", "14801013", "14801014", "14801015", "14801016", "14801017",
           "14801021", "14801022", "14801031", "14801032", "14801033", "14801034", "14801035",
           "14801036", "14801037", "14801038", "14801041", "14801042", "14801043", "14801044", "14801045"]

fel = []
info = []


def rapport(kat, text):
    (fel if kat == "FEL" else info).append(text)
    print(f"[{kat}] {text}")


def las_csv(namn):
    with open(os.path.join(UTMAPP, namn), encoding="utf-8", newline="") as f:
        rader = list(csv.DictReader(f, delimiter=";"))
    return rader


def las_flik(fil, flik):
    wb = openpyxl.load_workbook(fil, read_only=True)
    ws = wb[flik]
    rader = list(ws.iter_rows(values_only=True))
    return [str(h) if h is not None else "" for h in rader[0]], rader[1:]


def i0(v):
    return int(v) if v not in (None, "") else 0


def normalisera(p):
    return {"FP": "L", "DEM": "D", "KP": "K"}.get(p, p.upper())


# ------------------------------------------------------------------ 1. kallfiler sjalv
kalla = {}   # val -> dict kod -> {parti: antal, ...} + fält
kalla_rad = {}  # val -> kod -> radnummer i excel (1-baserat inkl. rubrik)
partikol = {}
for val in ("rd", "rf", "kf"):
    hdr, rader = las_flik(FIL[val], FLIK[val][0])
    hdr_p, rader_p = las_flik(FIL[val], FLIK[val][1])
    ix = {h: i for i, h in enumerate(hdr)}
    pk = list(range(ix["VALDISTRIKTSNAMN"] + 1, ix["OGEJ"]))
    partikol[val] = {hdr[i]: i for i in pk}
    d, dr = {}, {}
    for n, r in enumerate(rader, start=2):
        if r[ix["LÄNSKOD"]] != 14 or r[ix["KOMMUNKOD"]] != 80:
            continue
        kod = f"{int(r[ix['LÄNSKOD']]):02d}{int(r[ix['KOMMUNKOD']]):02d}{int(r[ix['VALDISTRIKTSKOD']]):04d}"
        if kod in d:
            rapport("FEL", f"{val}: dubblettkod {kod} i kallan")
        d[kod] = {"namn": r[ix["VALDISTRIKTSNAMN"]], "valkrets": r[ix["VALKRETSNAMN"]],
                  "OGEJ": i0(r[ix["OGEJ"]]), "BLANK": i0(r[ix["BLANK"]]), "OG": i0(r[ix["OG"]]),
                  "giltiga": i0(r[ix["RÖSTER GILTIGA"]]), "rostande": i0(r[ix["RÖSTANDE"]]),
                  "rostber": i0(r[ix["RÖSTBERÄTTIGADE"]]), "valdelt": r[ix["VALDELTAGANDE"]],
                  "partier": {hdr[i]: i0(r[i]) for i in pk}}
        dr[kod] = n
    # procentflik
    ixp = {h: i for i, h in enumerate(hdr_p)}
    for r in rader_p:
        if r[ixp["LÄNSKOD"]] != 14 or r[ixp["KOMMUNKOD"]] != 80:
            continue
        kod = f"1480{int(r[ixp['VALDISTRIKTSKOD']]):04d}"
        d[kod]["procent"] = {hdr_p[i]: r[i] for i in pk}
    kalla[val] = d
    kalla_rad[val] = dr
    print(f"kalla {val}: {len(d)} Goteborgsrader, {len(pk)} partikolumner")

# ------------------------------------------------------------------ 2. lasta CSV-filer
roster = {v: las_csv(f"roster_2018_{v}.csv") for v in ("rd", "rf", "kf")}
distrikt = {v: las_csv(f"distrikt_2018_{v}.csv") for v in ("rd", "rf", "kf")}
ogilt = {v: las_csv(f"ogiltiga_2018_{v}.csv") for v in ("rd", "rf", "kf")}
aggregat = {v: las_csv(f"aggregat_2018_{v}.csv") for v in ("rd", "rf", "kf")}
mandat = las_csv("mandat_2018.csv")
partier = las_csv("partier_2018.csv")

# format: koder 8 siffror text
for namn, filer in [("roster", roster), ("distrikt", distrikt), ("ogiltiga", ogilt)]:
    for v, rader in filer.items():
        dåliga = [r["kod"] for r in rader if not re.fullmatch(r"\d{8}", r["kod"])]
        if dåliga:
            rapport("FEL", f"{namn}_{v}: koder som inte ar 8 siffror: {dåliga[:5]}")
        for r in rader:
            if r["ar"] != "2018" or r["val"] != v:
                rapport("FEL", f"{namn}_{v}: fel ar/val i rad {r}")
                break
        for k in ("roster", "giltiga", "blanka", "ogiltiga", "rostande", "rostberattigade",
                  "ogiltiga_ej_anmalda", "ogiltiga_ovriga"):
            if k in rader[0]:
                dåliga = [r[k] for r in rader if not re.fullmatch(r"\d+", r[k])]
                if dåliga:
                    rapport("FEL", f"{namn}_{v}: kolumn {k} ej heltal: {dåliga[:5]}")
        for k in ("andel", "valdeltagande"):
            if k in rader[0]:
                dåliga = [r[k] for r in rader if not re.fullmatch(r"\d+\.\d\d", r[k])]
                if dåliga:
                    rapport("FEL", f"{namn}_{v}: kolumn {k} ej tal med tva decimaler och punkt: {dåliga[:5]}")
# radbrytning och BOM
for namn in ["roster_2018_rd.csv", "distrikt_2018_kf.csv", "aggregat_2018_rf.csv", "mandat_2018.csv"]:
    b = open(os.path.join(UTMAPP, namn), "rb").read()
    if b.startswith(b"\xef\xbb\xbf"):
        rapport("FEL", f"{namn}: BOM i borjan")
    if b"\r\n" in b:
        rapport("INFO", f"{namn}: CRLF radbrytning")
    if "\t" in b.decode("utf-8") or b.count(b",") and False:
        pass

# ------------------------------------------------------------------ 3. distriktsantal och namn mot shapefil
shp = {r["VD"]: r["VD_NAMN"] for r in DBF(SHP_DBF, encoding="utf-8") if str(r["VD"]).startswith("1480")}
print(f"shapefil: {len(shp)} Goteborgsdistrikt")
for v in ("rd", "rf", "kf"):
    koder = {r["kod"] for r in distrikt[v]}
    if len(distrikt[v]) != 351 or len(koder) != 351:
        rapport("FEL", f"distrikt_{v}: {len(distrikt[v])} rader, {len(koder)} unika koder, vantat 351")
    extra = koder - set(shp)
    saknas = set(shp) - koder
    if extra != {"14800000"} or saknas:
        rapport("FEL", f"distrikt_{v}: koder utover shapefil {sorted(extra)}, saknas {sorted(saknas)}")
    namn_diff = [(r["kod"], r["namn"], shp[r["kod"]]) for r in distrikt[v] if r["kod"] in shp and r["namn"] != shp[r["kod"]]]
    if namn_diff:
        rapport("FEL", f"distrikt_{v}: namnskillnad mot shapefil {namn_diff[:5]}")
    if set(kalla[v]) != koder:
        rapport("FEL", f"distrikt_{v}: kodmangd skiljer fran kallan")
    rkoder = {r["kod"] for r in roster[v]}
    if rkoder != koder:
        rapport("FEL", f"roster_{v}: kodmangd skiljer fran distriktsfilen")
    # namn i roster och ogiltiga lika med distrikt
    dn = {r["kod"]: r["namn"] for r in distrikt[v]}
    for namn, rader in (("roster", roster[v]), ("ogiltiga", ogilt[v])):
        nd = [r["kod"] for r in rader if r["namn"] != dn[r["kod"]]]
        if nd:
            rapport("FEL", f"{namn}_{v}: namn skiljer fran distriktsfil for {nd[:5]}")
ml = sorted(k for k in shp if shp[k].startswith("Majorna-Linné"))
print(f"Majorna-Linne i shapefil: {len(ml)} distrikt ({ml[0]}-{ml[-1]})")
ml_fil = sorted(r["kod"] for r in distrikt["rd"] if r["namn"].startswith("Majorna-Linné"))
if ml != ml_fil:
    rapport("FEL", f"Majorna-Linne: shapefil {len(ml)} vs distriktsfil {len(ml_fil)}")

# ------------------------------------------------------------------ 4. per distrikt: roster mot kalla, summor
for v in ("rd", "rf", "kf"):
    # roster per distrikt och parti mot kallan
    rs = collections.defaultdict(dict)
    for r in roster[v]:
        rs[r["kod"]][r["parti_kalla"]] = r
    n_jmf = 0
    for kod, k in kalla[v].items():
        for p, antal in k["partier"].items():
            r = rs[kod].get(p)
            if r is None:
                if antal != 0:
                    rapport("FEL", f"roster_{v} {kod} {p}: saknas i CSV men kallan har {antal}")
                continue
            n_jmf += 1
            if int(r["roster"]) != antal:
                rapport("FEL", f"roster_{v} {kod} {p}: CSV {r['roster']} != kalla {antal}")
            # andel mot procentflik och mot omraknad
            pa = k["procent"].get(p)
            if pa is None:
                pa = 0.0 if antal == 0 else None
            if pa is not None and abs(float(r["andel"]) - float(pa)) > 0.005:
                rapport("FEL", f"roster_{v} {kod} {p}: andel CSV {r['andel']} != procentflik {pa}")
            if k["giltiga"]:
                om = 100 * antal / k["giltiga"]
                if abs(float(r["andel"]) - om) > 0.0101:
                    rapport("FEL", f"roster_{v} {kod} {p}: andel {r['andel']} != omraknad {om:.3f}")
            if r["parti"] != normalisera(p):
                rapport("FEL", f"roster_{v} {kod} {p}: parti {r['parti']} != normaliserad {normalisera(p)}")
        # partier i CSV som inte finns i kallan
        for p in rs[kod]:
            if p not in k["partier"]:
                rapport("FEL", f"roster_{v} {kod}: parti {p} finns inte i kallan")
    print(f"{v}: {n_jmf} (distrikt, parti)-par jamforda mot kallan")
    # summa partier = giltiga, giltiga + ogiltiga = rostande, valdeltagande
    d_by = {r["kod"]: r for r in distrikt[v]}
    o_by = {r["kod"]: r for r in ogilt[v]}
    for kod, k in kalla[v].items():
        d = d_by[kod]
        s = sum(int(r["roster"]) for r in rs[kod].values())
        if s != int(d["giltiga"]):
            rapport("FEL", f"{v} {kod}: partisumma {s} != giltiga {d['giltiga']}")
        if int(d["giltiga"]) + int(d["ogiltiga"]) != int(d["rostande"]):
            rapport("FEL", f"{v} {kod}: giltiga+ogiltiga != rostande")
        if int(d["blanka"]) + int(d["ogiltiga_ovriga"]) != int(d["ogiltiga"]):
            rapport("FEL", f"{v} {kod}: blanka+ogiltiga_ovriga != ogiltiga")
        rb = int(d["rostberattigade"])
        if rb:
            vd = round(100 * int(d["rostande"]) / rb, 2)
            if abs(vd - float(d["valdeltagande"])) > 0.0101:
                rapport("FEL", f"{v} {kod}: valdeltagande {d['valdeltagande']} != {vd}")
        elif d["valdeltagande"] != "0.00":
            rapport("FEL", f"{v} {kod}: rostberattigade 0 men valdeltagande {d['valdeltagande']}")
        # mot kallan
        for k_csv, k_src in (("giltiga", "giltiga"), ("rostande", "rostande"), ("rostberattigade", "rostber")):
            if int(d[k_csv]) != k[k_src]:
                rapport("FEL", f"{v} {kod}: {k_csv} CSV {d[k_csv]} != kalla {k[k_src]}")
        if int(d["blanka"]) != k["BLANK"] or int(d["ogiltiga_ovriga"]) != k["OGEJ"] + k["OG"]:
            rapport("FEL", f"{v} {kod}: blanka/ogiltiga_ovriga skiljer fran kallan")
        if abs(float(d["valdeltagande"]) - float(k["valdelt"] or 0)) > 0.005:
            rapport("FEL", f"{v} {kod}: valdeltagande CSV {d['valdeltagande']} != kalla {k['valdelt']}")
        if d["namn"] != k["namn"]:
            rapport("FEL", f"{v} {kod}: namn CSV {d['namn']} != kalla {k['namn']}")
        if (d["valkrets"] or "") != (k["valkrets"] or ""):
            rapport("FEL", f"{v} {kod}: valkrets CSV {d['valkrets']!r} != kalla {k['valkrets']!r}")
        o = o_by[kod]
        if (int(o["ogiltiga_ej_anmalda"]), int(o["blanka"]), int(o["ogiltiga_ovriga"])) != (k["OGEJ"], k["BLANK"], k["OG"]):
            rapport("FEL", f"ogiltiga_{v} {kod}: skiljer fran kallan")

# ------------------------------------------------------------------ 5. Goteborgstotaler mot aggregat
tot_kalla = {}
for v in ("rd", "rf", "kf"):
    t = collections.Counter()
    g = r_ = rb = 0
    for k in kalla[v].values():
        for p, a in k["partier"].items():
            t[p] += a
        g += k["giltiga"]; r_ += k["rostande"]; rb += k["rostber"]
    tot_kalla[v] = (t, g, r_, rb)
    agg = {r["parti_kalla"]: r for r in aggregat[v] if r["niva"] == "goteborg"}
    for p, a in t.items():
        if a == 0:
            if p in agg:
                rapport("FEL", f"aggregat_{v} goteborg: {p} har 0 roster men finns i aggregatet")
            continue
        if p not in agg:
            rapport("FEL", f"aggregat_{v} goteborg: {p} saknas ({a} roster)")
            continue
        r = agg[p]
        if int(r["roster"]) != a:
            rapport("FEL", f"aggregat_{v} goteborg {p}: {r['roster']} != summa ur kalla {a}")
        if (int(r["giltiga"]), int(r["rostande"]), int(r["rostberattigade"])) != (g, r_, rb):
            rapport("FEL", f"aggregat_{v} goteborg {p}: giltiga/rostande/rostber {r['giltiga']}/{r['rostande']}/{r['rostberattigade']} != {g}/{r_}/{rb}")
        if abs(float(r["andel"]) - 100 * a / g) > 0.0101:
            rapport("FEL", f"aggregat_{v} goteborg {p}: andel {r['andel']} != {100*a/g:.3f}")
    # summa av roster-CSV = aggregat
    tcsv = collections.Counter()
    for r in roster[v]:
        tcsv[r["parti_kalla"]] += int(r["roster"])
    for p in agg:
        if tcsv[p] != int(agg[p]["roster"]):
            rapport("FEL", f"aggregat_{v} goteborg {p}: roster-CSV summa {tcsv[p]} != aggregat {agg[p]['roster']}")
    # summa av distrikt-CSV
    sd = (sum(int(r["giltiga"]) for r in distrikt[v]), sum(int(r["rostande"]) for r in distrikt[v]),
          sum(int(r["rostberattigade"]) for r in distrikt[v]))
    if sd != (g, r_, rb):
        rapport("FEL", f"distrikt_{v}: summor {sd} != kalla {(g, r_, rb)}")
    print(f"{v} Goteborg ur kallan: giltiga {g}, rostande {r_}, rostberattigade {rb}, "
          f"valdeltagande {100*r_/rb:.2f}; topp: {t.most_common(5)}")
    # riket och vgregion: rakna om ur hela filen
    hdr, rader = las_flik(FIL[v], FLIK[v][0])
    ix = {h: i for i, h in enumerate(hdr)}
    for niva, urval in (("riket", lambda r: True), ("vgregion", lambda r: r[ix["LÄNSKOD"]] == 14)):
        tt = collections.Counter(); gg = rr = bb = 0
        for r in rader:
            if not urval(r):
                continue
            for p, i in partikol[v].items():
                tt[p] += i0(r[i])
            gg += i0(r[ix["RÖSTER GILTIGA"]]); rr += i0(r[ix["RÖSTANDE"]]); bb += i0(r[ix["RÖSTBERÄTTIGADE"]])
        agg = {r["parti_kalla"]: r for r in aggregat[v] if r["niva"] == niva}
        for p, a in tt.items():
            if a and (p not in agg or int(agg[p]["roster"]) != a):
                rapport("FEL", f"aggregat_{v} {niva} {p}: {agg.get(p, {}).get('roster')} != {a}")
            if a and (int(agg[p]["giltiga"]), int(agg[p]["rostande"]), int(agg[p]["rostberattigade"])) != (gg, rr, bb):
                rapport("FEL", f"aggregat_{v} {niva} {p}: totaler skiljer")
        if sum(tt.values()) != gg:
            rapport("FEL", f"{v} {niva}: partisumma {sum(tt.values())} != giltiga {gg}")
        print(f"{v} {niva} ur kallan: giltiga {gg}, rostande {rr}, rostberattigade {bb}")

# ------------------------------------------------------------------ 6. andra kallor
# 6a. jamforande statistik 2018-2022 (riksdag): Roster 2018 per riket, valkrets Goteborgs kommun, kommun Goteborg
NAMN2KOD = {"Moderaterna": "M", "Centerpartiet": "C", "Liberalerna (tidigare Folkpartiet)": "L",
            "Kristdemokraterna": "KD", "Arbetarepartiet-Socialdemokraterna": "S", "Vänsterpartiet": "V",
            "Miljöpartiet de gröna": "MP", "Sverigedemokraterna": "SD", "Feministiskt initiativ": "FI",
            "Alternativ för Sverige": "AFS", "Medborgerlig Samling": "MED", "Piratpartiet": "PP",
            "Kristna Värdepartiet": "KRVP", "Djurens parti": "DJUP", "Enhet": "ENH",
            "Landsbygdspartiet Oberoende": "LPO", "Sveriges Kommunistiska Parti": "SKP",
            "Direktdemokraterna": "DD", "Klassiskt liberala partiet": "KLP", "Trygghetspartiet": "TRP",
            "Nordiska motståndsrörelsen": "NMR", "Europeiska Arbetarpartiet-EAP": "EAP",
            "Demokraterna": "DEM", "Kommunistiska Partiet": "K", "Vägvalet": "VägV"}
wb = openpyxl.load_workbook(FIL_JAMF, read_only=True)
agg_rd = {(r["niva"], r["parti_kalla"]): r for r in aggregat["rd"]}
n = 0
for r in wb["Riket"].iter_rows(values_only=True, min_row=2):
    if r[1] in NAMN2KOD and r[6] is not None:
        p = NAMN2KOD[r[1]]
        a = agg_rd.get(("riket", p))
        n += 1
        if a is None or int(a["roster"]) != int(r[6]):
            rapport("FEL", f"jamforande Riket 2018 {p}: {r[6]} != aggregat {a and a['roster']}")
        elif abs(float(a["andel"]) - 100 * float(r[7])) > 0.0101:
            rapport("FEL", f"jamforande Riket 2018 {p}: andel {100*float(r[7]):.2f} != aggregat {a['andel']}")
    elif r[1] and r[1] not in NAMN2KOD and r[6] is not None:
        rapport("INFO", f"jamforande Riket: rad utan partikod: {r[1]} {r[6]}")
print(f"jamforande statistik Riket: {n} partier jamforda")
n = 0
for r in wb["Valkrets"].iter_rows(values_only=True, min_row=2):
    if r[1] == "Göteborgs kommun" and r[2] in NAMN2KOD and r[7] is not None:
        p = NAMN2KOD[r[2]]; n += 1
        a = agg_rd.get(("goteborg", p))
        if a is None or int(a["roster"]) != int(r[7]):
            rapport("FEL", f"jamforande Valkrets Goteborg 2018 {p}: {r[7]} != aggregat {a and a['roster']}")
        elif abs(float(a["andel"]) - 100 * float(r[8])) > 0.0101:
            rapport("FEL", f"jamforande Valkrets Goteborg 2018 {p}: andel {100*float(r[8]):.2f} != {a['andel']}")
    elif r[1] == "Göteborgs kommun" and r[7] is not None:
        rapport("INFO", f"jamforande Valkrets Goteborg: rad utan partikod {r[2]} {r[7]}")
print(f"jamforande statistik Valkrets Goteborgs kommun: {n} partier jamforda")
n = 0
for r in wb["Kommun"].iter_rows(values_only=True, min_row=2):
    if r[1] == "Göteborg" and r[2] in NAMN2KOD and r[7] is not None:
        p = NAMN2KOD[r[2]]; n += 1
        a = agg_rd.get(("goteborg", p))
        if a is None or int(a["roster"]) != int(r[7]):
            rapport("FEL", f"jamforande Kommun Goteborg 2018 {p}: {r[7]} != aggregat {a and a['roster']}")
    elif r[1] == "Göteborg" and r[7] is not None:
        rapport("INFO", f"jamforande Kommun Goteborg: rad utan partikod {r[2]} {r[7]}")
print(f"jamforande statistik Kommun Goteborg: {n} partier jamforda")

# 6b. mandat 2018 ur Mandatfordelning-jamforelser
wb = openpyxl.load_workbook(FIL_MANDAT_JAMF, read_only=True)
m_csv = {(r["val"], r["niva"], r["parti_kalla"]): r for r in mandat}
def jmf_mandat(flik, rubrik, val, niva):
    rows = list(wb[flik].iter_rows(values_only=True, min_row=2))
    i = next(i for i, r in enumerate(rows) if r[0] == rubrik)
    tot = rows[i][1]
    s = 0
    # blocket for ett valomrade slutar nar partiernas 2018-mandat summerar till totalen
    for r in rows[i + 1:]:
        if r[0] is None or s >= tot:
            break
        p = NAMN2KOD.get(r[0])
        m18 = i0(r[1])
        if p is None:
            if m18:
                rapport("INFO", f"mandatjamforelse {flik} {rubrik}: {r[0]} {m18} mandat, ingen partikod")
            continue
        s += m18
        c = m_csv.get((val, niva, p))
        if m18 and (c is None or int(c["mandat"]) != m18):
            rapport("FEL", f"mandatjamforelse {flik} {rubrik} {p}: {m18} != mandat_2018.csv {c and c['mandat']}")
        if not m18 and c is not None and int(c["mandat"]):
            rapport("FEL", f"mandatjamforelse {flik} {rubrik} {p}: 0 != mandat_2018.csv {c['mandat']}")
    csum = sum(int(r["mandat"]) for r in mandat if r["val"] == val and r["niva"] == niva)
    print(f"mandatjamforelse {flik} {rubrik}: total 2018 {tot}, summa partier {s}, mandat_2018.csv summa {csum}")
    if tot != csum:
        rapport("FEL", f"mandatjamforelse {flik} {rubrik}: total {tot} != mandat_2018.csv {csum}")
jmf_mandat("Riksdag", "Sverige", "rd", "riket")
jmf_mandat("Region", "Västra Götaland", "rf", "vgregion")
jmf_mandat("Kommun", "Göteborg", "kf", "goteborg")
# mandat_2018.csv: Goteborgsrader mot aggregat
for r in mandat:
    if r["niva"] == "goteborg":
        a = next((x for x in aggregat[r["val"]] if x["niva"] == "goteborg" and x["parti_kalla"] == r["parti_kalla"]), None)
        if a is None or a["roster"] != r["roster"] or a["andel"] != r["andel"]:
            rapport("FEL", f"mandat_2018 {r['val']} goteborg {r['parti_kalla']}: roster/andel {r['roster']}/{r['andel']} != aggregat {a and (a['roster'], a['andel'])}")
        if int(r["fasta_mandat"]) + int(r["utjamningsmandat"]) != int(r["mandat"]):
            rapport("FEL", f"mandat_2018 {r['val']} {r['niva']} {r['parti_kalla']}: fasta+utjamning != mandat")
# mandat_2018.csv mot 2018_mandat.xlsx direkt (Goteborg rd/rf/kf)
hdr, rader = las_flik(FIL_MANDAT, "Mandatfördelning")
ix = {h: i for i, h in enumerate(hdr)}
for r in rader:
    vt = r[ix["VALTYP"]]
    if vt == "R" and r[ix["VALKRETSKOD"]] == 16 and r[ix["VALKRETS"]] == "Göteborgs kommun":
        key = ("rd", "goteborg")
    elif vt == "L" and r[ix["LÄN"]] == 14 and r[ix["VALKRETSKOD"]] == 1 and r[ix["VALKRETS"]] == "Göteborgs kommun":
        key = ("rf", "goteborg")
    elif vt == "K" and r[ix["LÄN"]] == 14 and r[ix["VALOMRÅDE"]] == "Göteborg":
        key = ("kf", "goteborg")
    else:
        continue
    c = m_csv.get(key + (r[ix["PARTIFÖRKORTNING"]],))
    if c is None or int(c["roster"]) != r[ix["RÖSTER"]] or int(c["mandat"]) != i0(r[ix["SUMMA MANDAT"]]) \
            or int(c["fasta_mandat"]) != i0(r[ix["FASTA MANDAT"]]) or int(c["utjamningsmandat"]) != i0(r[ix["UTJÄMNINGSMANDAT"]]):
        rapport("FEL", f"mandat_2018 {key} {r[ix['PARTIFÖRKORTNING']]}: skiljer fran 2018_mandat.xlsx")

# 6c. rostberattigade per distrikt ur reparerade filer
for v in ("rd", "rf", "kf"):
    wb = openpyxl.load_workbook(FIL_ROSTBER[v], read_only=True)
    ws = wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    hdr = rows[0]
    kolv = [i for i, h in enumerate(hdr) if isinstance(h, str) and h.startswith(("Män/", "Kvinnor/"))]
    ivd = hdr.index("valdistrikt_id"); ikom = hdr.index("kommun_id")
    tot = {}
    for r in rows[1:]:
        if str(r[ikom]) == "1480":
            tot[str(r[ivd]).zfill(8)] = (sum(i0(r[i]) for i in kolv), r[hdr.index("valdistrikt_namn")])
    if not tot:
        rapport("INFO", f"rostberattigade {v}: inga Goteborgsrader hittade, kommun_id exempel {rows[1][ikom]!r}")
        continue
    d_by = {r["kod"]: int(r["rostberattigade"]) for r in distrikt[v]}
    diff = [(k, d_by.get(k), t[0]) for k, t in tot.items() if d_by.get(k) != t[0]]
    saknas = sorted(set(d_by) - set(tot) - {"14800000"})
    print(f"rostberattigade {v}: {len(tot)} Goteborgsdistrikt i reparerad fil, summa {sum(t[0] for t in tot.values())}, "
          f"distriktsfil summa {sum(d_by.values())}; avvikelser {len(diff)}; saknas i reparerad {saknas[:5]}")
    for k, a, b in diff[:10]:
        rapport("FEL", f"rostberattigade {v} {k}: distriktsfil {a} != reparerad fil {b}")
    if len(diff) > 10:
        rapport("FEL", f"rostberattigade {v}: ytterligare {len(diff)-10} avvikelser")

# ------------------------------------------------------------------ 7. stickprov: tre Majornadistrikt, tre partital mot rakcell
random.seed(2018)
urval = random.sample(MAJORNA, 3)
print(f"stickprov distrikt: {urval}")
KOL_HTML = {"OGEJ": "OGEJ", "BLANK": "BLANK", "OG": "OG"}
for v in ("rd", "rf", "kf"):
    hdr, rader = las_flik(FIL[v], FLIK[v][0])
    wb = openpyxl.load_workbook(FIL[v], read_only=True)
    ws = wb[FLIK[v][0]]
    rs = {(r["kod"], r["parti_kalla"]): r for r in roster[v]}
    for kod in urval:
        radnr = kalla_rad[v][kod]
        pval = [p for p, a in kalla[v][kod]["partier"].items() if a > 0]
        pp = random.sample(pval, 3)
        for p in pp:
            kol = partikol[v][p]
            cell = ws.cell(row=radnr, column=kol + 1).value
            csvv = rs[(kod, p)]["roster"]
            ok = i0(cell) == int(csvv)
            print(f"  stickprov {v} {kod} {p}: {os.path.basename(FIL[v])} flik '{FLIK[v][0]}' rad {radnr} "
                  f"kolumn {get_column_letter(kol+1)} ({hdr[kol]}) = {cell!r}; CSV {csvv} {'OK' if ok else 'AVVIKELSE'}")
            if not ok:
                rapport("FEL", f"stickprov {v} {kod} {p}: cell {cell} != CSV {csvv}")

# ------------------------------------------------------------------ 8. historik.val.se distriktssidor (HTML) for stickprovsdistrikten
def hamta(V, kod):
    vd = kod[4:]
    fil = os.path.join(HTML_CACHE, f"{V}_14_80_{vd}.html")
    if not os.path.exists(fil):
        url = HTML_URL.format(V=V, vd=vd)
        try:
            with urllib.request.urlopen(url, timeout=30) as u:
                data = u.read()
        except Exception as e:
            rapport("INFO", f"HTML {url}: {e}")
            return None
        open(fil, "wb").write(data)
    return open(fil, "rb").read().decode("utf-8", "replace")

def tolka_html(t, kanda):
    """Tolkar distriktssidan. Returnerar (huvud, ovriga, tal) dar huvud och ovriga ar dict
    forkortning -> (antal, andel-strang) for de tva tabellerna, och tal har GILTIGA, ROSTBER,
    VDT (rostande) och valdeltagande. Tomma celler (0 roster) ar borttagna i sidan."""
    t = html.unescape(re.sub(r"<[^>]+>", "|", t))
    t = re.sub(r"\s+", " ", t)
    t = re.sub(r"( ?\| ?)+", "|", t)
    start = t.find("Röstfördelning - valdistrikt")  # fore tabellen finns en teckenforklaring med bara forkortningar
    slut = t.find("Röster på partier som ej beställt valsedlar - valdistrikt")
    if start > 0:
        t = t[start:slut if slut > start else None]
    tok = t.split("|")
    STOR = "Röstfördelning övriga partier"
    NOISE = {"Röster på partier som ej beställt valsedlar 2018"}
    huvud, ovriga, tal = {}, {}, {}
    ut = huvud
    i = 0
    while i < len(tok):
        x = tok[i]
        if x.startswith(STOR):
            ut = ovriga
        if x in kanda or x in ("OGEJ", "BLANK", "OG", "VDT", "ÖVR"):
            j = i + 2  # hoppa over partinamnet
            while j < len(tok) and tok[j] in NOISE:
                j += 1
            antal, andel = 0, "0,00%"
            if j < len(tok) and re.fullmatch(r"\d[\d ]*", tok[j]):
                antal = int(tok[j].replace(" ", ""))
                if j + 1 < len(tok) and re.fullmatch(r"\d+,\d\d%", tok[j + 1]):
                    andel = tok[j + 1]
            ut.setdefault(x, (antal, andel))
            i += 1
            continue
        m = re.fullmatch(r"Giltiga röster", x)
        if m and i + 1 < len(tok) and tok[i + 1].isdigit():
            tal["GILTIGA"] = int(tok[i + 1])
        if x == "Antal röstberättigade" and i + 1 < len(tok) and tok[i + 1].isdigit():
            tal["ROSTBER"] = int(tok[i + 1])
        i += 1
    if "VDT" in huvud:
        tal["ROSTANDE"], tal["VALDELT"] = huvud["VDT"]
    return huvud, ovriga, tal


def pct(v):
    return f"{float(v):.2f}".replace(".", ",") + "%"


for v, V in (("rd", "R"), ("rf", "L"), ("kf", "K")):
    d_by = {r["kod"]: r for r in distrikt[v]}
    o_by = {r["kod"]: r for r in ogilt[v]}
    kanda = set(partikol[v])
    for kod in urval:
        t = hamta(V, kod)
        if t is None:
            continue
        huvud, ovriga, tal = tolka_html(t, kanda)
        if not huvud:
            rapport("INFO", f"HTML {V} {kod}: kunde inte tolka sidan")
            continue
        n_ok = n_fel = 0
        rs = {r["parti_kalla"]: r for r in roster[v] if r["kod"] == kod}
        # partier i bada tabellerna mot CSV (antal och andel)
        for tabell in (huvud, ovriga):
            for p, (antal, andel) in tabell.items():
                if p in ("OGEJ", "BLANK", "OG", "VDT") or (p == "ÖVR" and tabell is huvud):
                    continue
                r = rs.get(p)
                if r is None:
                    if antal:
                        rapport("FEL", f"HTML {V} {kod}: parti {p} {antal} saknas i roster-CSV"); n_fel += 1
                elif int(r["roster"]) != antal or pct(r["andel"]) != andel:
                    rapport("FEL", f"HTML {V} {kod} {p}: {antal} {andel} != CSV {r['roster']} {r['andel']}"); n_fel += 1
                else:
                    n_ok += 1
        # OVR i huvudtabellen = summan av ovriga-tabellen = summan i CSV av de partier som inte star i huvudtabellen
        if "ÖVR" in huvud:
            csv_ovr = sum(int(r["roster"]) for p, r in rs.items() if p not in huvud)
            if huvud["ÖVR"][0] != csv_ovr:
                rapport("FEL", f"HTML {V} {kod} ÖVR: {huvud['ÖVR'][0]} != CSV-summa av ovriga partier {csv_ovr}"); n_fel += 1
            else:
                n_ok += 1
        # CSV-partier med roster som inte finns pa sidan
        for p, r in rs.items():
            if int(r["roster"]) and p not in huvud and p not in ovriga:
                rapport("FEL", f"HTML {V} {kod}: CSV-parti {p} {r['roster']} finns inte pa sidan"); n_fel += 1
        d = d_by[kod]; o = o_by[kod]
        jm = [("GILTIGA", tal.get("GILTIGA"), int(d["giltiga"])), ("ROSTBER", tal.get("ROSTBER"), int(d["rostberattigade"])),
              ("ROSTANDE", tal.get("ROSTANDE"), int(d["rostande"])),
              ("BLANK", huvud.get("BLANK", (0,))[0], int(o["blanka"])), ("OG", huvud.get("OG", (0,))[0], int(o["ogiltiga_ovriga"])),
              ("OGEJ", huvud.get("OGEJ", (0,))[0], int(o["ogiltiga_ej_anmalda"]))]
        for k, h, c in jm:
            if h != c:
                rapport("FEL", f"HTML {V} {kod} {k}: {h} != CSV {c}"); n_fel += 1
            else:
                n_ok += 1
        if tal.get("VALDELT") != pct(d["valdeltagande"]):
            rapport("FEL", f"HTML {V} {kod} valdeltagande {tal.get('VALDELT')} != CSV {d['valdeltagande']}"); n_fel += 1
        else:
            n_ok += 1
        print(f"HTML {V} {kod}: {n_ok} varden lika, {n_fel} avvikelser; huvudtabell {sorted(huvud)}; ovriga {sorted(ovriga)}")

# ------------------------------------------------------------------ 9. partier_2018.csv och normalisering
pset = {(r["val"], r["parti_kalla"]) for r in partier}
for v in ("rd", "rf", "kf"):
    used = {r["parti_kalla"] for r in roster[v]}
    saknas = used - {p for vv, p in pset if vv == v}
    extra = {p for vv, p in pset if vv == v} - used
    if saknas or extra:
        rapport("FEL", f"partier_2018 {v}: saknas {sorted(saknas)}, extra {sorted(extra)}")
for r in partier:
    if r["parti"] != normalisera(r["parti_kalla"]):
        rapport("FEL", f"partier_2018 {r['val']} {r['parti_kalla']}: parti {r['parti']}")
utan_namn = [(r["val"], r["parti_kalla"]) for r in partier if not r["namn"]]
print(f"partier_2018 utan namn: {utan_namn}")

# ------------------------------------------------------------------ 10. Majornasummor i noten
t = collections.Counter(); g = r_ = rb = 0
for r in roster["rd"]:
    if r["kod"] in MAJORNA:
        t[r["parti_kalla"]] += int(r["roster"])
for r in distrikt["rd"]:
    if r["kod"] in MAJORNA:
        g += int(r["giltiga"]); r_ += int(r["rostande"]); rb += int(r["rostberattigade"])
print(f"22 Majornadistrikt rd: giltiga {g}, rostande {r_}, rostber {rb}, valdeltagande {100*r_/rb:.2f}, {dict(t.most_common(9))}")
for v in ("rf", "kf"):
    g = r_ = rb = 0
    for r in distrikt[v]:
        if r["kod"] in MAJORNA:
            g += int(r["giltiga"]); r_ += int(r["rostande"]); rb += int(r["rostberattigade"])
    t = collections.Counter()
    for r in roster[v]:
        if r["kod"] in MAJORNA:
            t[r["parti_kalla"]] += int(r["roster"])
    print(f"22 Majornadistrikt {v}: giltiga {g}, rostande {r_}, rostber {rb}; {dict(t.most_common(12))}")

print("=== FEL:", len(fel))
for x in fel:
    print("  ", x)
print("=== INFO:", len(info))
for x in info:
    print("  ", x)
sys.exit(1 if fel else 0)
