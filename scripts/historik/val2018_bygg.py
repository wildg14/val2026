# -*- coding: utf-8 -*-
"""
val2018_bygg.py - valet 2018 for Goteborgs kommun (1480) per valdistrikt.

Laser Valmyndighetens xlsx-filer (2018-10-02) for riksdag (R), landsting (L)
och kommun (K) per valdistrikt samt 2018_mandat.xlsx (2018-11-06) och skriver
langa CSV-filer till data/historik/. Kontrollerar koder och namn mot
shapefilen 2018_valgeografi_valdistrikt (alla_valdistrikt.dbf).

Kors med:
  <venv>/bin/python scripts/historik/val2018_bygg.py

Alla tal kommer direkt ur kallfilerna. Aggregat ar summor av distriktsraderna
i samma fil, andelar i aggregaten ar berakade som 100*roster/giltiga.
"""
import csv
import collections
import os
import sys

import openpyxl
from dbfread import DBF

# ---------------------------------------------------------------- kallor
KALLMAPP = "/Users/daniel/code/Temp/Historiska dokument"
FIL_R = os.path.join(KALLMAPP, "2018_R_per_valdistrikt.xlsx")
FIL_L = os.path.join(KALLMAPP, "2018_L_per_valdistrikt.xlsx")
FIL_K = os.path.join(KALLMAPP, "2018_K_per_valdistrikt.xlsx")
FIL_MANDAT = os.path.join(KALLMAPP, "2018_mandat.xlsx")
SHP_DBF = ("/private/tmp/claude-501/-Users-daniel-code-Temp/"
           "f347baf2-1af3-43b3-b36d-e80b868ede0e/scratchpad/unz/"
           "2018_valgeografi_valdistrikt/alla_valdistrikt.dbf")
# Deltagande partier 2018 (partibeteckning per forkortning), hamtad 2026-09-03 fran
# https://historik.val.se/val/val2018/valsedlar/partier/deltagande_partier.skv (ISO-8859-1)
FIL_DELTAGANDE = ("/private/tmp/claude-501/-Users-daniel-code-Temp/"
                  "f347baf2-1af3-43b3-b36d-e80b868ede0e/scratchpad/dl2018/deltagande_partier.skv")
UTMAPP = "/Users/daniel/code/Temp/data/historik"

AR = 2018
LAN_GBG = 14
KOM_GBG = 80

# val: (kod i utdata, bokstav i kallan, flik antal, flik procent, fil)
VAL = [
    ("rd", "R", "R antal", "R procent", FIL_R),
    ("rf", "L", "L antal", "L procent", FIL_L),
    ("kf", "K", "K antal", "K procent", FIL_K),
]

# Kolumnnamn i kallfilerna (identiska i R, L, K)
KOL_LAN, KOL_KOM, KOL_VKK, KOL_VDK = "LÄNSKOD", "KOMMUNKOD", "VALKRETSKOD", "VALDISTRIKTSKOD"
KOL_VKN, KOL_VDN = "VALKRETSNAMN", "VALDISTRIKTSNAMN"
KOL_OGEJ, KOL_BLANK, KOL_OG = "OGEJ", "BLANK", "OG"
KOL_GILTIGA, KOL_ROSTANDE, KOL_ROSTBER, KOL_VALDELT = (
    "RÖSTER GILTIGA", "RÖSTANDE", "RÖSTBERÄTTIGADE", "VALDELTAGANDE")

# Normalisering av partikod
NORMALISERING = {"FP": "L", "DEM": "D", "KP": "K"}


def normalisera(parti_kalla):
    p = str(parti_kalla)
    return NORMALISERING.get(p, p.upper())


def f2(v):
    """Tal med tva decimaler, punkt som decimaltecken; tomt for None."""
    if v is None or v == "":
        return ""
    return f"{float(v):.2f}"


def i0(v):
    return int(v) if v is not None else 0


def skriv_csv(namn, kolumner, rader):
    vag = os.path.join(UTMAPP, namn)
    with open(vag, "w", encoding="utf-8", newline="") as f:
        w = csv.writer(f, delimiter=";", lineterminator="\n")
        w.writerow(kolumner)
        w.writerows(rader)
    print(f"skrev {vag} ({len(rader)} rader)")
    return vag


def las_flik(fil, flik):
    wb = openpyxl.load_workbook(fil, read_only=True)
    ws = wb[flik]
    rader = list(ws.iter_rows(values_only=True))
    hdr = [str(h) if h is not None else "" for h in rader[0]]
    return hdr, rader[1:]


def kod8(r, ix):
    return f"{int(r[ix[KOL_LAN]]):02d}{int(r[ix[KOL_KOM]]):02d}{int(r[ix[KOL_VDK]]):04d}"


def las_shapefil():
    d = DBF(SHP_DBF, encoding="utf-8")
    return {r["VD"]: r["VD_NAMN"] for r in d if str(r["VD"]).startswith("1480")}


def bearbeta_val(val, bokstav, flik_antal, flik_procent, fil, shp, mandat, avvikelser):
    print(f"=== {val} ({os.path.basename(fil)})")
    hdr, rader = las_flik(fil, flik_antal)
    hdr_p, rader_p = las_flik(fil, flik_procent)
    if hdr != hdr_p:
        avvikelser.append(f"{val}: rubrikerna skiljer sig mellan {flik_antal} och {flik_procent}")
    ix = {h: i for i, h in enumerate(hdr)}
    i_ogej = ix[KOL_OGEJ]
    partikol = list(range(ix[KOL_VDN] + 1, i_ogej))  # partikolumner: efter VALDISTRIKTSNAMN till fore OGEJ
    partier_alla = [hdr[i] for i in partikol]

    # nyckel per rad for att para antal och procent
    nyckel = lambda r: (r[ix[KOL_LAN]], r[ix[KOL_KOM]], r[ix[KOL_VDK]])
    proc = {nyckel(r): r for r in rader_p}
    if len(proc) != len(rader_p) or set(proc) != set(nyckel(r) for r in rader):
        avvikelser.append(f"{val}: antal- och procentflikarna har olika distriktsnycklar")

    gbg = [r for r in rader if r[ix[KOL_LAN]] == LAN_GBG and r[ix[KOL_KOM]] == KOM_GBG]
    lan14 = [r for r in rader if r[ix[KOL_LAN]] == LAN_GBG]
    print(f"rader riket {len(rader)}, lan 14 {len(lan14)}, Goteborg {len(gbg)}")

    # --- kontroller per distrikt
    for r in rader:
        g = i0(r[ix[KOL_GILTIGA]])
        s = sum(i0(r[i]) for i in partikol)
        if s != g:
            avvikelser.append(f"{val} {kod8(r, ix)}: partisumma {s} != giltiga {g}")
        og = i0(r[i_ogej]) + i0(r[ix[KOL_BLANK]]) + i0(r[ix[KOL_OG]])
        if g + og != i0(r[ix[KOL_ROSTANDE]]):
            avvikelser.append(f"{val} {kod8(r, ix)}: giltiga+ogiltiga {g + og} != rostande {r[ix[KOL_ROSTANDE]]}")
        rb = i0(r[ix[KOL_ROSTBER]])
        if rb:
            vd = round(100 * i0(r[ix[KOL_ROSTANDE]]) / rb, 2)
            if abs(vd - float(r[ix[KOL_VALDELT]] or 0)) > 0.011:
                avvikelser.append(f"{val} {kod8(r, ix)}: valdeltagande {r[ix[KOL_VALDELT]]} != {vd}")

    # --- kontroll mot shapefil (Goteborg)
    koder_fil = {kod8(r, ix): r[ix[KOL_VDN]] for r in gbg}
    saknas_i_shp = sorted(k for k in koder_fil if k not in shp)
    saknas_i_fil = sorted(k for k in shp if k not in koder_fil)
    namn_diff = sorted(k for k in koder_fil if k in shp and koder_fil[k] != shp[k])
    print(f"koder i fil men inte i shapefil: {saknas_i_shp}")
    print(f"koder i shapefil men inte i fil: {saknas_i_fil}")
    print(f"namnskillnader fil/shapefil: {namn_diff}")
    if saknas_i_shp != ["14800000"] or saknas_i_fil or namn_diff:
        avvikelser.append(f"{val}: kod/namn-avvikelse mot shapefil: {saknas_i_shp} {saknas_i_fil} {namn_diff}")

    # --- partier som forekommer i Goteborg
    partier_gbg = [hdr[i] for i in partikol if any(i0(r[i]) for r in gbg)]
    print(f"partier med roster i Goteborg: {partier_gbg}")

    # --- roster per distrikt (Goteborg)
    roster = []
    for r in sorted(gbg, key=lambda r: kod8(r, ix)):
        kod = kod8(r, ix)
        rp = proc[nyckel(r)]
        for i in partikol:
            p = hdr[i]
            if p not in partier_gbg:
                continue
            antal = i0(r[i])
            andel = rp[i]
            if andel is None:
                if antal == 0:
                    andel = 0.0
                else:
                    avvikelser.append(f"{val} {kod} {p}: antal {antal} men ingen procent")
            roster.append([AR, val, kod, r[ix[KOL_VDN]], p, normalisera(p), antal, f2(andel)])
    skriv_csv(f"roster_{AR}_{val}.csv",
              ["ar", "val", "kod", "namn", "parti_kalla", "parti", "roster", "andel"], roster)

    # --- distrikt (Goteborg)
    distrikt, ogiltiga = [], []
    for r in sorted(gbg, key=lambda r: kod8(r, ix)):
        kod = kod8(r, ix)
        ogej, blank, og = i0(r[i_ogej]), i0(r[ix[KOL_BLANK]]), i0(r[ix[KOL_OG]])
        distrikt.append([AR, val, kod, r[ix[KOL_VDN]], r[ix[KOL_VKN]] or "",
                         i0(r[ix[KOL_GILTIGA]]), blank, ogej + og, ogej + blank + og,
                         i0(r[ix[KOL_ROSTANDE]]), i0(r[ix[KOL_ROSTBER]]),
                         f2(r[ix[KOL_VALDELT]]), os.path.basename(fil)])
        ogiltiga.append([AR, val, kod, r[ix[KOL_VDN]], ogej, blank, og, os.path.basename(fil)])
    skriv_csv(f"distrikt_{AR}_{val}.csv",
              ["ar", "val", "kod", "namn", "valkrets", "giltiga", "blanka", "ogiltiga_ovriga",
               "ogiltiga", "rostande", "rostberattigade", "valdeltagande", "kalla_fil"], distrikt)
    skriv_csv(f"ogiltiga_{AR}_{val}.csv",
              ["ar", "val", "kod", "namn", "ogiltiga_ej_anmalda", "blanka", "ogiltiga_ovriga", "kalla_fil"],
              ogiltiga)

    # --- aggregat: riket, vgregion (lan 14), goteborg
    aggregat = []
    summor = {}
    for niva, delm in [("riket", rader), ("vgregion", lan14), ("goteborg", gbg)]:
        giltiga = sum(i0(r[ix[KOL_GILTIGA]]) for r in delm)
        rostande = sum(i0(r[ix[KOL_ROSTANDE]]) for r in delm)
        rostber = sum(i0(r[ix[KOL_ROSTBER]]) for r in delm)
        summor[niva] = {"giltiga": giltiga, "rostande": rostande, "rostberattigade": rostber}
        for i in partikol:
            p = hdr[i]
            s = sum(i0(r[i]) for r in delm)
            if s == 0:
                continue
            summor[niva][p] = s
            aggregat.append([AR, val, niva, p, normalisera(p), s, f2(100 * s / giltiga) if giltiga else "",
                             giltiga, rostande, rostber])
        print(f"{niva}: giltiga {giltiga}, rostande {rostande}, rostberattigade {rostber}, "
              f"valdeltagande {f2(100 * rostande / rostber) if rostber else ''}")
    skriv_csv(f"aggregat_{AR}_{val}.csv",
              ["ar", "val", "niva", "parti_kalla", "parti", "roster", "andel", "giltiga", "rostande",
               "rostberattigade"], aggregat)

    # --- avstamning mot 2018_mandat.xlsx (Goteborg)
    m = mandat[val]["goteborg"]
    for parti, (roster_m, andel_m) in sorted(m.items()):
        s = summor["goteborg"].get(parti)
        a = round(100 * s / summor["goteborg"]["giltiga"], 2) if s else None
        ok = (s == roster_m) and (a is not None and abs(a - andel_m) < 0.011)
        print(f"  mandatfil Goteborg {parti}: roster {roster_m} andel {andel_m} | summerat {s} {a} {'OK' if ok else 'AVVIKELSE'}")
        if not ok:
            avvikelser.append(f"{val} {parti}: mandatfilen {roster_m}/{andel_m} vs summerat {s}/{a}")
    if val == "rd":
        # Mandatfilen listar per valkrets bara partier som fick mandat dar. Summan over
        # valkretsarna ar darfor bara lika med rikssumman for partier med mandat i alla 29.
        for parti, (roster_m, antal_vk) in sorted(mandat["rd"]["riket_roster"].items()):
            s = summor["riket"].get(parti)
            if antal_vk == mandat["rd"]["antal_valkretsar"]:
                ok = s == roster_m
                print(f"  mandatfil riket {parti} ({antal_vk} valkretsar): {roster_m} | summerat {s} {'OK' if ok else 'AVVIKELSE'}")
                if not ok:
                    avvikelser.append(f"rd riket {parti}: mandatfilen {roster_m} vs summerat {s}")
            else:
                print(f"  mandatfil riket {parti} ({antal_vk} valkretsar, ofullstandig): {roster_m} | summerat {s}")
    return summor, partier_gbg, partier_alla


def las_mandat():
    """Laser fliken Mandatfordelning. Returnerar strukturer for avstamning och mandat_2018.csv."""
    hdr, rader = las_flik(FIL_MANDAT, "Mandatfördelning")
    ix = {h: i for i, h in enumerate(hdr)}
    c = lambda r, k: r[ix[k]]
    ut = {"rd": {}, "rf": {}, "kf": {}, "rader": rader, "ix": ix}
    ut["rd"]["goteborg"] = {c(r, "PARTIFÖRKORTNING"): (c(r, "RÖSTER"), c(r, "ANDEL AV GILTIGA RÖSTER"))
                            for r in rader if c(r, "VALTYP") == "R" and c(r, "VALKRETSKOD") == 16
                            and c(r, "VALKRETS") == "Göteborgs kommun"}
    riket = collections.defaultdict(lambda: [0, 0])
    valkretsar = set()
    for r in rader:
        if c(r, "VALTYP") == "R":
            riket[c(r, "PARTIFÖRKORTNING")][0] += i0(c(r, "RÖSTER"))
            riket[c(r, "PARTIFÖRKORTNING")][1] += 1
            valkretsar.add(c(r, "VALKRETSKOD"))
    ut["rd"]["riket_roster"] = {p: tuple(v) for p, v in riket.items()}
    ut["rd"]["antal_valkretsar"] = len(valkretsar)
    ut["rf"]["goteborg"] = {c(r, "PARTIFÖRKORTNING"): (c(r, "RÖSTER"), c(r, "ANDEL AV GILTIGA RÖSTER"))
                            for r in rader if c(r, "VALTYP") == "L" and c(r, "LÄN") == LAN_GBG
                            and c(r, "VALKRETSKOD") == 1 and c(r, "VALKRETS") == "Göteborgs kommun"}
    ut["kf"]["goteborg"] = {c(r, "PARTIFÖRKORTNING"): (c(r, "RÖSTER"), c(r, "ANDEL AV GILTIGA RÖSTER"))
                            for r in rader if c(r, "VALTYP") == "K" and c(r, "LÄN") == LAN_GBG
                            and c(r, "VALOMRÅDE") == "Göteborg"}
    return ut


def skriv_mandat(mandat, summor):
    hdr_ix = mandat["ix"]
    rader = mandat["rader"]
    c = lambda r, k: r[hdr_ix[k]]
    fil = os.path.basename(FIL_MANDAT)
    ut = []
    namn = {}

    def rad(val, niva, valkrets, r):
        namn[c(r, "PARTIFÖRKORTNING")] = c(r, "PARTIBETECKNING")
        return [AR, val, niva, valkrets, c(r, "PARTIFÖRKORTNING"), normalisera(c(r, "PARTIFÖRKORTNING")),
                c(r, "RÖSTER"), f2(c(r, "ANDEL AV GILTIGA RÖSTER")), i0(c(r, "FASTA MANDAT")),
                i0(c(r, "UTJÄMNINGSMANDAT")), i0(c(r, "SUMMA MANDAT")), fil]

    # riksdagen totalt: summa over alla valkretsar
    tot = collections.OrderedDict()
    for r in rader:
        if c(r, "VALTYP") != "R":
            continue
        p = c(r, "PARTIFÖRKORTNING")
        namn[p] = c(r, "PARTIBETECKNING")
        t = tot.setdefault(p, [0, 0, 0, 0])
        t[0] += i0(c(r, "RÖSTER")); t[1] += i0(c(r, "FASTA MANDAT"))
        t[2] += i0(c(r, "UTJÄMNINGSMANDAT")); t[3] += i0(c(r, "SUMMA MANDAT"))
    g_riket = summor["rd"]["riket"]["giltiga"]
    for p, t in tot.items():
        roster = summor["rd"]["riket"][p]  # rikssumma ur R-filen, mandatfilens summa ar ofullstandig
        ut.append([AR, "rd", "riket", "HELA LANDET", p, normalisera(p), roster, f2(100 * roster / g_riket),
                   t[1], t[2], t[3], fil + " (mandat), " + os.path.basename(FIL_R) + " (roster)"])
        print(f"riksdagen totalt {p}: roster {roster} (mandatfilens summa {t[0]}), mandat {t[3]}")
    print("riksdagen totalt:", {p: t[3] for p, t in tot.items()}, "summa", sum(t[3] for t in tot.values()))
    for r in rader:
        if c(r, "VALTYP") == "R" and c(r, "VALKRETSKOD") == 16 and c(r, "VALKRETS") == "Göteborgs kommun":
            ut.append(rad("rd", "goteborg", c(r, "VALKRETS"), r))
    # landsting: Vastra Gotaland totalt (summa over 5 valkretsar) och Goteborgs kommun
    tot = collections.OrderedDict()
    for r in rader:
        if c(r, "VALTYP") == "L" and c(r, "LÄN") == LAN_GBG:
            p = c(r, "PARTIFÖRKORTNING")
            namn[p] = c(r, "PARTIBETECKNING")
            t = tot.setdefault(p, [0, 0, 0, 0])
            t[0] += i0(c(r, "RÖSTER")); t[1] += i0(c(r, "FASTA MANDAT"))
            t[2] += i0(c(r, "UTJÄMNINGSMANDAT")); t[3] += i0(c(r, "SUMMA MANDAT"))
    g_vg = summor["rf"]["vgregion"]["giltiga"]
    for p, t in tot.items():
        roster = summor["rf"]["vgregion"][p]  # lanssumma ur L-filen, mandatfilens summa ar ofullstandig
        ut.append([AR, "rf", "vgregion", "Västra Götalands läns landsting", p, normalisera(p), roster,
                   f2(100 * roster / g_vg), t[1], t[2], t[3],
                   fil + " (mandat), " + os.path.basename(FIL_L) + " (roster)"])
        print(f"VG-regionen totalt {p}: roster {roster} (mandatfilens summa {t[0]}), mandat {t[3]}")
    print("VG-regionen totalt:", {p: t[3] for p, t in tot.items()}, "summa", sum(t[3] for t in tot.values()))
    for r in rader:
        if c(r, "VALTYP") == "L" and c(r, "LÄN") == LAN_GBG and c(r, "VALKRETSKOD") == 1 \
                and c(r, "VALKRETS") == "Göteborgs kommun":
            ut.append(rad("rf", "goteborg", c(r, "VALKRETS"), r))
    for r in rader:
        if c(r, "VALTYP") == "K" and c(r, "LÄN") == LAN_GBG and c(r, "VALOMRÅDE") == "Göteborg":
            ut.append(rad("kf", "goteborg", c(r, "VALKRETS"), r))
    skriv_csv(f"mandat_{AR}.csv",
              ["ar", "val", "niva", "valkrets", "parti_kalla", "parti", "roster", "andel", "fasta_mandat",
               "utjamningsmandat", "mandat", "kalla_fil"], ut)
    return namn


def las_deltagande():
    """Partibeteckning per (val, forkortning eller partikod) ur deltagande_partier.skv.
    rd: VALTYP R (riksdagen), rf: VALTYP L och LÄNSKOD 14, kf: VALTYP K och VALOMRÅDESKOD 1480."""
    ut = {"rd": collections.defaultdict(set), "rf": collections.defaultdict(set), "kf": collections.defaultdict(set)}
    namn_k = set()
    if not os.path.exists(FIL_DELTAGANDE):
        print("deltagande_partier.skv saknas, partinamn bara ur mandatfilen")
        return ut
    with open(FIL_DELTAGANDE, encoding="iso-8859-1", newline="") as f:
        lasare = csv.DictReader(f, delimiter=";")
        for r in lasare:
            vt = r["VALTYP"]
            if vt == "R":
                val = "rd"
            elif vt == "L" and r["LÄNSKOD"] == "14":
                val = "rf"
            elif vt == "K" and r["VALOMRÅDESKOD"] == "1480":  # kommunen Goteborg
                val = "kf"
                namn_k.add(r["VALOMRÅDESNAMN"])
            else:
                continue
            if r["PARTIFÖRKORTNING"]:
                ut[val][r["PARTIFÖRKORTNING"]].add(r["PARTIBETECKNING"])
            ut[val][r["PARTIKOD"]].add(r["PARTIBETECKNING"])
    print(f"deltagande_partier.skv: kommunvalomrade 1480 = {sorted(namn_k)}, "
          f"antal koder rd {len(ut['rd'])}, rf {len(ut['rf'])}, kf {len(ut['kf'])}")
    return ut


def skriv_partier(partier_per_val, namn_mandat):
    """Partikoder som forekommer i Goteborgsfilerna, med partibeteckning dar kallorna har en."""
    # numeriska partikoder i K-filen forklaras pa fliken "Per valdistrikt"
    wb = openpyxl.load_workbook(FIL_K, read_only=True)
    numeriska = {}
    for r in wb["Per valdistrikt"].iter_rows(values_only=True):
        if isinstance(r[0], int) and r[1]:
            numeriska[str(r[0])] = r[1]
    deltagande = las_deltagande()
    ut = []
    for val, partier in partier_per_val.items():
        for p in partier:
            p = str(p)
            d = deltagande[val].get(p, set())
            if len(d) > 1:
                print(f"obs: {val} {p} har flera partibeteckningar i deltagande_partier.skv: {sorted(d)}")
            if namn_mandat.get(p):
                n, kalla = namn_mandat[p], "2018_mandat.xlsx PARTIBETECKNING"
            elif len(d) == 1:
                n, kalla = next(iter(d)), "deltagande_partier.skv PARTIBETECKNING"
            elif numeriska.get(p):
                n, kalla = numeriska[p], "2018_K_per_valdistrikt.xlsx flik Per valdistrikt"
            else:
                n, kalla = "", ""
            ut.append([AR, val, p, normalisera(p), n, kalla])
    skriv_csv(f"partier_{AR}.csv", ["ar", "val", "parti_kalla", "parti", "namn", "namn_kalla"], ut)


def main():
    os.makedirs(UTMAPP, exist_ok=True)
    avvikelser = []
    shp = las_shapefil()
    print(f"shapefil: {len(shp)} Goteborgsdistrikt")
    mandat = las_mandat()
    summor, partier_per_val = {}, {}
    for val, bokstav, fa, fp, fil in VAL:
        s, partier_gbg, _ = bearbeta_val(val, bokstav, fa, fp, fil, shp, mandat, avvikelser)
        summor[val] = s
        partier_per_val[val] = partier_gbg
    namn_mandat = skriv_mandat(mandat, summor)
    skriv_partier(partier_per_val, namn_mandat)
    print("=== avvikelser:", len(avvikelser))
    for a in avvikelser:
        print(" -", a)
    return 0 if not avvikelser else 1


if __name__ == "__main__":
    sys.exit(main())
