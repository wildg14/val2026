#!/usr/bin/env python
"""mandat_valkretsar.py - mandat och valkretsar for Goteborg 2002-2022.

Skriver tre filer till data/historik/:
  mandat_riksdag_riket.csv      ar;parti;mandat;kalla
  mandat_valkrets_goteborg.csv  ar;val;valkrets;parti;fasta;utjamning;totalt;kalla
  valkretsar_goteborg.csv       ar;val;valkrets;antal_mandat;rostberattigade;kalla

Alla tal kommer ur de kallfiler som anges nedan. Skriptet raknar bara summor
och skillnader (fasta = totalt - utjamning) och avbryter om kontrollsummorna
inte stammer. Kors med venv-pythonen som har xlrd, openpyxl och lxml.
"""
import collections
import csv
import json
import os
import re
import sys

import openpyxl
import xlrd
from lxml import etree

ROT = "/Users/daniel/code/Temp"
HIST = ROT + "/Historiska dokument"
SCRATCH = "/private/tmp/claude-501/-Users-daniel-code-Temp/f347baf2-1af3-43b3-b36d-e80b868ede0e/scratchpad"
UT = ROT + "/data/historik"
WEBB = ROT + "/docs/historik/kallor/webb"

# Valmyndighetens filer i Historiska dokument
F_VKM_R_2010 = HIST + "/valkretsmandat_R_2010.xls"
F_VKM_L_2010 = HIST + "/valkretsmandat_L_2010.xls"
F_VKM_K_2010 = HIST + "/valkretsmandat_K_2010.xls"
F_VKM_R_2014 = HIST + "/Valkretsmandat riksdag 2014.xls"
F_VKM_L_2014 = HIST + "/Valkretsmandat landsting 2014.xls"
F_VKM_R_1988_2014 = HIST + "/Valkretsmandat riksdag 1988-2014.xls"
F_MF_R_2010 = HIST + "/mandatfordelning_per_valkrets_R (1).xls"
F_MF_R_2014 = HIST + "/mandatfordelning_per_valkrets_R.xls"
F_MF_K_2010 = HIST + "/mandatfordelning_per_valkrets_K.xls"      # ar avgors i skriptet
F_MF_L_2011 = HIST + "/mandatfordelning_per_valkrets_L.xls"      # ar avgors i skriptet
F_MANDAT_2018 = HIST + "/2018_mandat.xlsx"
F_2018_R = HIST + "/2018_R_per_valdistrikt.xlsx"
F_2018_L = HIST + "/2018_L_per_valdistrikt.xlsx"
F_2018_K = HIST + "/2018_K_per_valdistrikt.xlsx"
# 2022
F_JAMF_2022 = ROT + "/Mandatfordelning-jamforelser-mellan-2018-och-2022.xlsx"
F_JAMF_RD_2022 = ROT + "/slutligt-valresultat-riksdagen-jamforande-statistik-2018-2022.xlsx"
F_VALDATA_2022 = ROT + "/data/valdata_2022.json"
F_RAW_2022 = {
    "rd": ROT + "/Roster-per-distrikt-slutligt-antal-roster-inklusive-totalt-valdeltagande-riksdagsvalet-2022.xlsx",
    "rf": ROT + "/roster-per-distrikt-slutligt-antal-roster-inklusive-totalt-valdeltagande-regionval-2022.xlsx",
    "kf": ROT + "/roster-per-distrikt-slutligt-antal-roster-inklusive-totalt-valdeltagande-kommunval-2022.xlsx",
}
# Valmyndighetens XML (ISO-8859-1)
XML_2006 = SCRATCH + "/dl2006"
XML_2010 = SCRATCH + "/unz/slutresultat__1_"
XML_2014 = SCRATCH + "/unz/slutresultat"
# Sparade webbsidor fran historik.val.se (kontroll och 2002 ars fasta/utjamning for Goteborg)
F_WEBB_2002 = WEBB + "/val_02_slutresultat_00R_00-text.html"
F_WEBB_2006 = WEBB + "/val2006_slutlig_R_rike_roster.html"

PARTI_NORM = {"FP": "L", "DEM": "D", "KP": "K"}


def norm(p):
    p = str(p).strip()
    return PARTI_NORM.get(p, p.upper())


def fail(msg):
    print("FEL:", msg)
    sys.exit(1)


def check(cond, msg):
    if not cond:
        fail(msg)
    print("OK:", msg)


def xls_rows(path, sheet=0):
    sh = xlrd.open_workbook(path).sheet_by_index(sheet)
    return [[sh.cell_value(r, c) for c in range(sh.ncols)] for r in range(sh.nrows)]


def xlsx_rows(path, sheet):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet]
    return [list(r) for r in ws.iter_rows(values_only=True)]


def as_int(v):
    if v in ("", None):
        return 0
    return int(v)


def xml_root(path):
    return etree.parse(path).getroot()


def giltiga(node, attr="MANDAT", utj="VARAV_UTJÄMNING"):
    """{parti_kalla: (mandat, utjamning)} for en XML-nod med GILTIGA-barn."""
    out = collections.OrderedDict()
    for g in node.findall("GILTIGA"):
        m = g.get(attr)
        if m is None:
            continue
        u = g.get(utj)
        out[g.get("PARTI")] = (int(m), None if u is None else int(u))
    return out


def rostberattigade(node, attr="RÖSTBERÄTTIGADE"):
    v = node.find("VALDELTAGANDE")
    return None if v is None or v.get(attr) is None else int(v.get(attr))


# ---------------------------------------------------------------- mandatfordelning_per_valkrets (xls, 2010 och 2014)

def las_mandatfordelning(path, rad_kod, sheet=0):
    """Laser Valmyndighetens mandatfordelning_per_valkrets-fil (R eller L).
    Returnerar {(lan, valkrets): {"namn":..., "fasta_antal":..., "parti": {p: (fasta, utj)}}}."""
    rows = xls_rows(path, sheet)
    hdr_i = [i for i, r in enumerate(rows) if r[0] == "LAN"][0]
    hdr = rows[hdr_i]
    out = collections.OrderedDict()
    for r in rows[hdr_i + 1:]:
        if r[0] in ("", None):
            continue
        d = {"namn": r[3], "fasta_antal": as_int(r[4]), "parti": collections.OrderedDict()}
        for c in range(5, len(hdr)):
            h = hdr[c]
            if not h or h.endswith(" utj"):
                continue
            fasta = as_int(r[c])
            utj = as_int(r[c + 1]) if c + 1 < len(hdr) and hdr[c + 1] == h + " utj" else 0
            if fasta or utj:
                d["parti"][h] = (fasta, utj)
        out[(int(r[0]), int(r[1]))] = d
    return out


def summera_riket(mf):
    tot = collections.Counter()
    fasta = collections.Counter()
    utj = collections.Counter()
    for d in mf.values():
        for p, (f, u) in d["parti"].items():
            fasta[p] += f
            utj[p] += u
            tot[p] += f + u
    return tot, fasta, utj


def las_mandatfordelning_K(path):
    """mandatfordelning_per_valkrets_K.xls: LAN, KOM, VALKRETS, NAMN_LAN, NAMN_KOM, NAMN, ANT_MAND, sedan en kolumn per parti."""
    rows = xls_rows(path)
    hdr = rows[0]
    out = collections.OrderedDict()
    for r in rows[1:]:
        if r[0] in ("", None):
            continue
        d = {"namn": r[5], "antal": as_int(r[6]), "parti": collections.OrderedDict()}
        for c in range(7, len(hdr)):
            if r[c] not in ("", None):
                d["parti"][hdr[c]] = as_int(r[c])
        out[(int(r[0]), int(r[1]), int(r[2]))] = d
    return out


# ---------------------------------------------------------------- huvudprogram

def main():
    os.makedirs(UT, exist_ok=True)
    riket = []      # ar;parti;mandat;kalla
    vkm = []        # ar;val;valkrets;parti;fasta;utjamning;totalt;kalla
    vkr = []        # ar;val;valkrets;antal_mandat;rostberattigade;kalla

    def add_riket(ar, parti_kalla, mandat, kalla):
        if mandat > 0:
            riket.append([ar, norm(parti_kalla), mandat, kalla])

    def add_vkm(ar, val, valkrets, parti_kalla, fasta, utj, totalt, kalla):
        if totalt in (0, None) and not fasta:
            return
        vkm.append([ar, val, valkrets, norm(parti_kalla), fasta, utj, totalt, kalla])

    def add_vkr(ar, val, valkrets, antal, rb, kalla):
        vkr.append([ar, val, valkrets, antal, rb, kalla])

    # ------------------------------------------------ 2006 och 2002 ur XML 2006 (FGVAL = 2002)
    r06 = xml_root(XML_2006 + "/slutresultat_00R.xml")
    check(r06.get("VALDAG") == "20060917" and r06.get("VALDAG_FGVAL") == "20020915",
          "2006 00R.xml ar valet 2006-09-17 med foregaende val 2002-09-15")
    nat = r06.find(".//NATION")
    k06 = "dl2006/slutresultat_00R.xml NATION/GILTIGA MANDAT"
    k02 = "dl2006/slutresultat_00R.xml NATION/GILTIGA MANDAT_FGVAL (foregaende val 2002)"
    m06 = giltiga(nat)
    m02 = giltiga(nat, "MANDAT_FGVAL", "INGEN")
    check(sum(m for m, u in m06.values()) == 349, "2006 riksdag summerar till 349 (XML)")
    check(sum(m for m, u in m02.values()) == 349, "2002 riksdag summerar till 349 (XML MANDAT_FGVAL)")
    for p, (m, u) in m06.items():
        add_riket(2006, p, m, k06)
    for p, (m, u) in m02.items():
        add_riket(2002, p, m, k02)

    # kontroll mot sparade webbsidor
    if os.path.exists(F_WEBB_2002):
        txt = re.sub(r"<[^>]*>", " ", open(F_WEBB_2002, encoding="utf-8").read())
        txt = re.sub(r"[ \t]+", " ", txt)
        # "M C FP KD S V MP NBP ÖVR OG Vdt%" ... "M(U) 55(5) 22(6) 48(4) 33(4) 144(2) 30(7) 17(11) 0(0) 0(0)"
        m = re.search(r"Sverige Summa.*?M\(U\) ([0-9()\s]+)", txt, re.S)
        webb02 = [int(x) for x in re.findall(r"(\d+)\(\d+\)", m.group(1))][:7]
        xml02 = [m02[p][0] for p in ("M", "C", "FP", "KD", "S", "V", "MP")]
        check(webb02 == xml02, "2002 riksdag riket: XML MANDAT_FGVAL = historik.val.se val_02 00-text.html M(U) %s" % webb02)
        m = re.search(r"Göteborgs kommun Summa.*?M\(U\) ([0-9()\s]+)", txt, re.S)
        gbg02 = re.findall(r"(\d+)\((\d+)\)", m.group(1))[:7]
        webb_gbg02 = collections.OrderedDict(zip(("M", "C", "FP", "KD", "S", "V", "MP"), [(int(a), int(b)) for a, b in gbg02]))
    else:
        webb_gbg02 = None
        print("VARNING: %s saknas, hoppar over webbkontrollen 2002" % F_WEBB_2002)
    if os.path.exists(F_WEBB_2006):
        txt = re.sub(r"<[^>]*>", " ", open(F_WEBB_2006, encoding="utf-8").read())
        txt = re.sub(r"\s+", " ", txt)
        w = {}
        for p in ("M", "C", "FP", "KD", "S", "V", "MP"):
            mm = re.search(r" %s [A-Za-zÅÄÖåäö\- ]+? (\d+) [\d,]+ (\d+) (\d+) [\d,]+ (\d+) " % p, txt)
            w[p] = (int(mm.group(2)), int(mm.group(4)))
        check(all(w[p] == (m06[p][0], m02[p][0]) for p in w),
              "2006 och 2002 riksdag riket: XML = historik.val.se val2006 R/rike/roster.html")
    else:
        print("VARNING: %s saknas, hoppar over webbkontrollen 2006" % F_WEBB_2006)

    # Goteborgs riksdagsvalkrets 1416: 2006 och 2002
    kr = [k for k in r06.iter("KRETS_RIKSDAG") if k.get("KOD") == "1416"][0]
    check(kr.get("NAMN") == "Göteborgs kommun" and kr.get("MANDAT_VALKRETS") == "17", "2006 riksdagsvalkrets 1416 Goteborgs kommun, 17 fasta mandat")
    g06 = giltiga(kr)
    g02 = giltiga(kr, "MANDAT_FGVAL", "INGEN")
    for p, (m, u) in g06.items():
        add_vkm(2006, "rd", "Göteborgs kommun", p, m - u, u, m, "dl2006/slutresultat_00R.xml KRETS_RIKSDAG 1416 GILTIGA MANDAT, VARAV_UTJÄMNING")
    for p, (m, u) in g02.items():
        if webb_gbg02 and p in webb_gbg02:
            check(webb_gbg02[p][0] == m, "2002 riksdag Goteborg %s: XML MANDAT_FGVAL %d = webbsidan" % (p, m))
            add_vkm(2002, "rd", "Göteborgs kommun", p, m - webb_gbg02[p][1], webb_gbg02[p][1], m,
                    "dl2006/slutresultat_00R.xml KRETS_RIKSDAG 1416 MANDAT_FGVAL; utjamning ur historik.val.se val_02 00-text.html M(U)")
        else:
            add_vkm(2002, "rd", "Göteborgs kommun", p, "", "", m, "dl2006/slutresultat_00R.xml KRETS_RIKSDAG 1416 MANDAT_FGVAL")
    vd = kr.find("VALDELTAGANDE")
    add_vkr(2006, "rd", "Göteborgs kommun", 17, int(vd.get("RÖSTBERÄTTIGADE")),
            "mandat: dl2006/slutresultat_00R.xml KRETS_RIKSDAG 1416 MANDAT_VALKRETS; rostberattigade valdagen: VALDELTAGANDE RÖSTBERÄTTIGADE")
    add_vkr(2002, "rd", "Göteborgs kommun", 17, int(vd.get("RÖSTBERÄTTIGADE_KLARA_VALDISTRIKT_FGVAL")),
            "mandat: Valkretsmandat riksdag 1988-2014.xls kolumn 2002; rostberattigade valdagen: dl2006/slutresultat_00R.xml KRETS_RIKSDAG 1416 VALDELTAGANDE RÖSTBERÄTTIGADE_KLARA_VALDISTRIKT_FGVAL")

    # Goteborgs landstingsvalkrets 1401: 2006 och 2002
    l06 = xml_root(XML_2006 + "/slutresultat_00L.xml")
    kl = [k for k in l06.iter("KRETS_LANDSTING") if k.get("KOD") == "1401"][0]
    check(kl.get("NAMN") == "Göteborgs kommun" and kl.get("MANDAT_VALKRETS") == "44", "2006 landstingsvalkrets 1401 Goteborgs kommun, 44 fasta mandat")
    for p, (m, u) in giltiga(kl).items():
        add_vkm(2006, "rf", "Göteborgs kommun", p, m - u, u, m, "dl2006/slutresultat_00L.xml KRETS_LANDSTING 1401 GILTIGA MANDAT, VARAV_UTJÄMNING")
    for p, (m, u) in giltiga(kl, "MANDAT_FGVAL", "INGEN").items():
        add_vkm(2002, "rf", "Göteborgs kommun", p, "", "", m, "dl2006/slutresultat_00L.xml KRETS_LANDSTING 1401 MANDAT_FGVAL (utjamning okand)")
    vd = kl.find("VALDELTAGANDE")
    add_vkr(2006, "rf", "Göteborgs kommun", 44, int(vd.get("RÖSTBERÄTTIGADE")),
            "mandat: dl2006/slutresultat_00L.xml KRETS_LANDSTING 1401 MANDAT_VALKRETS; rostberattigade valdagen: VALDELTAGANDE RÖSTBERÄTTIGADE")
    add_vkr(2002, "rf", "Göteborgs kommun", "", int(vd.get("RÖSTBERÄTTIGADE_KLARA_VALDISTRIKT_FGVAL")),
            "fasta mandat 2002 saknas i kallorna; rostberattigade valdagen: dl2006/slutresultat_00L.xml KRETS_LANDSTING 1401 VALDELTAGANDE RÖSTBERÄTTIGADE_KLARA_VALDISTRIKT_FGVAL")

    # Goteborgs kommunvalkretsar: 2006 och 2002
    k06x = xml_root(XML_2006 + "/slutresultat_1480K.xml")
    kom = k06x.find(".//KOMMUN")
    check(kom.get("KOD") == "1480" and kom.get("MANDAT_VALOMRÅDE") == "81", "2006 kommunfullmaktige Goteborg 81 mandat")
    for p, (m, u) in giltiga(kom).items():
        add_vkm(2006, "kf", "Göteborg", p, m, 0, m, "dl2006/slutresultat_1480K.xml KOMMUN 1480 GILTIGA MANDAT (hela kommunen)")
    for p, (m, u) in giltiga(kom, "MANDAT_FGVAL", "INGEN").items():
        add_vkm(2002, "kf", "Göteborg", p, m, 0, m, "dl2006/slutresultat_1480K.xml KOMMUN 1480 GILTIGA MANDAT_FGVAL (hela kommunen)")
    majorna_2006 = {}
    for kk in k06x.iter("KRETS_KOMMUN"):
        namn = kk.get("NAMN")
        for p, (m, u) in giltiga(kk).items():
            add_vkm(2006, "kf", namn, p, m, 0, m, "dl2006/slutresultat_1480K.xml KRETS_KOMMUN %s GILTIGA MANDAT" % kk.get("KOD"))
        fg = giltiga(kk, "MANDAT_FGVAL", "INGEN")
        for p, (m, u) in fg.items():
            add_vkm(2002, "kf", namn, p, m, 0, m, "dl2006/slutresultat_1480K.xml KRETS_KOMMUN %s GILTIGA MANDAT_FGVAL" % kk.get("KOD"))
        vd = kk.find("VALDELTAGANDE")
        add_vkr(2006, "kf", namn, int(kk.get("MANDAT_VALKRETS")), int(vd.get("RÖSTBERÄTTIGADE")),
                "mandat: dl2006/slutresultat_1480K.xml KRETS_KOMMUN %s MANDAT_VALKRETS; rostberattigade valdagen: VALDELTAGANDE RÖSTBERÄTTIGADE" % kk.get("KOD"))
        add_vkr(2002, "kf", namn, sum(m for m, u in fg.values()), int(vd.get("RÖSTBERÄTTIGADE_KLARA_VALDISTRIKT_FGVAL")),
                "mandat: summa MANDAT_FGVAL per parti i dl2006/slutresultat_1480K.xml KRETS_KOMMUN %s (kommunval utan utjamningsmandat); rostberattigade valdagen: VALDELTAGANDE RÖSTBERÄTTIGADE_KLARA_VALDISTRIKT_FGVAL" % kk.get("KOD"))
        majorna_2006[namn] = sorted((v.get("KOD"), v.get("NAMN")) for v in kk.iter("VALDISTRIKT")
                                    if v.get("KOD")[4:6] in ("84", "85", "59"))
    check(sum(int(kk.get("MANDAT_VALKRETS")) for kk in k06x.iter("KRETS_KOMMUN")) == 81, "2006 kommunvalkretsarnas fasta mandat summerar till 81")

    # ------------------------------------------------ 2010
    mf10 = las_mandatfordelning(F_MF_R_2010, None)
    tot, fasta, utj = summera_riket(mf10)
    check(sum(tot.values()) == 349 and sum(fasta.values()) == 310 and sum(utj.values()) == 39,
          "2010 riksdag: mandatfordelning_per_valkrets_R (1).xls summerar till 349 = 310 fasta + 39 utjamning")
    r10 = xml_root(XML_2010 + "/slutresultat_00R.xml")
    check(r10.get("VALDAG") == "20100919", "2010 00R.xml ar valet 2010-09-19")
    xn = giltiga(r10.find(".//NATION"))
    check(all(xn[p][0] == tot[p] for p in tot), "2010 riksdag riket: xls-summa = XML NATION MANDAT")
    for p in tot:
        add_riket(2010, p, tot[p], "mandatfordelning_per_valkrets_R (1).xls, summa per parti av <parti> och <parti> utj over alla valkretsar")
    g = mf10[(14, 16)]
    check(g["namn"] == "Göteborgs kommun" and g["fasta_antal"] == 17, "2010 riksdagsvalkrets 14/16 = Goteborgs kommun, 17 fasta")
    kr10 = [k for k in r10.iter("KRETS_RIKSDAG") if k.get("KOD") == "1416"][0]
    xg = giltiga(kr10)
    check(all(xg[p] == (f + u, u) for p, (f, u) in g["parti"].items()), "2010 riksdag Goteborg: xls = XML KRETS_RIKSDAG 1416")
    for p, (f, u) in g["parti"].items():
        add_vkm(2010, "rd", "Göteborgs kommun", p, f, u, f + u, "mandatfordelning_per_valkrets_R (1).xls rad LAN 14 VALKRETS 16, kolumner <parti> och <parti> utj")
    vkm_r10 = {r[0]: r for r in xls_rows(F_VKM_R_2010)[5:]}
    check(vkm_r10["Göteborgs kommun"][5] == 17.0, "2010 valkretsmandat_R_2010.xls Goteborgs kommun 17 mandat")
    add_vkr(2010, "rd", "Göteborgs kommun", 17, rostberattigade(kr10),
            "mandat: valkretsmandat_R_2010.xls kolumn Mandat (rostberattigade dar: %d); rostberattigade valdagen: slutresultat__1_/slutresultat_00R.xml KRETS_RIKSDAG 1416 VALDELTAGANDE RÖSTBERÄTTIGADE" % int(vkm_r10["Göteborgs kommun"][1]))

    # landsting 2010: septemberresultatet i XML, omvalet 2011 i mandatfordelning_per_valkrets_L.xls
    l10 = xml_root(XML_2010 + "/slutresultat_00L.xml")
    kl10 = [k for k in l10.iter("KRETS_LANDSTING") if k.get("KOD") == "1401"][0]
    check(kl10.get("MANDAT_VALKRETS") == "44", "2010 landstingsvalkrets 1401, 44 fasta mandat (XML)")
    for p, (m, u) in giltiga(kl10).items():
        add_vkm(2010, "rf", "Göteborgs kommun", p, m - u, u, m, "slutresultat__1_/slutresultat_00L.xml KRETS_LANDSTING 1401 GILTIGA MANDAT, VARAV_UTJÄMNING (rostrakningen efter valet 2010-09-19, upphavd genom omvalet 2011-05-15)")
    vkm_l10 = {(int(r[0]), int(r[3])): r for r in xls_rows(F_VKM_L_2010)[5:] if r[0] != ""}
    check(vkm_l10[(14, 1)][4] == "Göteborgs kommun" and vkm_l10[(14, 1)][9] == 44.0, "2010 valkretsmandat_L_2010.xls Goteborgs kommun 44 mandat")
    add_vkr(2010, "rf", "Göteborgs kommun", 44, rostberattigade(kl10),
            "mandat: valkretsmandat_L_2010.xls kolumn Mandat (rostberattigade dar: %d); rostberattigade valdagen: slutresultat__1_/slutresultat_00L.xml KRETS_LANDSTING 1401 VALDELTAGANDE RÖSTBERÄTTIGADE" % int(vkm_l10[(14, 1)][5]))
    mfl = las_mandatfordelning(F_MF_L_2011, None)
    tot_l, fasta_l, utj_l = summera_riket({k: v for k, v in mfl.items() if k[0] == 14})
    l14 = xml_root(XML_2014 + "/slutresultat_00L.xml")
    lan14 = [k for k in l14.iter("LÄN") if k.get("KOD") == "14"][0]
    fg14 = giltiga(lan14, "MANDAT_FGVAL", "INGEN")
    lan10 = [k for k in l10.iter("LÄN") if k.get("KOD") == "14"][0]
    x10 = giltiga(lan10)
    check(all(fg14.get(p, (0,))[0] == tot_l[p] for p in tot_l) and sum(tot_l.values()) == 149,
          "mandatfordelning_per_valkrets_L.xls lan 14 = MANDAT_FGVAL i 2014 ars 00L.xml (omvalet 2011-05-15), inte septemberresultatet 2010")
    check(not all(x10.get(p, (0,))[0] == tot_l[p] for p in tot_l), "mandatfordelning_per_valkrets_L.xls skiljer sig fran 2010 ars 00L.xml (SPVG saknas, SD fler)")
    g = mfl[(14, 1)]
    check(g["namn"] == "Göteborgs kommun" and g["fasta_antal"] == 44, "omval 2011 landstingsvalkrets 14/1 = Goteborgs kommun, 44 fasta")
    kl14 = [k for k in l14.iter("KRETS_LANDSTING") if k.get("KOD") == "1401"][0]
    fg = giltiga(kl14, "MANDAT_FGVAL", "INGEN")
    check(all(fg[p][0] == f + u for p, (f, u) in g["parti"].items()), "omval 2011 Goteborg: L-xls = MANDAT_FGVAL i 2014 ars 00L.xml KRETS_LANDSTING 1401")
    for p, (f, u) in g["parti"].items():
        add_vkm(2011, "rf", "Göteborgs kommun", p, f, u, f + u, "mandatfordelning_per_valkrets_L.xls rad LAN 14 VALKRETS 1 (omvalet till regionfullmaktige 2011-05-15)")
    add_vkr(2011, "rf", "Göteborgs kommun", 44, "", "mandat: mandatfordelning_per_valkrets_L.xls ANT_MAND_FAST (omvalet 2011-05-15); rostberattigade saknas i kallorna")

    # kommun 2010: mandatfordelning_per_valkrets_K.xls, kontroll mot 1480K.xml
    mfk = las_mandatfordelning_K(F_MF_K_2010)
    k10x = xml_root(XML_2010 + "/slutresultat_1480K.xml")
    kom10 = k10x.find(".//KOMMUN")
    check(kom10.get("MANDAT_VALOMRÅDE") == "81", "2010 kommunfullmaktige Goteborg 81 mandat (XML)")
    xk = {kk.get("KOD"): kk for kk in k10x.iter("KRETS_KOMMUN")}
    vkm_k10 = {(int(r[2]), int(r[4])): r for r in xls_rows(F_VKM_K_2010)[5:] if r[0] != ""}
    fg_k14 = {kk.get("KOD"): giltiga(kk, "MANDAT_FGVAL", "INGEN") for kk in xml_root(XML_2014 + "/slutresultat_1480K.xml").iter("KRETS_KOMMUN")}
    for (lan, kom_, vk), d in mfk.items():
        if (lan, kom_) != (14, 80):
            continue
        kod = "1480%02d" % vk
        xm = giltiga(xk[kod])
        check(all(xm[p][0] == m for p, m in d["parti"].items()) and sum(d["parti"].values()) == d["antal"] == int(xk[kod].get("MANDAT_VALKRETS")),
              "mandatfordelning_per_valkrets_K.xls %s = 2010 ars 1480K.xml KRETS_KOMMUN %s (alltsa 2010)" % (d["namn"], kod))
        check(all(fg_k14[kod][p][0] == m for p, m in d["parti"].items()), "  ... och = MANDAT_FGVAL i 2014 ars 1480K.xml")
        for p, m in d["parti"].items():
            add_vkm(2010, "kf", d["namn"], p, m, 0, m, "mandatfordelning_per_valkrets_K.xls rad LAN 14 KOM 80 VALKRETS %d, kolumn %s" % (vk, p))
        r = vkm_k10[(1480, vk)]
        check(r[5] == d["namn"] and int(r[10]) == d["antal"], "2010 valkretsmandat_K_2010.xls %s %d mandat" % (d["namn"], d["antal"]))
        add_vkr(2010, "kf", d["namn"], d["antal"], rostberattigade(xk[kod]),
                "mandat: valkretsmandat_K_2010.xls kolumn Mandat (rostberattigade dar: %d); rostberattigade valdagen: slutresultat__1_/slutresultat_1480K.xml KRETS_KOMMUN %s VALDELTAGANDE RÖSTBERÄTTIGADE" % (int(r[6]), kod))
    for p, (m, u) in giltiga(kom10).items():
        add_vkm(2010, "kf", "Göteborg", p, m, 0, m, "slutresultat__1_/slutresultat_1480K.xml KOMMUN 1480 GILTIGA MANDAT (hela kommunen)")
    majorna_2010 = {xk[k].get("KOD"): sorted(v.get("KOD") for v in xk[k].iter("VALDISTRIKT") if v.get("KOD")[4:6] == "09") for k in xk}

    # ------------------------------------------------ 2014
    mf14 = las_mandatfordelning(F_MF_R_2014, None)
    tot, fasta, utj = summera_riket(mf14)
    check(sum(tot.values()) == 349 and sum(fasta.values()) == 310 and sum(utj.values()) == 39,
          "2014 riksdag: mandatfordelning_per_valkrets_R.xls summerar till 349 = 310 fasta + 39 utjamning")
    r14 = xml_root(XML_2014 + "/slutresultat_00R.xml")
    check(r14.get("VALDAG") == "20140914", "2014 00R.xml ar valet 2014-09-14")
    xn = giltiga(r14.find(".//NATION"))
    check(all(xn[p][0] == tot[p] for p in tot), "2014 riksdag riket: xls-summa = XML NATION MANDAT")
    for p in tot:
        add_riket(2014, p, tot[p], "mandatfordelning_per_valkrets_R.xls (Valmyndigheten 2014-10-31), summa per parti av <parti> och <parti> utj over alla valkretsar")
    g = mf14[(14, 16)]
    kr14 = [k for k in r14.iter("KRETS_RIKSDAG") if k.get("KOD") == "1416"][0]
    xg = giltiga(kr14)
    check(g["fasta_antal"] == 17 and all(xg[p] == (f + u, u) for p, (f, u) in g["parti"].items()), "2014 riksdag Goteborg: xls = XML KRETS_RIKSDAG 1416, 17 fasta")
    for p, (f, u) in g["parti"].items():
        add_vkm(2014, "rd", "Göteborgs kommun", p, f, u, f + u, "mandatfordelning_per_valkrets_R.xls rad LAN 14 VALKRETS 16, kolumner <parti> och <parti> utj")
    vkm_r14 = {r[0]: r for r in xls_rows(F_VKM_R_2014)[5:]}
    check(vkm_r14["Göteborgs kommun"][5] == 17.0, "2014 Valkretsmandat riksdag 2014.xls Goteborgs kommun 17 mandat")
    add_vkr(2014, "rd", "Göteborgs kommun", 17, rostberattigade(kr14),
            "mandat: Valkretsmandat riksdag 2014.xls kolumn MANDAT (rostberattigade dar: %d); rostberattigade valdagen: slutresultat/slutresultat_00R.xml KRETS_RIKSDAG 1416 VALDELTAGANDE RÖSTBERÄTTIGADE" % int(vkm_r14["Göteborgs kommun"][1]))
    for p, (m, u) in giltiga(kl14).items():
        add_vkm(2014, "rf", "Göteborgs kommun", p, m - u, u, m, "slutresultat/slutresultat_00L.xml KRETS_LANDSTING 1401 GILTIGA MANDAT, VARAV_UTJÄMNING")
    vkm_l14 = {(r[0], int(r[3])): r for r in xls_rows(F_VKM_L_2014)[5:] if r[0] != ""}
    check(vkm_l14[("14", 1)][4] == "Göteborgs kommun" and vkm_l14[("14", 1)][9] == 44.0 and kl14.get("MANDAT_VALKRETS") == "44", "2014 Valkretsmandat landsting 2014.xls Goteborgs kommun 44 mandat = XML")
    add_vkr(2014, "rf", "Göteborgs kommun", 44, rostberattigade(kl14),
            "mandat: Valkretsmandat landsting 2014.xls kolumn MANDAT (rostberattigade dar: %d); rostberattigade valdagen: slutresultat/slutresultat_00L.xml KRETS_LANDSTING 1401 VALDELTAGANDE RÖSTBERÄTTIGADE" % int(vkm_l14[("14", 1)][5]))
    k14x = xml_root(XML_2014 + "/slutresultat_1480K.xml")
    kom14 = k14x.find(".//KOMMUN")
    check(kom14.get("MANDAT_VALOMRÅDE") == "81", "2014 kommunfullmaktige Goteborg 81 mandat (XML)")
    for p, (m, u) in giltiga(kom14).items():
        add_vkm(2014, "kf", "Göteborg", p, m, 0, m, "slutresultat/slutresultat_1480K.xml KOMMUN 1480 GILTIGA MANDAT (hela kommunen)")
    majorna_2014 = {}
    for kk in k14x.iter("KRETS_KOMMUN"):
        for p, (m, u) in giltiga(kk).items():
            add_vkm(2014, "kf", kk.get("NAMN"), p, m, 0, m, "slutresultat/slutresultat_1480K.xml KRETS_KOMMUN %s GILTIGA MANDAT" % kk.get("KOD"))
        add_vkr(2014, "kf", kk.get("NAMN"), int(kk.get("MANDAT_VALKRETS")), rostberattigade(kk),
                "mandat: slutresultat/slutresultat_1480K.xml KRETS_KOMMUN %s MANDAT_VALKRETS (ingen valkretsmandat-fil for kommun 2014 bland kallorna); rostberattigade valdagen: VALDELTAGANDE RÖSTBERÄTTIGADE" % kk.get("KOD"))
        majorna_2014[kk.get("NAMN")] = sorted((v.get("KOD"), v.get("NAMN")) for v in kk.iter("VALDISTRIKT") if v.get("KOD")[4:6] == "10")
    check(sum(int(kk.get("MANDAT_VALKRETS")) for kk in k14x.iter("KRETS_KOMMUN")) == 81, "2014 kommunvalkretsarnas fasta mandat summerar till 81")

    # ------------------------------------------------ 2018
    rows = [r for r in xlsx_rows(F_MANDAT_2018, "Mandatfördelning")[1:] if r[0]]
    tot = collections.Counter(); fasta = collections.Counter(); utj = collections.Counter()
    for r in rows:
        if r[0] == "R":
            tot[r[5]] += r[11] or 0; fasta[r[5]] += r[9] or 0; utj[r[5]] += r[10] or 0
    check(sum(tot.values()) == 349 and sum(fasta.values()) == 310 and sum(utj.values()) == 39,
          "2018 riksdag: 2018_mandat.xlsx summerar till 349 = 310 fasta + 39 utjamning")
    for p in tot:
        add_riket(2018, p, tot[p], "2018_mandat.xlsx flik Mandatfördelning, VALTYP R, summa SUMMA MANDAT per PARTIFÖRKORTNING over alla valkretsar")
    rb18 = {}
    for val, f, sheet in (("rd", F_2018_R, "R antal"), ("rf", F_2018_L, "L antal"), ("kf", F_2018_K, "K antal")):
        rr = xlsx_rows(f, sheet)
        hdr = rr[0]
        irb = hdr.index("RÖSTBERÄTTIGADE")
        gb = [r for r in rr[1:] if r[0] == 14 and r[1] == 80]
        kretsar = collections.Counter((r[2], r[6]) for r in gb)
        rb18[val] = (sum(r[irb] or 0 for r in gb), len(gb), dict(kretsar))
    check(rb18["kf"][2] == {(None, None): rb18["kf"][1]} and rb18["rd"][2] == {(16, "Göteborgs kommun"): rb18["rd"][1]} and rb18["rf"][2] == {(1, "Göteborgs kommun"): rb18["rf"][1]},
          "2018 per_valdistrikt: Goteborgs %d rader (350 distrikt + Uppsamlingsdistrikt) har tom kommunvalkrets i K-filen, riksdagsvalkrets 16 och landstingsvalkrets 1" % rb18["kf"][1])
    for r in rows:
        if r[0] == "R" and r[4] == "Göteborgs kommun":
            add_vkm(2018, "rd", "Göteborgs kommun", r[5], r[9] or 0, r[10] or 0, r[11], "2018_mandat.xlsx flik Mandatfördelning, VALTYP R VALKRETSKOD 16, kolumner FASTA MANDAT, UTJÄMNINGSMANDAT, SUMMA MANDAT")
        if r[0] == "L" and r[2] == "Västra Götalands läns landsting" and r[4] == "Göteborgs kommun":
            add_vkm(2018, "rf", "Göteborgs kommun", r[5], r[9] or 0, r[10] or 0, r[11], "2018_mandat.xlsx flik Mandatfördelning, VALTYP L VALOMRÅDE Västra Götalands läns landsting VALKRETSKOD 1")
        if r[0] == "K" and r[2] == "Göteborg":
            check(r[3] == 0 and (r[10] or 0) == 0 and r[9] == r[11], "2018 kommun Goteborg %s: valkretskod 0, inga utjamningsmandat" % r[5])
            add_vkm(2018, "kf", "Göteborg", r[5], r[9], 0, r[11], "2018_mandat.xlsx flik Mandatfördelning, VALTYP K VALOMRÅDE Göteborg VALKRETSKOD 0")
    f_rd = sum(r[9] or 0 for r in rows if r[0] == "R" and r[4] == "Göteborgs kommun")
    f_rf = sum(r[9] or 0 for r in rows if r[0] == "L" and r[2] == "Västra Götalands läns landsting" and r[4] == "Göteborgs kommun")
    f_kf = sum(r[9] or 0 for r in rows if r[0] == "K" and r[2] == "Göteborg")
    check(f_rd == 17 and f_kf == 81, "2018 Goteborg fasta mandat: riksdag 17, kommun 81 (summa FASTA MANDAT)")
    add_vkr(2018, "rd", "Göteborgs kommun", f_rd, rb18["rd"][0], "mandat: summa FASTA MANDAT i 2018_mandat.xlsx VALTYP R VALKRETSKOD 16; rostberattigade valdagen: summa RÖSTBERÄTTIGADE i 2018_R_per_valdistrikt.xlsx flik R antal, LÄNSKOD 14 KOMMUNKOD 80")
    add_vkr(2018, "rf", "Göteborgs kommun", f_rf, rb18["rf"][0], "mandat: summa FASTA MANDAT i 2018_mandat.xlsx VALTYP L VALKRETSKOD 1 (Västra Götalands läns landsting); rostberattigade valdagen: summa RÖSTBERÄTTIGADE i 2018_L_per_valdistrikt.xlsx flik L antal")
    add_vkr(2018, "kf", "Göteborg", f_kf, rb18["kf"][0], "mandat: summa FASTA MANDAT i 2018_mandat.xlsx VALTYP K Göteborg VALKRETSKOD 0 (hela kommunen en valkrets); rostberattigade valdagen: summa RÖSTBERÄTTIGADE i 2018_K_per_valdistrikt.xlsx flik K antal")

    # ------------------------------------------------ 2022
    NAMN2022 = {"Arbetarepartiet-Socialdemokraterna": "S", "Centerpartiet": "C", "Kristdemokraterna": "KD",
                "Liberalerna (tidigare Folkpartiet)": "L", "Miljöpartiet de gröna": "MP", "Moderaterna": "M",
                "Sverigedemokraterna": "SD", "Vänsterpartiet": "V", "Demokraterna": "D", "Feministiskt initiativ": "FI"}
    rows = xlsx_rows(F_JAMF_2022, "Riksdag")
    check(rows[0][1] == 2018 and rows[0][2] == 2022 and rows[1][0] == "Sverige" and rows[1][2] == 349, "2022 Mandatfordelning-jamforelser: flik Riksdag, kolumn 2022, Sverige 349")
    m22 = collections.OrderedDict((NAMN2022[r[0]], r[2] or 0) for r in rows[2:] if r[0])
    m18 = collections.OrderedDict((NAMN2022[r[0]], r[1] or 0) for r in rows[2:] if r[0])
    check(sum(m22.values()) == 349 and all(m18[p] == tot[p] for p in tot), "2022-filen: summa 349 och kolumn 2018 = 2018_mandat.xlsx")
    vd22 = json.load(open(F_VALDATA_2022))["mandat"]["riksdag_verklig"]
    check(all(vd22[p] == m22[p] for p in m22), "2022 riksdag = data/valdata_2022.json mandat.riksdag_verklig")
    jr = xlsx_rows(F_JAMF_RD_2022, "Riket")
    j22 = {NAMN2022[r[11]]: r[12] for r in jr[1:] if r[11] in NAMN2022}
    check(all(j22[p] == m22[p] for p in m22), "2022 riksdag = slutligt-valresultat-riksdagen-jamforande-statistik-2018-2022.xlsx flik Riket, Mandat 2022")
    for p, m in m22.items():
        add_riket(2022, p, m, "Mandatfordelning-jamforelser-mellan-2018-och-2022.xlsx flik Riksdag kolumn 2022")
    # riksdagsvalkretsen Goteborg 2022 ur jamforande-filen, flik Valkrets (bara totalt)
    jv = xlsx_rows(F_JAMF_RD_2022, "Valkrets")
    g22 = collections.OrderedDict()
    aktiv = None
    for r in jv[1:]:
        if r[1]:
            aktiv = r[1]
        if aktiv == "Göteborgs kommun" and r[11] in NAMN2022:
            g22[NAMN2022[r[11]]] = (r[12] or 0, r[14] or 0)
        if aktiv == "Göteborgs kommun" and r[11] == "Summa":
            summa22, summa18 = r[12], r[14]
    check(sum(v[0] for v in g22.values()) == summa22 and sum(v[1] for v in g22.values()) == summa18,
          "2022 riksdag Goteborgs kommun i jamforande-filen: partierna summerar till Summa-raden (2022: %d, 2018: %d)" % (summa22, summa18))
    x18 = {r[5]: r[11] for r in xlsx_rows(F_MANDAT_2018, "Mandatfördelning")[1:] if r[0] == "R" and r[4] == "Göteborgs kommun"}
    avvik = {p: (x18.get(p, 0), g22[p][1]) for p in g22 if x18.get(p, 0) != g22[p][1]}
    if avvik:
        print("AVVIKELSE: jamforande-filens kolumn Mandat 2018 for Goteborgs kommun skiljer sig fran 2018_mandat.xlsx (parti: 2018_mandat.xlsx, jamforande):", avvik)
    for p, (m, m18_) in g22.items():
        add_vkm(2022, "rd", "Göteborgs kommun", p, "", "", m, "slutligt-valresultat-riksdagen-jamforande-statistik-2018-2022.xlsx flik Valkrets, block Göteborgs kommun, kolumn Mandat 2022 (fasta/utjamning saknas; filens kolumn Mandat 2018 avviker fran 2018_mandat.xlsx for SD)")
    # kommun 2022
    rows = xlsx_rows(F_JAMF_2022, "Kommun")
    i0 = [i for i, r in enumerate(rows) if r[0] == "Göteborg"][0]
    check(rows[i0][2] == 81, "2022 Mandatfordelning-jamforelser flik Kommun: Göteborg 81 mandat")
    k22 = collections.OrderedDict()
    for r in rows[i0 + 1:]:
        if r[0] in NAMN2022:
            k22[NAMN2022[r[0]]] = (r[1] or 0, r[2] or 0)
        else:
            break
    check(sum(v[1] for v in k22.values()) == 81 and sum(v[0] for v in k22.values()) == 81, "2022 kommun Goteborg summerar till 81 bada aren")
    k18x = {r[5]: r[11] for r in xlsx_rows(F_MANDAT_2018, "Mandatfördelning")[1:] if r[0] == "K" and r[2] == "Göteborg"}
    check(all(k22[norm(p)][0] == m for p, m in k18x.items()), "2022-filens kolumn 2018 for Goteborg = 2018_mandat.xlsx")
    for p, (m18_, m) in k22.items():
        add_vkm(2022, "kf", "Göteborg", p, m, 0, m, "Mandatfordelning-jamforelser-mellan-2018-och-2022.xlsx flik Kommun, block Göteborg, kolumn 2022")
    rb22 = {}
    for val, f in F_RAW_2022.items():
        wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
        ws = [w for w in wb.worksheets if w.title.startswith("roster")][0]
        rr = list(ws.iter_rows(values_only=True))
        hdr = [str(h).strip() for h in rr[0]]
        iv, ivn, idi, irb = hdr.index("Valkretskod"), hdr.index("Valkretsnamn"), hdr.index("Valdistriktskod"), hdr.index("Röstberättigade")
        per = collections.defaultdict(dict)
        for r in rr[1:]:
            if r[4] == "Göteborg":
                per[(str(r[iv]).strip(), r[ivn])][str(r[idi]).strip()] = r[irb] or 0
        check(len(per) == 1, "2022 %s: Goteborgs distrikt ligger i en enda valkrets %s" % (val, list(per)[0]))
        (kod, namn), d = list(per.items())[0]
        rb22[val] = (kod, namn, sum(d.values()), len(d))
    add_vkr(2022, "rd", "Göteborgs kommun", "", rb22["rd"][2], "fasta mandat 2022 saknas i kallorna (jamforande-filen ger bara summa mandat); rostberattigade valdagen: summa Röstberättigade per valdistrikt i %s flik roster_RD, Kommun Göteborg, Valkretskod %s" % (os.path.basename(F_RAW_2022["rd"]), rb22["rd"][0]))
    add_vkr(2022, "rf", "Göteborgs kommun", "", rb22["rf"][2], "fasta mandat 2022 saknas i kallorna; rostberattigade valdagen: summa Röstberättigade per valdistrikt i %s flik roster_RF, Kommun Göteborg, Valkretskod %s" % (os.path.basename(F_RAW_2022["rf"]), rb22["rf"][0]))
    add_vkr(2022, "kf", "Göteborg", 81, rb22["kf"][2], "mandat: Mandatfordelning-jamforelser-mellan-2018-och-2022.xlsx flik Kommun Göteborg (hela kommunen en valkrets, Valkretskod %s); rostberattigade valdagen: summa Röstberättigade per valdistrikt i %s flik roster_KF" % (rb22["kf"][0], os.path.basename(F_RAW_2022["kf"])))

    # ------------------------------------------------ kontroller pa utdata och skrivning
    for ar in (2002, 2006, 2010, 2014, 2018, 2022):
        s = sum(r[2] for r in riket if r[0] == ar)
        check(s == 349, "mandat_riksdag_riket %d summerar till 349" % ar)
    for ar, val, vk, vantat in ((2006, "rd", "Göteborgs kommun", 18), (2010, "rd", "Göteborgs kommun", 18), (2014, "rd", "Göteborgs kommun", 17),
                                (2018, "rd", "Göteborgs kommun", 19), (2002, "rd", "Göteborgs kommun", 18)):
        s = sum(r[6] for r in vkm if r[0] == ar and r[1] == val and r[2] == vk)
        check(s == vantat, "mandat_valkrets_goteborg %d %s %s totalt %d" % (ar, val, vk, s))
    for ar in (2002, 2006, 2010, 2014, 2018, 2022):
        s = sum(r[6] for r in vkm if r[0] == ar and r[1] == "kf" and r[2] == "Göteborg")
        check(s == 81, "mandat_valkrets_goteborg %d kf Göteborg (hela kommunen) summerar till 81" % ar)
    for ar in (2002, 2006, 2010, 2014):
        s = sum(r[6] for r in vkm if r[0] == ar and r[1] == "kf" and r[2] != "Göteborg")
        check(s == 81, "mandat_valkrets_goteborg %d kf kretsarna summerar till 81" % ar)
        s = sum(r[3] for r in vkr if r[0] == ar and r[1] == "kf")
        check(s == 81, "valkretsar_goteborg %d kf antal_mandat summerar till 81" % ar)

    def skriv(namn, hdr, rader):
        p = os.path.join(UT, namn)
        with open(p, "w", encoding="utf-8", newline="") as fh:
            w = csv.writer(fh, delimiter=";", lineterminator="\n")
            w.writerow(hdr)
            for r in rader:
                # semikolon i kalla-texten byts mot komma sa att filen aldrig behover citattecken
                w.writerow(["" if v is None else (v.replace(";", ",") if isinstance(v, str) else v) for v in r])
        print("skrev", p, len(rader), "rader")

    ordning = {"rd": 0, "rf": 1, "kf": 2}
    riket.sort(key=lambda r: (r[0], -r[2], r[1]))
    vkm.sort(key=lambda r: (r[0], ordning[r[1]], r[2] != "Göteborg", r[2], -(r[6] or 0), r[3]))
    vkr.sort(key=lambda r: (r[0], ordning[r[1]], r[2]))
    skriv("mandat_riksdag_riket.csv", ["ar", "parti", "mandat", "kalla"], riket)
    skriv("mandat_valkrets_goteborg.csv", ["ar", "val", "valkrets", "parti", "fasta", "utjamning", "totalt", "kalla"], vkm)
    skriv("valkretsar_goteborg.csv", ["ar", "val", "valkrets", "antal_mandat", "rostberattigade", "kalla"], vkr)

    print("\nMajornaomradets distrikt per kommunvalkrets:")
    for namn, v in majorna_2006.items():
        print(" 2006", namn, len(v), [x[1] for x in v])
    for kod, v in majorna_2010.items():
        print(" 2010", kod, len(v), v)
    for namn, v in majorna_2014.items():
        print(" 2014", namn, len(v), [x[1] for x in v])
    print(" 2018 kommunvalkrets 0 Göteborg:", rb18["kf"][1], "distrikt, rostberattigade", rb18["kf"][0])
    print(" 2022 valkretsar:", rb22)


if __name__ == "__main__":
    main()
