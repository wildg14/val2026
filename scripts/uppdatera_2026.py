#!/usr/bin/env python3
"""Tar in valresultat för ett nytt val (2026) och skriver data/valdata_<år>.json/.js samt swing mot 2022.

Valnatten, JSON-vägen (rekommenderad):
    .venv/bin/python scripts/uppdatera_2026.py --hamta --status preliminar --valnatt
hämtar Valmyndighetens resultatfiler (scripts/hamta_2026.py, --tillfalle p eller s, --genrep för
simuleringarna, --utan-signatur skickas vidare till hamta_2026.py som reservläge om val.se bekräftar
problem med certifikatet), packar upp dem och läser rd/, rf/ och kf/ direkt. Har mappen redan hämtats:
    .venv/bin/python scripts/uppdatera_2026.py --valnatt-mapp data/valnatt/senaste --status preliminar

Slutlig-vägen 2022 (en, två eller tre filer):
    .venv/bin/python scripts/uppdatera_2026.py --rd RD-FIL.xlsx [--rf RF-FIL.xlsx] [--kf KF-FIL.xlsx] \\
        --status preliminar [--tid 2026-09-13T21:30:00]

Reservväg om Valmyndighetens filformat ändrats: fyll i en CSV för hand (val;kod;parti;roster,
med raderna giltiga, rostande och rostberattigade per distrikt) och kör
    .venv/bin/python scripts/uppdatera_2026.py --csv valnatt.csv --status preliminar
Mall: --skriv-mall valnatt.csv

Generalrepetition: --repetera kör parsern på 2022 års råfiler och jämför med data/valdata_2022.json.

Jämförbarhet mot 2022 (jamforbar_mot_bas, grans_andrad per distrikt): båda källorna läses, JSON-
filernas statusJamforelse och Valmyndighetens xlsx via --jamforelsefil (standard: filen under
Historiska dokument/dl_webb/2026/). Säger någon av dem "ej jämförbart" gäller det, med en varning
vid konflikt mellan källorna. Saknas båda antas alla distrikt jämförbara, med en varning.

Skriptet filtrerar på de 23 distriktskoderna, larmar (VARNING:) om distrikt saknas, bytt namn,
om Göteborgs distriktsindelning ser ändrad ut eller om okända partier får röster, och avbryter (FEL:)
om filformatet inte känns igen eller om summorna inte går ihop. En tom import (inget Majornadistrikt
räknat i något val än, som kvällens första körningar) skriver filerna med 0 räknade i stället för att
stoppas. --tvinga skriver ändå vid färre räknade distrikt än den befintliga filen; kommandot efter ett
sådant stopp är --valnatt-mapp data/valnatt/senaste --tvinga.
"""
import argparse
import csv
import glob
import json
import sys
from pathlib import Path

import openpyxl

ROT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROT))

from scripts import hamta_2026, schema, valnatt  # noqa: E402
from scripts.mandat import jamkade_uddatal  # noqa: E402
from scripts.valmyndigheten import (MAJORNA_KODER, NYCKELPARTIER, OVRIGA, VAL, FormatFel, SummaFel,  # noqa: E402
                                    aggregat_andelar, las_rafil, till_distrikt)

RAFIL_MONSTER = {"rd": "*oster-per-distrikt*riksdagsval*.xlsx", "rf": "*oster-per-distrikt*regionval*.xlsx",
                 "kf": "*oster-per-distrikt*kommunval*.xlsx"}
GOTEBORG_DISTRIKT = {"xlsx_2022": 411, "json_2026": 397}   # 2022: 410 valdistrikt plus uppsamling; 2026: 396 plus uppsamling (genrep-filerna)
JAMFORELSEFIL_STANDARD = ROT / "Historiska dokument" / "dl_webb" / "2026" / "valdistrikt-jamforelser-mellan-2022-och-2026.xlsx"
JAMFORBAR = {}        # kod -> bool, fylls av JSON-vägen eller jämförelsefilen
META_TEST = False     # sätts när någon fil har test: true
METOD = "Räkneexempel: 4 %-spärr och jämkade uddatalsmetoden tillämpade på Majornas riksdagsröster {ar}."
VARNINGAR = []
NAMNVARNADE = set()   # koder som redan fått en namnvarning i den här körningen (en gång per distrikt, inte en per val)


def varning(text):
    VARNINGAR.append(text)
    print(f"VARNING: {text}", file=sys.stderr)


def fel(text):
    print(f"FEL: {text}", file=sys.stderr)
    sys.exit(1)


def tomma_distrikt(bas):
    namn = {d["kod"]: d["namn"] for d in bas["distrikt"]} if bas else {}
    return {kod: {"kod": kod, "namn": namn.get(kod, kod), "raknat": False, "rd": {}, "rf": {}, "kf": {},
                  "giltiga": {}, "rostande": {}, "rostberattigade": {}} for kod in MAJORNA_KODER}


def fyll(distrikt, kod, val, post):
    d = distrikt[kod]
    d[val] = post["roster"]
    d["giltiga"][val] = post["giltiga"]
    d["rostande"][val] = post["rostande"]
    d["rostberattigade"][val] = post["rostberattigade"]
    d["raknat"] = True
    if post.get("namn"):
        d["namn"] = post["namn"]


def _lagg_till_okant(agg, etikett, n, andel):
    """Ackumulerar ett okänt parti över distrikten, för en varning per parti i stället för per distrikt."""
    post = agg.setdefault(etikett, {"roster": 0, "distrikt": 0, "over_troskel": False})
    post["roster"] += n
    post["distrikt"] += 1
    if andel >= 0.005:
        post["over_troskel"] = True


def _varna_okanda(val, agg):
    for etikett, info in agg.items():
        if info["over_troskel"]:
            varning(f"{val}: okänt parti '{etikett}' med {info['roster']} röster i {info['distrikt']} distrikt, läggs i {OVRIGA}")


def las_rafiler(filer, distrikt, bas):
    """filer: {val: sökväg}. Fyller distrikt och returnerar jämförelseaggregat {omrade: {val: ...}}."""
    jamforelser = {"goteborg": {}, "riket": {}}
    for val, path in filer.items():
        print(f"Läser {val}: {path}")
        ra = las_rafil(path, MAJORNA_KODER)
        if ra["val"] != val:
            fel(f"{path}: bladet gäller {ra['val'].upper()}, men filen angavs som --{val}")
        gbg = ra["aggregat"]["goteborg"]["antal_distrikt"]
        if gbg != GOTEBORG_DISTRIKT["xlsx_2022"]:
            varning(f"{val}: Göteborg har {gbg} valdistrikt i filen (2022: {GOTEBORG_DISTRIKT['xlsx_2022']}). "
                    "Distriktsindelningen kan ha ändrats, kontrollera att de 23 koderna fortfarande täcker Majorna.")
        okanda_agg = {}
        for kod in MAJORNA_KODER:
            post = ra["distrikt"].get(kod)
            if post is None:
                varning(f"{val}: distrikt {kod} ({distrikt[kod]['namn']}) saknas i filen, markeras som oräknat")
                continue
            if bas and post["namn"] != bas_namn(bas, kod) and kod not in NAMNVARNADE:
                varning(f"{val}: {kod} heter '{post['namn']}' i filen men '{bas_namn(bas, kod)}' 2022")
                NAMNVARNADE.add(kod)
            try:
                d = till_distrikt(post, val)
            except SummaFel as ex:
                fel(f"{val}: {ex}")
            for etikett, n in d["okanda"].items():
                _lagg_till_okant(okanda_agg, etikett, n, n / d["giltiga"] if d["giltiga"] else 0)
            fyll(distrikt, kod, val, d)
        _varna_okanda(val, okanda_agg)
        try:
            agg = aggregat_andelar(ra)
        except SummaFel as ex:
            fel(f"{val}: {ex}")
        for omrade in ("goteborg", "riket"):
            if omrade in agg:
                jamforelser[omrade][val] = agg[omrade]
    return jamforelser


def bas_namn(bas, kod):
    return next((d["namn"] for d in bas["distrikt"] if d["kod"] == kod), kod)


SUMMANYCKLAR = ("giltiga", "rostande", "rostberattigade")


def las_csv(path, distrikt):
    try:
        text = Path(path).read_text("utf-8-sig")
    except UnicodeDecodeError as ex:
        fel(f"{path}: kunde inte läsas: {ex}. Spara som CSV UTF-8 i Excel.")
    except OSError as ex:
        fel(f"{path}: kunde inte läsas: {ex}")
    forsta = text.splitlines()[0] if text.strip() else ""
    avgransare = ";" if forsta.count(";") >= forsta.count(",") else ","
    lasare = csv.DictReader(text.splitlines(), delimiter=avgransare)
    krav = {"val", "kod", "parti", "roster"}
    if not lasare.fieldnames or not krav <= {(k or "").strip().lower() for k in lasare.fieldnames if k}:
        fel(f"{path}: CSV-filen måste ha kolumnerna val;kod;parti;roster")
    tillatna_per_val = {v: {q.upper(): q for q in NYCKELPARTIER[v] + [OVRIGA]} for v in VAL}
    poster = {}
    for r in lasare:
        i = lasare.line_num
        ra = r
        r = {(k or "").strip().lower(): (v or "").strip() for k, v in r.items() if k is not None}
        if r.get("val", "").startswith("#") or not any(r.values()):
            continue   # kommentarrad eller tom rad ur mallen, före kolumnvakten så att en kommentarrad
                       # med för många semikolon inte avvisas som "fler kolumner"
        if None in ra:
            fel(f"{path} rad {i}: fler kolumner än rubriken (ett semikolon för mycket?)")
        if not r.get("val") or not r.get("kod"):
            if r.get("roster"):
                fel(f"{path} rad {i}: raden har ett tal men saknar val eller kod")
            continue   # tom rad ur mallen
        val, kod, parti = r["val"].lower(), r["kod"], r["parti"]
        if val not in VAL:
            fel(f"{path} rad {i}: okänt val '{val}' (rd, rf eller kf)")
        if kod not in MAJORNA_KODER:
            fel(f"{path} rad {i}: koden {kod} är inte ett av Majornas 23 distrikt")
        if r["roster"] == "":
            continue
        try:
            n = int("".join(r["roster"].split()))
        except ValueError:
            fel(f"{path} rad {i}: '{r['roster']}' är inte ett heltal")
        if n < 0:
            fel(f"{path} rad {i}: negativt tal {n}")
        p = poster.setdefault((val, kod), {"roster": {}, "giltiga": None, "rostande": None, "rostberattigade": 0})
        if parti.lower() in SUMMANYCKLAR:
            p[parti.lower()] = n
            continue
        tillatna = tillatna_per_val[val]
        if parti.upper() not in tillatna:
            fel(f"{path} rad {i}: okänt parti '{parti}' för {val} (tillåtna: {', '.join(NYCKELPARTIER[val] + [OVRIGA])})")
        kanon = tillatna[parti.upper()]
        p["roster"][kanon] = p["roster"].get(kanon, 0) + n
    for (val, kod), p in sorted(poster.items()):
        if p["giltiga"] is None:
            fel(f"{path}: {val} {kod} saknar raden 'giltiga'")
        if sum(p["roster"].values()) != p["giltiga"]:
            fel(f"{path}: {val} {kod}: partiröster {sum(p['roster'].values())} != giltiga {p['giltiga']}")
        if p["rostande"] is None:
            varning(f"{val} {kod}: 'rostande' saknas, sätts lika med giltiga. Valdeltagandet visas då cirka en "
                    "procentenhet för lågt eftersom ogiltiga röster inte räknas med.")
            p["rostande"] = p["giltiga"]
        if p["rostande"] < p["giltiga"]:
            fel(f"{path}: {val} {kod}: röstande {p['rostande']} < giltiga {p['giltiga']}")
        if not p["rostberattigade"]:
            befintlig = distrikt[kod]["rostberattigade"].get(val)
            if befintlig:
                varning(f"{val} {kod}: 'rostberattigade' saknas i CSV-filen, behåller {befintlig} från den andra källan")
                p["rostberattigade"] = befintlig
        if p["rostberattigade"] and p["rostande"] > p["rostberattigade"]:
            fel(f"{path}: {val} {kod}: röstande {p['rostande']} > röstberättigade {p['rostberattigade']}")
        fyll(distrikt, kod, val, p)


def kontrollera_rostberattigade(distrikt):
    """Allt-eller-inget för 'rostberattigade' gäller slutläget över alla källor tillsammans (JSON,
    xlsx och CSV), inte bara en enskild källas egna rader: annars kan en komplettering som saknar
    'rostberattigade' skriva över ett tal en tidigare källa redan satt, eller så kan varje källa se
    komplett ut för sig men ändå ge ett blandat slutläge tillsammans med de andra källornas distrikt.
    Anropas av main på slutläget, efter att alla källor fyllt distrikt, före bygg."""
    for val in VAL:
        raknade = [d for d in distrikt.values() if d.get(val)]
        utan = sorted(d["kod"] for d in raknade if not d["rostberattigade"].get(val))
        med = sorted(d["kod"] for d in raknade if d["rostberattigade"].get(val))
        if med and utan:
            fel(f"{val}: 'rostberattigade' saknas för {', '.join(utan)} men finns för andra distrikt. "
                "Fyll i eller kontrollera källan, annars blir Majornas valdeltagande fel.")
        elif utan:
            varning(f"{val}: 'rostberattigade' saknas, valdeltagande kan inte visas")


def las_jamforelsefil(path):
    """Valmyndighetens valdistrikt-jamforelser-mellan-2022-och-2026.xlsx -> {kod: True/False} för Majornas koder."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb["Jämförelser"] if "Jämförelser" in wb.sheetnames else wb.worksheets[0]
        rader = ws.iter_rows(values_only=True)
        huvud = [str(h).strip() if h is not None else "" for h in next(rader)]
        if "Valdistriktskod 2026" not in huvud or "Jämförbarhet" not in huvud:
            fel(f"{path}: väntade kolumnerna 'Valdistriktskod 2026' och 'Jämförbarhet', fick {huvud}")
        ik, ij = huvud.index("Valdistriktskod 2026"), huvud.index("Jämförbarhet")
        ut = {}
        for r in rader:
            kod = str(r[ik]).strip() if r[ik] is not None else ""
            if kod in MAJORNA_KODER:
                ut[kod] = str(r[ij]).strip() == valnatt.STATUS_JAMFORBAR
        return ut
    finally:
        wb.close()


def satt_jamforbar(kod, varde, kalla):
    """Samma distrikt får inte bedömas olika av två filer; då vinner 'inte jämförbart' med en varning."""
    if kod in JAMFORBAR and JAMFORBAR[kod] != varde:
        varning(f"{kalla}: {kod} bedöms {'jämförbart' if varde else 'ej jämförbart'} men en tidigare fil sa tvärtom, sätts till ej jämförbart")
        JAMFORBAR[kod] = False
    else:
        JAMFORBAR[kod] = varde


def las_valnattsmapp(mapp, distrikt, bas):
    """Mappen från hamta_2026.py (rd/, rf/, kf/ med JSON). Fyller distrikt, returnerar jämförelseaggregat och riksdagens mandat.

    Fel i jämförelseaggregaten eller riksdagens mandat får aldrig stoppa distriktsimporten: de är bara
    jämförelsemarkörer, distrikten ligger i en annan fil. Ett sådant fel ger varning och lämnar
    aggregatet respektive mandaten tomma för det valet.
    """
    global META_TEST
    mapp = Path(mapp)
    jamforelser = {"goteborg": {}, "riket": {}}
    verklig = {}
    hittade = 0
    for val in VAL:
        sub = mapp / val
        if not sub.is_dir():
            varning(f"{val}: ingen mapp {sub}, valet markeras som oräknat")
            continue
        rost = sorted(sub.glob("*_rostfordelning_*.json"))
        mandat = sorted(sub.glob("*_mandatfordelning_*.json"))
        summering = sorted(sub.glob("*_summering_*.json"))
        if not rost:
            fel(f"{val}: ingen röstfördelningsfil i {sub}")
        if len(rost) > 1:
            varning(f"{val}: {len(rost)} röstfördelningsfiler i {sub}, läser {rost[0].name}, "
                    f"ignorerar {', '.join(p.name for p in rost[1:])}")
        if len(mandat) > 1:
            varning(f"{val}: {len(mandat)} mandatfördelningsfiler i {sub}, läser {mandat[0].name}, "
                    f"ignorerar {', '.join(p.name for p in mandat[1:])}")
        if len(summering) > 1:
            varning(f"{val}: {len(summering)} summeringsfiler i {sub}, läser {summering[0].name}, "
                    f"ignorerar {', '.join(p.name for p in summering[1:])}")
        print(f"Läser {val}: {rost[0].name}")
        hittade += 1
        try:
            ra = valnatt.las_rostfordelning(rost[0], MAJORNA_KODER, kommunkod=None if val == "kf" else valnatt.KOMMUNKOD_GOTEBORG)
        except valnatt.FormatFel as ex:
            fel(f"{val}: {ex}")
        except SummaFel as ex:
            fel(f"{val}: {ex}")
        if ra["val"] != val:
            fel(f"{rost[0]}: filen gäller {ra['val'].upper()} men ligger i mappen {val}")
        if ra["meta"]["test"]:
            META_TEST = True
        if ra["meta"]["antal_i_omradet"] != GOTEBORG_DISTRIKT["json_2026"]:
            varning(f"{val}: Göteborg har {ra['meta']['antal_i_omradet']} distrikt i filen (genrep 2026: {GOTEBORG_DISTRIKT['json_2026']}). "
                    "Kontrollera att de 23 koderna fortfarande täcker Majorna.")
        raknade_har = 0
        okanda_agg = {}
        for kod in MAJORNA_KODER:
            post = ra["distrikt"].get(kod)
            if post is None:
                varning(f"{val}: distrikt {kod} ({distrikt[kod]['namn']}) saknas i filen, markeras som oräknat")
                continue
            satt_jamforbar(kod, post["jamforbar"], f"{val}-filen")
            if bas and post["namn"] != bas_namn(bas, kod) and kod not in NAMNVARNADE:
                varning(f"{val}: {kod} heter '{post['namn']}' i filen men '{bas_namn(bas, kod)}' 2022")
                NAMNVARNADE.add(kod)
            if not post["raknat"]:
                continue
            for etikett, n in post["okanda"].items():
                _lagg_till_okant(okanda_agg, etikett, n, n / post["giltiga"] if post["giltiga"] else 0)
            fyll(distrikt, kod, val, post)
            raknade_har += 1
        _varna_okanda(val, okanda_agg)
        if raknade_har != min(ra["meta"]["raknade"], 23) and ra["meta"]["raknade"] >= ra["meta"]["totalt"]:
            varning(f"{val}: filhuvudet säger {ra['meta']['raknade']} av {ra['meta']['totalt']} räknade men {raknade_har} av 23 Majornadistrikt är räknade")
        if mandat:
            try:
                agg = valnatt.aggregat_2026(val, mandat[0], summering[0] if summering else None)
                for omrade, post in agg.items():
                    jamforelser[omrade][val] = post
            except Exception as ex:
                varning(f"{val}: jämförelseaggregat kunde inte läsas: {type(ex).__name__}: {ex}")
            if val == "rd":
                try:
                    verklig = valnatt.riksdag_verklig(mandat[0])
                except Exception as ex:
                    varning(f"rd: riksdagens mandat kunde inte läsas: {type(ex).__name__}: {ex}")
                    verklig = {}
                if verklig:
                    summa = sum(verklig.values())
                    if summa != 349:
                        varning(f"rd: riksdagens mandat summerar till {summa}, inte 349")
                    for kod, n in verklig.items():
                        if kod not in NYCKELPARTIER["rd"]:
                            varning(f"rd: partikoden {kod} med {n} mandat finns inte bland sidans riksdagspartier")
        else:
            varning(f"{val}: ingen mandatfördelningsfil, jämförelseaggregat saknas")
    if not hittade:
        fel(f"{mapp}: inga av mapparna rd, rf, kf finns")
    return jamforelser, verklig


def skriv_mall(path):
    with open(path, "w", encoding="utf-8", newline="") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(["val", "kod", "parti", "roster"])
        w.writerow(["#", "Fyll i bara sista kolumnen (roster). Rader som börjar med # ignoreras. giltiga ska vara "
                          "summan av partiraderna; rostande och rostberattigade hämtas från val.se."])
        w.writerow(["#", "Mallen täcker riksdagsvalet (rd). Region (rf) och kommun (kf) skrivs med samma format: "
                          "val;kod;parti;roster."])
        bas = las_bas(ROT / "data" / "valdata_2022.json", roll="bas")
        for kod in MAJORNA_KODER:
            w.writerow([])
            w.writerow(["#", kod, bas_namn(bas, kod) if bas else "", ""])
            for p in NYCKELPARTIER["rd"] + [OVRIGA, "giltiga", "rostande", "rostberattigade"]:
                w.writerow(["rd", kod, p, ""])
    print(f"Mall skriven: {path} (rader som börjar med # ignoreras)")


def las_bas(path, roll="ny"):
    """roll="bas": path är basåret för swing (t.ex. valdata_2022.json). Används av skriv_mall (för
    distriktsnamnen i mallen) och av repetera/main (--bas, basåret för swing). roll="ny" (standard):
    path är årets egen, redan skrivna fil; används bara av kontrollera_mot_befintlig, som jämför med
    filen från en tidigare körning innan den skrivs över. Rådet i felet skiljer sig därefter."""
    path = Path(path)
    if not path.exists():
        return None
    try:
        return json.loads(path.read_text("utf-8"))
    except (json.JSONDecodeError, UnicodeDecodeError, OSError) as ex:
        rad = ("Basårets fil måste vara hel, ingen swing beräknas annars." if roll == "bas"
               else "Ta bort eller flytta filen och kör om.")
        fel(f"{path}: har inte formen av en JSON-fil: {ex}. {rad}")


def bygg(ar, distrikt, status, jamforelser, tid, kalla, verklig=None, test=False):
    mandat = {"metod": METOD.format(ar=ar)}
    rd = [d for d in distrikt.values() if d.get("rd")]
    if rd:
        roster = {}
        for d in rd:
            for p, n in d["rd"].items():
                roster[p] = roster.get(p, 0) + n
        mandat["riksdag_majorna"] = jamkade_uddatal(roster, 349)
    if verklig:
        # riksdagens verkliga mandat (halvcirkeln) hör inte till Majornas röster: den ska följa med
        # så snart mandatfördelningsfilen finns, även innan något Majornadistrikt är räknat i rd.
        mandat["riksdag_verklig"] = dict(verklig)
    v = schema.bygg_valdata(ar, list(distrikt.values()), status, jamforelser, mandat, uppdaterad=tid, kalla=kalla)
    for d in v["distrikt"]:
        d["jamforbar_mot_bas"] = bool(JAMFORBAR.get(d["kod"], True))
        d["grans_andrad"] = not d["jamforbar_mot_bas"]
    if test:
        v["meta"]["test"] = True
    return v


def kontrollera_mot_befintlig(ny, path, tvinga):
    """Vägrar skriva över en fil med färre räknade distrikt, eller en skarp fil med testdata.

    En tom import (inget Majornadistrikt räknat i något val än, som kvällens första körningar) är
    bara en varning, aldrig ett stopp: filerna skrivs med 0 räknade och konfigen slår på
    valnattsläget. Fallet "vi hade räknade distrikt, nu noll" fångas ändå av spärren mot färre
    räknade nedan. Alla valens spärrproblem samlas och rapporteras i ett enda FEL, inte ett i taget.
    """
    gammal = las_bas(path)
    raknade_ny = {val: sum(1 for d in ny["distrikt"] if d[val]) for val in VAL}
    if sum(raknade_ny.values()) == 0:
        varning("inget Majornadistrikt räknat i något val än, filerna skrivs med 0 räknade")
    if gammal is None:
        return
    raknade_gammal = {val: sum(1 for d in gammal["distrikt"] if d[val]) for val in VAL}
    problem = []
    farre = [val for val in VAL if raknade_ny[val] < raknade_gammal[val]]
    if farre:
        tal = ", ".join(f"{val}: {raknade_ny[val]} mot {raknade_gammal[val]}" for val in farre)
        problem.append(
            f"färre räknade distrikt än i {path.name} ({tal}). "
            "Vill du skriva ändå: kör med --valnatt-mapp data/valnatt/senaste --tvinga. --tvinga ersätter "
            "hela filen (val som saknas i den nya blir tomma), skriptet slår inte ihop med den gamla. "
            "Kompletterar du för hand: kör om med --valnatt-mapp data/valnatt/senaste --csv FIL, då slås CSV:n "
            "ihop med JSON-vägen (CSV:n vinner per distrikt och val); --tvinga ersätter i stället hela filen.")
    if ny["meta"].get("test") and not gammal["meta"].get("test"):
        problem.append(f"{path.name} är skarp data men den nya filen är testdata")
    if not problem:
        return
    if tvinga:
        for text in problem:
            varning(text)
        return
    fel("\n".join(problem))


def repetera(ut_bas):
    filer = {}
    for val, monster in RAFIL_MONSTER.items():
        traffar = sorted(glob.glob(str(ROT / monster)))
        if not traffar:
            fel(f"hittar ingen råfil för {val} ({monster}) i {ROT}")
        filer[val] = traffar[0]
    bas = las_bas(ut_bas, roll="bas")
    if bas is None:
        fel(f"{ut_bas} saknas, kör scripts/bygg_data.py först")
    distrikt = tomma_distrikt(bas)
    jamforelser = las_rafiler(filer, distrikt, bas)
    v = bygg(2022, distrikt, "slutlig", jamforelser, bas["meta"]["uppdaterad"], bas["meta"]["kalla"])
    diffar = []
    for d, b in zip(v["distrikt"], bas["distrikt"]):
        for val in VAL:
            if d[val] != b[val]:
                diffar.append(f"{d['kod']} {val}: {d[val]} != {b[val]}")
            for f in ("giltiga", "rostande", "rostberattigade"):
                if d[f].get(val) != b[f].get(val):
                    diffar.append(f"{d['kod']} {val} {f}: {d[f].get(val)} != {b[f].get(val)}")
    if v["aggregat"]["majorna"] != bas["aggregat"]["majorna"]:
        diffar.append("aggregat.majorna skiljer sig")
    if v["mandat"]["riksdag_majorna"] != bas["mandat"]["riksdag_majorna"]:
        diffar.append(f"mandat: {v['mandat']['riksdag_majorna']} != {bas['mandat']['riksdag_majorna']}")
    for omrade in ("goteborg", "riket"):
        for val, post in bas["aggregat"][omrade].items():
            ny = v["aggregat"][omrade].get(val)
            if ny is None:
                diffar.append(f"aggregat.{omrade}.{val} saknas")
                continue
            for p, andel in post["andel"].items():
                if abs(ny["andel"].get(p, -1) - andel) > 1e-9:
                    diffar.append(f"aggregat.{omrade}.{val}.{p}: {ny['andel'].get(p)} != {andel}")
            tol = 5e-5 if omrade == "goteborg" else 1e-9
            if abs(ny["valdeltagande"] - post["valdeltagande"]) > tol:
                diffar.append(f"aggregat.{omrade}.{val}.valdeltagande: {ny['valdeltagande']} != {post['valdeltagande']}")
    for rad in diffar:
        print("DIFF " + rad)
    if diffar:
        print(f"REPETITION MISSLYCKADES: {len(diffar)} diffar mot {ut_bas}")
        return 1
    print(f"REPETITION OK: 2022 års råfiler ger exakt samma valdata som {Path(ut_bas).name} "
          f"(23 distrikt, 3 val, aggregat, mandat). Varningar: {len(VARNINGAR)}")
    return 0


def main():
    try:
        sys.stdout.reconfigure(line_buffering=True)   # så "2>&1 | tee" får rätt ordning på Läser- och VARNING-rader
    except AttributeError:
        pass
    VARNINGAR.clear()
    JAMFORBAR.clear()
    NAMNVARNADE.clear()
    global META_TEST
    META_TEST = False
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    for val in VAL:
        ap.add_argument(f"--{val}", help=f"Valmyndighetens xlsx-råfil för {VAL[val].lower()}valet (blad roster_{val.upper()}), slutlig-vägen")
    ap.add_argument("--valnatt-mapp", metavar="MAPP", help="mappen från hamta_2026.py med rd/, rf/ och kf/ (JSON)")
    ap.add_argument("--hamta", action="store_true", help="kör hamta_2026.py först och använd dess mapp")
    ap.add_argument("--tillfalle", choices=["p", "s"], default="p", help="för --hamta: p preliminär, s slutlig")
    ap.add_argument("--genrep", action="store_true", help="för --hamta: simuleringarna i stället för val2026")
    ap.add_argument("--utan-signatur", action="store_true",
                    help="för --hamta: hoppa över signaturkontrollen, reservläge om val.se bekräftar problem med certifikatet")
    ap.add_argument("--csv", help="reservväg: manuellt ifylld CSV (val;kod;parti;roster)")
    ap.add_argument("--skriv-mall", metavar="FIL", help="skriv en tom CSV-mall och avsluta")
    ap.add_argument("--jamforelsefil", default=JAMFORELSEFIL_STANDARD, help="Valmyndighetens xlsx med jämförbarhet 2022 mot 2026")
    ap.add_argument("--ar", type=int, default=2026)
    ap.add_argument("--status", choices=["preliminar", "slutlig"], default="preliminar")
    ap.add_argument("--tid", help="tidsstämpel ISO 8601, t.ex. 2026-09-13T21:30:00 (standard: nu)")
    ap.add_argument("--ut", default=ROT / "data")
    ap.add_argument("--bas", default=ROT / "data" / "valdata_2022.json", help="basår för swing")
    ap.add_argument("--repetera", action="store_true", help="kör 2022 års råfiler och jämför med --bas")
    ap.add_argument("--valnatt", action="store_true",
                    help="uppdatera data/konfig: lägg till året, gör det till standard och slå på valnattsläget")
    ap.add_argument("--tvinga", action="store_true",
                    help="skriv ändå vid färre räknade distrikt: ersätter hela filen, slår inte ihop med den gamla. "
                         "Efter ett sådant stopp: kör om med --valnatt-mapp data/valnatt/senaste --tvinga "
                         "(--hamta med samma filer ger kod 3, inget nytt att hämta)")
    a = ap.parse_args()

    if a.skriv_mall:
        skriv_mall(a.skriv_mall)
        return 0
    if a.repetera:
        return repetera(a.bas)
    filer = {val: getattr(a, val) for val in VAL if getattr(a, val)}
    if a.hamta:
        argv = (["hamta_2026.py", "--tillfalle", a.tillfalle, "--ut", str(Path(a.ut) / "valnatt"), "--bara-om-nytt"]
                + (["--genrep"] if a.genrep else []) + (["--utan-signatur"] if a.utan_signatur else []))
        sparad = sys.argv
        sys.argv = argv
        try:
            kod = hamta_2026.main()
        except SystemExit as ex:
            fel(f"hämtningen avbröts med kod {ex.code}")
        except Exception as ex:
            fel(f"hämtningen misslyckades: {type(ex).__name__}: {ex}")
        finally:
            sys.argv = sparad
        if kod == 3:
            print("Inget nytt att läsa in.")
            return 3
        if kod != 0:
            fel("hämtningen misslyckades, inget skrivet")
        a.valnatt_mapp = str(Path(a.ut) / "valnatt" / "senaste")
    if not filer and not a.csv and not a.valnatt_mapp:
        ap.error("ange --valnatt-mapp, --hamta, minst en av --rd/--rf/--kf, eller --csv (eller --repetera / --skriv-mall)")
    bas = las_bas(a.bas, roll="bas")
    if bas is None:
        varning(f"{a.bas} saknas: ingen swing beräknas och distriktsnamnen tas ur filen")
    distrikt = tomma_distrikt(bas)
    jamforelser = {"goteborg": {}, "riket": {}}
    verklig = {}
    if Path(a.jamforelsefil).exists():
        try:
            flaggor = las_jamforelsefil(a.jamforelsefil)
        except Exception as ex:
            fel(f"{a.jamforelsefil}: kunde inte läsas: {type(ex).__name__}: {ex}")
        for kod, ok in flaggor.items():
            satt_jamforbar(kod, ok, "jämförelsefilen")
    try:
        if a.valnatt_mapp:
            jamforelser, verklig = las_valnattsmapp(a.valnatt_mapp, distrikt, bas)
        if filer:
            try:
                jamforelser = las_rafiler(filer, distrikt, bas)
            except FormatFel:
                raise
            except Exception as ex:
                fel(f"{', '.join(str(p) for p in filer.values())}: kunde inte läsas: {type(ex).__name__}: {ex}")
        if a.csv:
            print(f"Läser CSV: {a.csv}")
            # CSV-vägen ger inga jämförelseaggregat (goteborg/riket): finns en JSON-väg körd innan (eller i
            # samma anrop via --valnatt-mapp) behålls dess jamforelser oförändrade här.
            las_csv(a.csv, distrikt)
    except FormatFel as ex:
        fel(str(ex))
    kontrollera_rostberattigade(distrikt)
    if not JAMFORBAR:
        varning("ingen uppgift om jämförbarhet mot 2022 (varken JSON eller jämförelsefil): alla distrikt antas jämförbara")
    if META_TEST:
        varning("TESTDATA: minst en fil har test: true. Filerna skrivs med meta.test och får inte publiceras som skarpa.")
    ut = Path(a.ut)
    if META_TEST and ut.resolve() == (ROT / "data").resolve() and not a.tvinga:
        fel("filerna är testdata (test: true) och --ut är repots data/: kör med --tvinga om det är avsiktligt, "
            "annars --ut till en annan mapp")
    kalla = f"Valmyndigheten, {'preliminär' if a.status == 'preliminar' else 'slutlig'} rösträkning per valdistrikt {a.ar}"
    v = bygg(a.ar, distrikt, a.status, jamforelser, a.tid, kalla, verklig, META_TEST)
    swing_obj = None
    if bas:
        jamforbara = [d["kod"] for d in v["distrikt"] if d["jamforbar_mot_bas"]]
        swing_obj = schema.swing(v, bas, jamforbara=jamforbara)
    konfig = None
    if a.valnatt:
        try:
            konfig = schema.las_konfig(ut)
        except (json.JSONDecodeError, UnicodeDecodeError, OSError) as ex:
            fel(f"{ut / 'konfig.json'}: kunde inte läsas: {type(ex).__name__}: {ex}")
        konfig["ar"] = sorted(set(konfig.get("ar", [])) | {str(a.ar)})
        konfig["standardAr"] = str(a.ar)
        konfig["valnatt"] = True
    kontrollera_mot_befintlig(v, ut / f"valdata_{a.ar}.json", a.tvinga)
    filer_ut = list(schema.skriv(ut / f"valdata_{a.ar}", v))
    if swing_obj is not None:
        filer_ut += list(schema.skriv(ut / f"swing_{a.ar}", swing_obj))
    if konfig is not None:
        filer_ut += list(schema.skriv_konfig(ut, konfig))
    raknade = {val: sum(1 for d in v["distrikt"] if d[val]) for val in VAL}
    n = len(v["distrikt"])
    print("Räknade distrikt: " + ", ".join(f"{VAL[val]} {k}/{n}" for val, k in raknade.items()))
    ej = [d["namn"] for d in v["distrikt"] if d["grans_andrad"]]
    print(f"Jämförbara mot {bas['meta']['ar'] if bas else 'basåret'}: {n - len(ej)} av {n}. Omritade: {', '.join(ej) or 'inga'}")
    for f in filer_ut:
        print(f"    {f.stat().st_size / 1024:6.1f} kB  {f}")
    print(f"{len(VARNINGAR)} varningar. Status: {a.status}. Uppdaterad: {v['meta']['uppdaterad']}")
    if a.valnatt:
        print(f"Konfigen uppdaterad: ar {konfig['ar']}, standardAr {a.ar}, valnatt på. Ladda upp data/ (valdata, swing och konfig).")
    else:
        print("Nästa steg: kör med --valnatt för att slå på valnattsläget i data/konfig.js, eller redigera filen för hand. Ladda sedan upp data/.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
