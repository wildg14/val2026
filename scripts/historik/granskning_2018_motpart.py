# -*- coding: utf-8 -*-
"""
granskning_2018_motpart.py - oberoende motgranskning av 2018 ars filer i data/historik/.

Skriven fran grunden av granskaren (etikett granskning_2018). Laser kallfilerna sjalv och
raknar om allt, i stallet for att aterbruka val2018_bygg.py eller granskning_2018_kontroll.py.

Kontroller:
  1  Filformat: UTF-8 utan BOM, LF, semikolon, punkt som decimaltecken, 8-siffriga koder som text.
  2  Egen inlasning av 2018_R/L/K_per_valdistrikt.xlsx (antal- och procentflik) och jamforelse
     rad for rad, parti for parti, mot roster_/distrikt_/ogiltiga_/aggregat_-filerna.
  3  Interna identiteter per distrikt: partisumma = giltiga, giltiga + ogiltiga = rostande,
     valdeltagande = 100 * rostande / rostberattigade.
  4  Andra kallor: shapefilen alla_valdistrikt.dbf (koder, namn), reparerade rostberattigade-filer,
     2018_mandat.xlsx, slutligt-valresultat-riksdagen-jamforande-statistik-2018-2022.xlsx,
     vd-indelning-2018.skv och vd-mappning-2014-2018.skv.
  5  Partitackning: varje partikolumn med minst en Goteborgsrost finns i roster-filen, och ingen
     kolumn med noll roster finns med.
  6  Stickprov: tre slumpade Majornadistrikt (annat fro an forra granskningen), tre partital var
     mot exakt cell i kallfilen, samt hela distriktssidan pa historik.val.se (HTML) med korrekt
     tabelltolkning.
  7  Partinamn utan forkortning (BASIP, INI, NYREF) sokes upp pa distriktssidorna.

Kors med:
  <venv>/bin/python scripts/historik/granskning_2018_motpart.py
Skriver bara till stdout och cachar HTML i scratchpad. Andrar inga CSV-filer.
"""
import collections
import csv
import html
import os
import random
import re
import sys
import urllib.request

import openpyxl
from openpyxl.utils import get_column_letter
from dbfread import DBF

KALLMAPP = "/Users/daniel/code/Temp/Historiska dokument"
SCRATCH = "/private/tmp/claude-501/-Users-daniel-code-Temp/f347baf2-1af3-43b3-b36d-e80b868ede0e/scratchpad"
UTMAPP = "/Users/daniel/code/Temp/data/historik"

FIL = {"rd": os.path.join(KALLMAPP, "2018_R_per_valdistrikt.xlsx"),
       "rf": os.path.join(KALLMAPP, "2018_L_per_valdistrikt.xlsx"),
       "kf": os.path.join(KALLMAPP, "2018_K_per_valdistrikt.xlsx")}
FLIK = {"rd": ("R antal", "R procent"), "rf": ("L antal", "L procent"), "kf": ("K antal", "K procent")}
VALBOKSTAV = {"rd": "R", "rf": "L", "kf": "K"}
FIL_MANDAT = os.path.join(KALLMAPP, "2018_mandat.xlsx")
FIL_JAMF = "/Users/daniel/code/Temp/slutligt-valresultat-riksdagen-jamforande-statistik-2018-2022.xlsx"
SHP_DBF = SCRATCH + "/unz/2018_valgeografi_valdistrikt/alla_valdistrikt.dbf"
FIL_ROSTBER = {"rd": SCRATCH + "/repaired/2018_rostberattigade_R.xlsx",
               "rf": SCRATCH + "/repaired/2018_rostberattigade_L.xlsx",
               "kf": SCRATCH + "/repaired/2018_rostberattigade_K.xlsx"}
SKV_INDELNING = SCRATCH + "/unz/mappning_2014_2018/vd-indelning-2018.skv"
SKV_MAPPNING = SCRATCH + "/unz/mappning_2014_2018/vd-mappning-2014-2018.skv"
HTML_CACHE = SCRATCH + "/dl2018_motpart"
HTML_URL = "https://historik.val.se/val/val2018/slutresultat/{V}/valdistrikt/14/80/{vd}/index.html"

MAJORNA = ["14801011", "14801012", "14801013", "14801014", "14801015", "14801016", "14801017",
           "14801021", "14801022", "14801031", "14801032", "14801033", "14801034", "14801035",
           "14801036", "14801037", "14801038", "14801041", "14801042", "14801043", "14801044",
           "14801045"]

fel, varning, info = [], [], []


def rap(kat, text):
    {"FEL": fel, "VARNING": varning}.get(kat, info).append(text)
    print(f"[{kat}] {text}")


def i0(v):
    return int(v) if v not in (None, "") else 0


def normalisera(p):
    return {"FP": "L", "DEM": "D", "KP": "K"}.get(p, p.upper())


# ---------------------------------------------------------------- 1. filformat
print("== 1. filformat")
FILER = [f"{s}_2018_{v}.csv" for s in ("roster", "distrikt", "ogiltiga", "aggregat") for v in ("rd", "rf", "kf")]
FILER += ["mandat_2018.csv", "partier_2018.csv"]
for namn in FILER:
    sokvag = os.path.join(UTMAPP, namn)
    if not os.path.exists(sokvag):
        rap("FEL", f"{namn}: filen saknas")
        continue
    b = open(sokvag, "rb").read()
    if b.startswith(b"\xef\xbb\xbf"):
        rap("FEL", f"{namn}: BOM")
    try:
        text = b.decode("utf-8")
    except UnicodeDecodeError as e:
        rap("FEL", f"{namn}: inte UTF-8 ({e})")
        continue
    if b"\r" in b:
        rap("FEL", f"{namn}: CR i radbrytningarna")
    if not text.endswith("\n"):
        rap("VARNING", f"{namn}: ingen radbrytning sist")
    rader = text.rstrip("\n").split("\n")
    n = rader[0].count(";") + 1
    dalig = [i for i, r in enumerate(rader, 1) if r.count(";") + 1 != n]
    if dalig:
        rap("FEL", f"{namn}: rader med fel antal falt: {dalig[:5]}")
    if "," in text and re.search(r"\d,\d", text):
        rap("FEL", f"{namn}: decimalkomma forekommer")
print(f"  {len(FILER)} filer kontrollerade")


def las_csv(namn):
    with open(os.path.join(UTMAPP, namn), encoding="utf-8", newline="") as f:
        return list(csv.DictReader(f, delimiter=";"))


roster = {v: las_csv(f"roster_2018_{v}.csv") for v in ("rd", "rf", "kf")}
distrikt = {v: las_csv(f"distrikt_2018_{v}.csv") for v in ("rd", "rf", "kf")}
ogilt = {v: las_csv(f"ogiltiga_2018_{v}.csv") for v in ("rd", "rf", "kf")}
aggregat = {v: las_csv(f"aggregat_2018_{v}.csv") for v in ("rd", "rf", "kf")}
mandat = las_csv("mandat_2018.csv")
partier = las_csv("partier_2018.csv")

for namn, filer in (("roster", roster), ("distrikt", distrikt), ("ogiltiga", ogilt)):
    for v, rader in filer.items():
        d = [r["kod"] for r in rader if not re.fullmatch(r"\d{8}", r["kod"])]
        if d:
            rap("FEL", f"{namn}_{v}: koder utan 8 siffror: {d[:5]}")
        if any(r["ar"] != "2018" or r["val"] != v for r in rader):
            rap("FEL", f"{namn}_{v}: fel varde i ar eller val")
for v in ("rd", "rf", "kf"):
    nycklar = [(r["kod"], r["parti_kalla"]) for r in roster[v]]
    dubbl = [k for k, c in collections.Counter(nycklar).items() if c > 1]
    if dubbl:
        rap("FEL", f"roster_{v}: dubbletter {dubbl[:5]}")

# ---------------------------------------------------------------- 2. egen inlasning av kallan
print("== 2. egen inlasning av kallfilerna")
kalla, kalla_rad, partikol, hdr_val = {}, {}, {}, {}
riket, vgr = {}, {}
for v in ("rd", "rf", "kf"):
    wb = openpyxl.load_workbook(FIL[v], read_only=True, data_only=True)
    ws = wb[FLIK[v][0]]
    rader = list(ws.iter_rows(values_only=True))
    hdr = [str(h) if h is not None else "" for h in rader[0]]
    hdr_val[v] = hdr
    ix = {h: i for i, h in enumerate(hdr)}
    pk = list(range(ix["VALDISTRIKTSNAMN"] + 1, ix["OGEJ"]))
    partikol[v] = {hdr[i]: i for i in pk}
    d, dr = {}, {}
    tr = collections.Counter()
    tv = collections.Counter()
    sr = collections.Counter()
    sv = collections.Counter()
    for radnr, r in enumerate(rader[1:], start=2):
        lan, kom = r[ix["LÄNSKOD"]], r[ix["KOMMUNKOD"]]
        for p, i in partikol[v].items():
            tr[p] += i0(r[i])
        sr["giltiga"] += i0(r[ix["RÖSTER GILTIGA"]])
        sr["rostande"] += i0(r[ix["RÖSTANDE"]])
        sr["rostber"] += i0(r[ix["RÖSTBERÄTTIGADE"]])
        if lan == 14:
            for p, i in partikol[v].items():
                tv[p] += i0(r[i])
            sv["giltiga"] += i0(r[ix["RÖSTER GILTIGA"]])
            sv["rostande"] += i0(r[ix["RÖSTANDE"]])
            sv["rostber"] += i0(r[ix["RÖSTBERÄTTIGADE"]])
        if lan != 14 or kom != 80:
            continue
        kod = f"{int(lan):02d}{int(kom):02d}{int(r[ix['VALDISTRIKTSKOD']]):04d}"
        if kod in d:
            rap("FEL", f"{v}: dubblettkod {kod} i kallfilen")
        d[kod] = {"namn": r[ix["VALDISTRIKTSNAMN"]] or "",
                  "valkrets": r[ix["VALKRETSNAMN"]] or "",
                  "OGEJ": i0(r[ix["OGEJ"]]), "BLANK": i0(r[ix["BLANK"]]), "OG": i0(r[ix["OG"]]),
                  "giltiga": i0(r[ix["RÖSTER GILTIGA"]]), "rostande": i0(r[ix["RÖSTANDE"]]),
                  "rostber": i0(r[ix["RÖSTBERÄTTIGADE"]]), "valdelt": r[ix["VALDELTAGANDE"]],
                  "partier": {hdr[i]: i0(r[i]) for i in pk}}
        dr[kod] = radnr
    ws2 = wb[FLIK[v][1]]
    rader2 = list(ws2.iter_rows(values_only=True))
    hdr2 = [str(h) if h is not None else "" for h in rader2[0]]
    ix2 = {h: i for i, h in enumerate(hdr2)}
    if hdr2[:len(hdr)] != hdr:
        rap("VARNING", f"{v}: procentflikens rubriker skiljer fran antalflikens")
    for r in rader2[1:]:
        if r[ix2["LÄNSKOD"]] != 14 or r[ix2["KOMMUNKOD"]] != 80:
            continue
        kod = f"1480{int(r[ix2['VALDISTRIKTSKOD']]):04d}"
        d[kod]["procent"] = {hdr2[i]: r[i] for i in pk}
    kalla[v], kalla_rad[v] = d, dr
    riket[v] = (tr, sr)
    vgr[v] = (tv, sv)
    print(f"  {v}: {len(rader)-1} rader i filen, {len(d)} Goteborgsrader, {len(pk)} partikolumner, "
          f"riket giltiga {sr['giltiga']}, VG giltiga {sv['giltiga']}")

# ---------------------------------------------------------------- 3. jamforelse mot CSV
print("== 3. distrikt och partital mot kallan")
for v in ("rd", "rf", "kf"):
    rs = collections.defaultdict(dict)
    for r in roster[v]:
        rs[r["kod"]][r["parti_kalla"]] = r
    d_by = {r["kod"]: r for r in distrikt[v]}
    o_by = {r["kod"]: r for r in ogilt[v]}
    if set(d_by) != set(kalla[v]) or set(rs) != set(kalla[v]) or set(o_by) != set(kalla[v]):
        rap("FEL", f"{v}: kodmangderna i distrikt/roster/ogiltiga och kallan skiljer sig")
    n_par = 0
    for kod, k in kalla[v].items():
        d = d_by.get(kod)
        if d is None:
            continue
        # partital
        for p, antal in k["partier"].items():
            r = rs[kod].get(p)
            if r is None:
                if antal:
                    rap("FEL", f"roster_{v} {kod} {p}: saknas i CSV, kallan har {antal}")
                continue
            n_par += 1
            if int(r["roster"]) != antal:
                rap("FEL", f"roster_{v} {kod} {p}: CSV {r['roster']} != kalla {antal}")
            if r["parti"] != normalisera(p):
                rap("FEL", f"roster_{v} {kod} {p}: normaliserat {r['parti']} != {normalisera(p)}")
            pa = k["procent"].get(p)
            pa = 0.0 if pa in (None, "") else float(pa)
            if abs(float(r["andel"]) - pa) > 0.005:
                rap("FEL", f"roster_{v} {kod} {p}: andel {r['andel']} != procentflik {pa}")
            if k["giltiga"] and abs(float(r["andel"]) - 100 * antal / k["giltiga"]) > 0.0101:
                rap("FEL", f"roster_{v} {kod} {p}: andel {r['andel']} != omraknad "
                           f"{100*antal/k['giltiga']:.4f}")
        for p in rs[kod]:
            if p not in k["partier"]:
                rap("FEL", f"roster_{v} {kod}: parti {p} finns inte som kolumn i kallan")
        # distriktsfalten
        if d["namn"] != k["namn"]:
            rap("FEL", f"distrikt_{v} {kod}: namn {d['namn']!r} != kalla {k['namn']!r}")
        if d["valkrets"] != k["valkrets"]:
            rap("FEL", f"distrikt_{v} {kod}: valkrets {d['valkrets']!r} != kalla {k['valkrets']!r}")
        if (int(d["giltiga"]), int(d["rostande"]), int(d["rostberattigade"])) != \
                (k["giltiga"], k["rostande"], k["rostber"]):
            rap("FEL", f"distrikt_{v} {kod}: giltiga/rostande/rostberattigade skiljer fran kallan")
        if int(d["blanka"]) != k["BLANK"]:
            rap("FEL", f"distrikt_{v} {kod}: blanka {d['blanka']} != kalla {k['BLANK']}")
        if int(d["ogiltiga_ovriga"]) != k["OGEJ"] + k["OG"]:
            rap("FEL", f"distrikt_{v} {kod}: ogiltiga_ovriga {d['ogiltiga_ovriga']} != OGEJ+OG "
                       f"{k['OGEJ'] + k['OG']}")
        if int(d["ogiltiga"]) != k["OGEJ"] + k["BLANK"] + k["OG"]:
            rap("FEL", f"distrikt_{v} {kod}: ogiltiga {d['ogiltiga']} != OGEJ+BLANK+OG")
        o = o_by[kod]
        if (int(o["ogiltiga_ej_anmalda"]), int(o["blanka"]), int(o["ogiltiga_ovriga"])) != \
                (k["OGEJ"], k["BLANK"], k["OG"]):
            rap("FEL", f"ogiltiga_{v} {kod}: skiljer fran kallan")
        # interna identiteter
        s = sum(int(r["roster"]) for r in rs[kod].values())
        if s != int(d["giltiga"]):
            rap("FEL", f"{v} {kod}: partisumma {s} != giltiga {d['giltiga']}")
        if int(d["giltiga"]) + int(d["ogiltiga"]) != int(d["rostande"]):
            rap("FEL", f"{v} {kod}: giltiga+ogiltiga != rostande")
        rb = int(d["rostberattigade"])
        if rb:
            vd = round(100 * int(d["rostande"]) / rb, 2)
            if abs(vd - float(d["valdeltagande"])) > 0.0101:
                rap("FEL", f"{v} {kod}: valdeltagande {d['valdeltagande']} != omraknat {vd}")
        elif float(d["valdeltagande"]) != 0:
            rap("FEL", f"{v} {kod}: rostberattigade 0 men valdeltagande {d['valdeltagande']}")
        if abs(float(d["valdeltagande"]) - float(k["valdelt"] or 0)) > 0.005:
            rap("FEL", f"{v} {kod}: valdeltagande {d['valdeltagande']} != kalla {k['valdelt']}")
    print(f"  {v}: {len(kalla[v])} distrikt, {n_par} (distrikt, parti)-par jamforda")

# ---------------------------------------------------------------- 4. partitackning
print("== 4. partitackning i roster-filerna")
for v in ("rd", "rf", "kf"):
    gbg = collections.Counter()
    for k in kalla[v].values():
        for p, a in k["partier"].items():
            gbg[p] += a
    i_csv = {r["parti_kalla"] for r in roster[v]}
    saknas = sorted(p for p, a in gbg.items() if a > 0 and p not in i_csv)
    tomma = sorted(p for p in i_csv if gbg[p] == 0)
    if saknas:
        rap("FEL", f"roster_{v}: partier med roster i Goteborg som saknas: {saknas}")
    if tomma:
        rap("VARNING", f"roster_{v}: partier utan en enda Goteborgsrost finns med: {tomma}")
    utan = sorted(p for p, a in gbg.items() if a == 0)
    print(f"  {v}: {len(i_csv)} partier i CSV, {len(gbg)} kolumner i kallan, "
          f"{len(utan)} kolumner utan Goteborgsroster (utelamnade): {utan[:12]}")

# ---------------------------------------------------------------- 5. aggregat
print("== 5. aggregat mot omraknade summor")
for v in ("rd", "rf", "kf"):
    gbg = collections.Counter()
    g = ro = rb = 0
    for k in kalla[v].values():
        for p, a in k["partier"].items():
            gbg[p] += a
        g += k["giltiga"]; ro += k["rostande"]; rb += k["rostber"]
    for niva, (t, s) in (("goteborg", (gbg, {"giltiga": g, "rostande": ro, "rostber": rb})),
                         ("riket", riket[v]), ("vgregion", vgr[v])):
        agg = {r["parti_kalla"]: r for r in aggregat[v] if r["niva"] == niva}
        for p, a in t.items():
            if a == 0:
                if p in agg:
                    rap("VARNING", f"aggregat_{v} {niva}: {p} med 0 roster finns med")
                continue
            r = agg.get(p)
            if r is None:
                rap("FEL", f"aggregat_{v} {niva}: {p} saknas ({a} roster)")
                continue
            if int(r["roster"]) != a:
                rap("FEL", f"aggregat_{v} {niva} {p}: {r['roster']} != omraknat {a}")
            if (int(r["giltiga"]), int(r["rostande"]), int(r["rostberattigade"])) != \
                    (s["giltiga"], s["rostande"], s["rostber"]):
                rap("FEL", f"aggregat_{v} {niva} {p}: totaler skiljer fran omraknat")
            if abs(float(r["andel"]) - 100 * a / s["giltiga"]) > 0.0101:
                rap("FEL", f"aggregat_{v} {niva} {p}: andel {r['andel']} != "
                           f"{100*a/s['giltiga']:.4f}")
        if sum(t.values()) != s["giltiga"]:
            rap("FEL", f"{v} {niva}: partisumma {sum(t.values())} != giltiga {s['giltiga']}")
        for p in agg:
            if t[p] == 0:
                rap("FEL", f"aggregat_{v} {niva}: {p} finns men har 0 roster i kallan")
    # summa av CSV-raderna
    scsv = collections.Counter()
    for r in roster[v]:
        scsv[r["parti_kalla"]] += int(r["roster"])
    for p, a in scsv.items():
        if a != gbg[p]:
            rap("FEL", f"roster_{v}: summa {p} {a} != kalla {gbg[p]}")
    sd = (sum(int(r["giltiga"]) for r in distrikt[v]), sum(int(r["rostande"]) for r in distrikt[v]),
          sum(int(r["rostberattigade"]) for r in distrikt[v]))
    if sd != (g, ro, rb):
        rap("FEL", f"distrikt_{v}: summor {sd} != kalla {(g, ro, rb)}")
    print(f"  {v} Goteborg: giltiga {g}, rostande {ro}, rostberattigade {rb}, "
          f"valdeltagande {100*ro/rb:.2f}")

# ---------------------------------------------------------------- 6. andra kallor
print("== 6. andra kallor")
shp = {r["VD"]: r["VD_NAMN"] for r in DBF(SHP_DBF, encoding="utf-8") if str(r["VD"]).startswith("1480")}
for v in ("rd", "rf", "kf"):
    koder = {r["kod"] for r in distrikt[v]}
    if koder - set(shp) != {"14800000"} or set(shp) - koder:
        rap("FEL", f"distrikt_{v}: kodmangd mot shapefil, extra {sorted(koder-set(shp))}, "
                   f"saknas {sorted(set(shp)-koder)}")
    nd = [(r["kod"], r["namn"], shp[r["kod"]]) for r in distrikt[v]
          if r["kod"] in shp and r["namn"] != shp[r["kod"]]]
    if nd:
        rap("FEL", f"distrikt_{v}: namnskillnad mot shapefil {nd[:5]}")
print(f"  shapefil: {len(shp)} Goteborgsdistrikt, distriktsfilerna 351 rader (inkl uppsamling)")

ind = {}
with open(SKV_INDELNING, encoding="iso-8859-1") as f:
    for rad in f:
        if rad.startswith("#"):
            continue
        k, i = rad.rstrip("\n").split(";")
        ind[k] = i
gbg_ind = sorted(k for k in ind if k.startswith("1480"))
if set(gbg_ind) != set(shp):
    rap("VARNING", f"vd-indelning-2018: {len(gbg_ind)} Goteborgskoder mot shapefilens {len(shp)}, "
                   f"differens {sorted(set(gbg_ind) ^ set(shp))[:5]}")
ml = sorted(k for k in shp if shp[k].startswith("Majorna-Linné"))
print(f"  vd-indelning-2018: {len(gbg_ind)} Goteborgskoder; Majorna-Linne i shapefilen "
      f"{len(ml)} distrikt {ml[0]}-{ml[-1]}")
mlf = sorted(r["kod"] for r in distrikt["rd"] if r["namn"].startswith("Majorna-Linné"))
if mlf != ml:
    rap("FEL", f"Majorna-Linne: distriktsfil {len(mlf)} != shapefil {len(ml)}")
if not set(MAJORNA) <= set(mlf):
    rap("FEL", "de 22 angivna Majornakoderna ingar inte alla i Majorna-Linne")

# Procenten i vd-mappning-2014-2018.skv ar andelen AV DET GAMLA distriktet som gick till det nya:
# summan per 2014-kod ar 100, medan summan per 2018-kod kan vara vad som helst.
mapp = collections.defaultdict(list)
per_gammal = collections.defaultdict(float)
with open(SKV_MAPPNING, encoding="iso-8859-1") as f:
    for rad in f:
        if rad.startswith("#"):
            continue
        fg, ny_, pct_ = rad.rstrip("\n").split(";")
        mapp[ny_].append((fg, float(pct_)))
        per_gammal[fg] += float(pct_)
avv = [(k, v) for k, v in per_gammal.items() if abs(v - 100) > 0.15]
if avv:
    rap("FEL", f"vd-mappning-2014-2018: {len(avv)} 2014-koder vars procent inte summerar till 100, "
               f"exempel {avv[:5]}")
utan_mapp = sorted(k for k in shp if k not in mapp)
nya = sorted(k for k in utan_mapp if ind.get(k) == "N")
if sorted(utan_mapp) != nya:
    rap("FEL", f"vd-mappning-2014-2018: Goteborgsdistrikt utan 2014-koppling som inte ar nya: "
               f"{[k for k in utan_mapp if k not in nya]}")
print(f"  vd-mappning-2014-2018: {len([k for k in shp if k in mapp])} av {len(shp)} "
      f"Goteborgsdistrikt har minst en 2014-koppling; {len(utan_mapp)} utan "
      f"({utan_mapp}), alla med indelning N (nytt); "
      f"{len(per_gammal)} 2014-koder summerar till 100 procent (avrundning hogst 0.1)")

# rostberattigade per distrikt ur de reparerade filerna
for v in ("rd", "rf", "kf"):
    wb = openpyxl.load_workbook(FIL_ROSTBER[v], read_only=True, data_only=True)
    rows = list(wb.worksheets[0].iter_rows(values_only=True))
    hdr = list(rows[0])
    kolv = [i for i, h in enumerate(hdr) if isinstance(h, str) and h.startswith(("Män/", "Kvinnor/"))]
    ivd, ikom = hdr.index("valdistrikt_id"), hdr.index("kommun_id")
    tot = {str(r[ivd]).zfill(8): sum(i0(r[i]) for i in kolv) for r in rows[1:] if str(r[ikom]) == "1480"}
    d_by = {r["kod"]: int(r["rostberattigade"]) for r in distrikt[v]}
    diff = [(k, d_by.get(k), a) for k, a in tot.items() if d_by.get(k) != a]
    for k, a, b in diff[:10]:
        rap("FEL", f"rostberattigade {v} {k}: distriktsfil {a} != aldersfil {b}")
    print(f"  rostberattigade {v}: {len(tot)} distrikt i aldersfilen, summa {sum(tot.values())}, "
          f"distriktsfil {sum(d_by.values())}, avvikelser {len(diff)}")

# mandatfilen
wb = openpyxl.load_workbook(FIL_MANDAT, read_only=True, data_only=True)
rader = list(wb["Mandatfördelning"].iter_rows(values_only=True))
hdr = [str(h) if h is not None else "" for h in rader[0]]
ix = {h: i for i, h in enumerate(hdr)}
m_csv = {(r["val"], r["niva"], r["parti_kalla"]): r for r in mandat}
n_m = 0
for r in rader[1:]:
    vt = r[ix["VALTYP"]]
    if vt == "R" and r[ix["VALKRETSKOD"]] == 16 and r[ix["VALKRETS"]] == "Göteborgs kommun":
        key = ("rd", "goteborg")
    elif vt == "L" and r[ix["LÄN"]] == 14 and r[ix["VALKRETSKOD"]] == 1 and r[ix["VALKRETS"]] == "Göteborgs kommun":
        key = ("rf", "goteborg")
    elif vt == "K" and r[ix["LÄN"]] == 14 and r[ix["VALOMRÅDE"]] == "Göteborg":
        key = ("kf", "goteborg")
    else:
        continue
    n_m += 1
    c = m_csv.get(key + (r[ix["PARTIFÖRKORTNING"]],))
    if c is None:
        rap("FEL", f"mandat_2018: {key} {r[ix['PARTIFÖRKORTNING']]} saknas")
        continue
    if (int(c["roster"]), int(c["mandat"]), int(c["fasta_mandat"]), int(c["utjamningsmandat"])) != \
            (r[ix["RÖSTER"]], i0(r[ix["SUMMA MANDAT"]]), i0(r[ix["FASTA MANDAT"]]), i0(r[ix["UTJÄMNINGSMANDAT"]])):
        rap("FEL", f"mandat_2018 {key} {r[ix['PARTIFÖRKORTNING']]}: skiljer fran 2018_mandat.xlsx")
for r in mandat:
    if int(r["fasta_mandat"]) + int(r["utjamningsmandat"]) != int(r["mandat"]):
        rap("FEL", f"mandat_2018 {r['val']} {r['niva']} {r['parti_kalla']}: fasta+utjamning != mandat")
for val, niva, vantat in (("rd", "riket", 349), ("rf", "vgregion", 149), ("kf", "goteborg", 81)):
    s = sum(int(r["mandat"]) for r in mandat if r["val"] == val and r["niva"] == niva)
    if s != vantat:
        rap("FEL", f"mandat_2018 {val} {niva}: {s} mandat, vantat {vantat}")
print(f"  2018_mandat.xlsx: {n_m} Goteborgsrader jamforda; mandatsummor 349/149/81 kontrollerade")

# jamforande statistik 2018-2022 (oberoende publikation)
NAMN2KOD = {"Moderaterna": "M", "Centerpartiet": "C", "Liberalerna (tidigare Folkpartiet)": "L",
            "Kristdemokraterna": "KD", "Arbetarepartiet-Socialdemokraterna": "S",
            "Vänsterpartiet": "V", "Miljöpartiet de gröna": "MP", "Sverigedemokraterna": "SD"}
wb = openpyxl.load_workbook(FIL_JAMF, read_only=True, data_only=True)
agg_rd = {(r["niva"], r["parti_kalla"]): r for r in aggregat["rd"]}
n = 0
for r in wb["Riket"].iter_rows(values_only=True, min_row=2):
    if r[1] in NAMN2KOD and r[6] is not None:
        p = NAMN2KOD[r[1]]
        a = agg_rd.get(("riket", p))
        n += 1
        if a is None or int(a["roster"]) != int(r[6]):
            rap("FEL", f"jamforande Riket {p}: {r[6]} != aggregat {a and a['roster']}")
for r in wb["Valkrets"].iter_rows(values_only=True, min_row=2):
    if r[1] == "Göteborgs kommun" and r[2] in NAMN2KOD and r[7] is not None:
        p = NAMN2KOD[r[2]]
        a = agg_rd.get(("goteborg", p))
        n += 1
        if a is None or int(a["roster"]) != int(r[7]):
            rap("FEL", f"jamforande Valkrets Goteborg {p}: {r[7]} != aggregat {a and a['roster']}")
print(f"  jamforande statistik 2018-2022: {n} partital jamforda (riket och valkrets Goteborg)")

# ---------------------------------------------------------------- 7. stickprov mot rakcell
print("== 7. stickprov: tre Majornadistrikt, tre partital var per val")
random.seed(20181)
urval = sorted(random.sample(MAJORNA, 3))
print(f"  slumpade distrikt: {urval}")
stickprov = []
for v in ("rd", "rf", "kf"):
    wb = openpyxl.load_workbook(FIL[v], read_only=True, data_only=True)
    ws = wb[FLIK[v][0]]
    rs = {(r["kod"], r["parti_kalla"]): r for r in roster[v]}
    for kod in urval:
        radnr = kalla_rad[v][kod]
        pval = sorted(p for p, a in kalla[v][kod]["partier"].items() if a > 0)
        for p in random.sample(pval, 3):
            kol = partikol[v][p]
            cell = ws.cell(row=radnr, column=kol + 1).value
            csvv = rs[(kod, p)]["roster"]
            ok = i0(cell) == int(csvv)
            rad = (f"{os.path.basename(FIL[v])} flik '{FLIK[v][0]}' rad {radnr} kolumn "
                   f"{get_column_letter(kol+1)} ({hdr_val[v][kol]}) = {cell!r}; "
                   f"roster_2018_{v}.csv {kod} {p} = {csvv}")
            stickprov.append((v, kod, p, rad, ok))
            print(f"  {'OK ' if ok else 'FEL'} {rad}")
            if not ok:
                rap("FEL", f"stickprov {v} {kod} {p}: cell {cell} != CSV {csvv}")

# ---------------------------------------------------------------- 8. distriktssidor pa historik.val.se
print("== 8. distriktssidor (HTML) med korrekt tabelltolkning")
# de tre distrikt som granskning_2018_kontroll.py flaggade, for att prova om avvikelserna var akta
TIDIGARE_URVAL = ["14801011", "14801015", "14801041"]
os.makedirs(HTML_CACHE, exist_ok=True)


def hamta(V, kod):
    fil = os.path.join(HTML_CACHE, f"{V}_{kod}.html")
    if not os.path.exists(fil):
        url = HTML_URL.format(V=V, vd=kod[4:])
        try:
            with urllib.request.urlopen(url, timeout=30) as u:
                open(fil, "wb").write(u.read())
        except Exception as e:
            rap("INFO", f"HTML {url}: {e}")
            return None
    rat = open(fil, "rb").read().decode("utf-8", "replace")
    # sidorna ar teckenkodade med HTML-entiteter (R&ouml;stf&ouml;rdelning); avkoda dem fore sokning
    return html.unescape(rat)


def cellvarde(td):
    return re.sub(r"<[^>]+>", " ", td).replace("\xa0", " ").strip()


def tabellrader(sidtext, rubrik):
    """Returnerar listan av rader (lista av cellstrangar) i tabellen efter rubriken."""
    i = sidtext.find(rubrik)
    if i < 0:
        return []
    j = sidtext.find("<table", i)
    k = sidtext.find("</table>", j)
    if j < 0 or k < 0:
        return []
    ut = []
    for tr in re.findall(r"<tr[^>]*>(.*?)</tr>", sidtext[j:k], re.S):
        celler = [cellvarde(td) for td in re.findall(r"<t[dh][^>]*>(.*?)</t[dh]>", tr, re.S)]
        if celler:
            ut.append(celler)
    return ut


def tal(s):
    s = s.replace(" ", "").replace(" ", "")
    return int(s) if re.fullmatch(r"\d+", s) else None


for v in ("rd", "rf", "kf"):
    V = VALBOKSTAV[v]
    d_by = {r["kod"]: r for r in distrikt[v]}
    o_by = {r["kod"]: r for r in ogilt[v]}
    for kod in urval + TIDIGARE_URVAL:
        sida = hamta(V, kod)
        if sida is None:
            continue
        namn = d_by[kod]["namn"]
        huvud = tabellrader(sida, "Röstfördelning - valdistrikt")
        if not huvud:
            rap("FEL", f"HTML {V} {kod}: kunde inte tolka huvudtabellen")
            continue
        ovriga = tabellrader(sida, "Röstfördelning övriga partier - valdistrikt")
        rs = {r["parti_kalla"]: r for r in roster[v] if r["kod"] == kod}
        d, o = d_by[kod], o_by[kod]
        n_ok = n_fel = 0

        def jmf(etikett, a, b):
            global n_ok, n_fel
            if a == b:
                n_ok += 1
            else:
                n_fel += 1
                rap("FEL", f"HTML {V} {kod} {etikett}: sidan {a} != filerna {b}")

        sid_huvud = {}
        sid_ovriga = {}
        for rad in huvud:
            if len(rad) < 4 or rad[0] == "Förk." or rad[1] == "Parti":
                continue
            fork, pnamn, antal = rad[0], rad[1], tal(rad[2])
            if pnamn == "Giltiga röster":
                jmf("giltiga", antal, int(d["giltiga"]))
            elif pnamn == "Antal röstberättigade":
                jmf("rostberattigade", antal, int(d["rostberattigade"]))
            elif fork == "VDT":
                jmf("rostande", antal, int(d["rostande"]))
                jmf("valdeltagande", rad[3], f"{float(d['valdeltagande']):.2f}".replace(".", ",") + "%")
            elif fork == "OGEJ":
                jmf("OGEJ", antal or 0, int(o["ogiltiga_ej_anmalda"]))
            elif fork == "BLANK":
                jmf("BLANK", antal or 0, int(o["blanka"]))
            elif fork == "OG":
                jmf("OG", antal or 0, int(o["ogiltiga_ovriga"]))
            elif fork:
                sid_huvud[fork] = antal or 0
        for rad in ovriga:
            if len(rad) < 4 or rad[1].startswith("Totalt övriga") or rad[0] == "Förk.":
                continue
            if rad[1].startswith("Röster på partier som ej beställt"):
                sid_ovriga["__EJBESTALLT__"] = tal(rad[2]) or 0
                continue
            nyckel = rad[0] if rad[0] else "namn:" + rad[1]
            sid_ovriga[nyckel] = tal(rad[2]) or 0
        # partier i huvudtabellen utom OVR
        for p, a in sid_huvud.items():
            if p == "ÖVR":
                continue
            r = rs.get(p)
            jmf(f"parti {p}", a, int(r["roster"]) if r else None)
        # OVR pa sidan = summan i CSV av alla partier som inte star i huvudtabellen
        if "ÖVR" in sid_huvud:
            csv_ovr = sum(int(r["roster"]) for p, r in rs.items() if p not in sid_huvud or p == "ÖVR")
            jmf("ÖVR (inkl CSV-kolumnen ÖVR)", sid_huvud["ÖVR"], csv_ovr)
        # ovriga tabellen: forkortningar som finns som kolumn i kallan
        namnlosa = []
        for p, a in sid_ovriga.items():
            if p == "__EJBESTALLT__":
                continue
            if p.startswith("namn:"):
                namnlosa.append((p[5:], a))
                continue
            r = rs.get(p)
            if r is None:
                if a:
                    rap("FEL", f"HTML {V} {kod}: parti {p} har {a} roster pa sidan men saknas i CSV")
                    n_fel += 1
                else:
                    n_ok += 1  # 0 roster och ingen kolumn i kallan: forvantat
                continue
            jmf(f"ovrigt parti {p}", a, int(r["roster"]))
        # summakontroll for ovriga tabellen
        s_ov = sum(a for p, a in sid_ovriga.items())
        if "ÖVR" in sid_huvud:
            jmf("summa ovriga tabellen = ÖVR", s_ov, sid_huvud["ÖVR"])
        print(f"  HTML {V} {kod} {namn}: {n_ok} lika, {n_fel} avvikelser; "
              f"ovriga utan forkortning pa sidan: {namnlosa}")

# ---------------------------------------------------------------- 9. partinamn utan forkortning
print("== 9. partier utan namn i partier_2018.csv")
utan_namn = [(r["val"], r["parti_kalla"]) for r in partier if not r["namn"]]
print(f"  {utan_namn}")
sok = [p for v, p in utan_namn if v == "rd" and p != "ÖVR"]
for p in sok:
    kandidater = sorted(((k["partier"][p], kod) for kod, k in kalla["rd"].items()
                         if k["partier"].get(p, 0) > 0 and kod != "14800000"), reverse=True)[:2]
    for antal, kod in kandidater:
        sida = hamta("R", kod)
        if sida is None:
            continue
        ovriga = tabellrader(sida, "Röstfördelning övriga partier - valdistrikt")
        namnlosa = [(rad[1], tal(rad[2])) for rad in ovriga
                    if len(rad) >= 3 and not rad[0] and not rad[1].startswith(("Totalt", "Röster på"))]
        traff = [n for n, a in namnlosa if a == antal]
        print(f"  {p}: {antal} roster i {kod}; namnlosa rader pa sidan {namnlosa} -> {traff}")

# ------------------------------------------------- 11. foregaende val (2014) pa distriktssidorna
# Distriktssidorna har kolumnerna "Antal 2014" och "Andel 2014", men bara for distrikt som ar
# oforandrade sedan 2014 (indelning O). Bland de 22 Majornadistrikten galler det tva: 14801032
# Godhem (kommer fran 2014 ars 14801031 till 100 procent) och 14801042 Gatenhielmska (fran
# 14801042 till 100 procent). Talen jamfors mot 2014-agentens roster_2014_<val>_xls.csv.
print("== 11. foregaende val 2014 pa distriktssidorna mot 2014 ars filer")
OFORANDRADE = {"14801032": "14801031", "14801042": "14801042"}
for v in ("rd", "rf", "kf"):
    fil2014 = os.path.join(UTMAPP, f"roster_2014_{v}_xls.csv")
    if not os.path.exists(fil2014):
        rap("INFO", f"roster_2014_{v}_xls.csv saknas, jamforelsen med 2014 far goras senare")
        continue
    r14 = collections.defaultdict(dict)
    with open(fil2014, encoding="utf-8", newline="") as f:
        for r in csv.DictReader(f, delimiter=";"):
            r14[r["kod"]][r["parti"]] = int(r["roster"])
    for kod18, kod14 in OFORANDRADE.items():
        sida = hamta(VALBOKSTAV[v], kod18)
        if sida is None:
            continue
        n_ok = n_saknas = 0
        for tabellrubrik in ("Röstfördelning - valdistrikt", "Röstfördelning övriga partier - valdistrikt"):
            for rad in tabellrader(sida, tabellrubrik):
                if len(rad) < 7 or rad[0] in ("Förk.", "", "VDT", "OGEJ", "BLANK", "OG", "ÖVR"):
                    continue
                a14 = tal(rad[6])
                if a14 is None:
                    continue
                p = normalisera(rad[0])
                if p not in r14[kod14]:
                    n_saknas += 1
                    rap("INFO", f"2014 {v} {kod14}: parti {p} finns pa 2018-sidan ({a14}) men inte "
                                f"i roster_2014_{v}_xls.csv")
                elif r14[kod14][p] != a14:
                    rap("FEL", f"2014 {v} {kod14} {p}: sidan for {kod18} anger {a14}, "
                               f"roster_2014_{v}_xls.csv har {r14[kod14][p]}")
                else:
                    n_ok += 1
        print(f"  {v} {kod18} (2014: {kod14}): {n_ok} partital lika, {n_saknas} utan motsvarighet")

# ---------------------------------------------------------------- 10. sammanfattning
print("== 10. summering")
for v in ("rd", "rf", "kf"):
    g = sum(int(r["giltiga"]) for r in distrikt[v] if r["kod"] in MAJORNA)
    ro = sum(int(r["rostande"]) for r in distrikt[v] if r["kod"] in MAJORNA)
    rb = sum(int(r["rostberattigade"]) for r in distrikt[v] if r["kod"] in MAJORNA)
    t = collections.Counter()
    for r in roster[v]:
        if r["kod"] in MAJORNA:
            t[r["parti_kalla"]] += int(r["roster"])
    print(f"  22 Majornadistrikt {v}: giltiga {g}, rostande {ro}, rostberattigade {rb}, "
          f"valdeltagande {100*ro/rb:.2f}; {dict(t.most_common(6))}")

print(f"=== FEL {len(fel)}, VARNING {len(varning)}, INFO {len(info)}")
for x in fel:
    print("  FEL:", x)
for x in varning:
    print("  VARNING:", x)
for x in info:
    print("  INFO:", x)
sys.exit(1 if fel else 0)
